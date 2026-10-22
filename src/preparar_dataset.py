"""Entrega 02: preparación reproducible Bronze -> Silver -> Gold en un solo archivo.

Ejecutar desde la raíz: .venv/Scripts/python.exe src/preparar_dataset.py
Revisar entradas:       .venv/Scripts/python.exe src/preparar_dataset.py --comprobar-entradas

Requiere Python 3.12, requirements.txt y la extensión oficial spatial de DuckDB.
El bloque ENTRADAS identifica las evidencias locales que acompañan a los originales.
No descarga fuentes, no importa otros scripts ni consume Silver/Gold anteriores.

Lectura recomendada: main() muestra el recorrido completo; las secciones # %%
contienen los lectores, transformaciones y controles utilizados en cada etapa.
Se conservan las reglas de la preparación validada, incluidos nulos, incidencias,
relaciones múltiples y escenarios. Las decisiones humanas se verifican como datos.

Cada ejecución escribe Bronze, 48 tablas Silver, 13 Gold y un manifiesto bajo
resultados/entrega02/<ejecución>/. Las tablas son Parquet/GeoParquet EPSG:25830.
Los originales, los scripts históricos y sus indicadores actual.json se conservan.
El manifiesto final solo aparece tras completar controles, exportaciones y relectura.
"""

import argparse
import csv
import hashlib
import io
import json
import math
import re
import sqlite3
import struct
import tempfile
import unicodedata
from collections import Counter, defaultdict
from contextlib import closing, contextmanager
from datetime import date, datetime, timezone
from decimal import Decimal
from html import unescape
from html.parser import HTMLParser
from pathlib import Path, PurePosixPath
from urllib.parse import parse_qs, urlparse
from xml.etree import ElementTree

import duckdb
from openpyxl import load_workbook


# %% 1. Configuración: rutas relativas a la raíz del proyecto

RAIZ = Path(__file__).resolve().parents[1]
ENTRADAS = {
    'catalogo': 'config/catalogo_bronze.json',
    'recarga_principal': 'resultados/revision_recarga/20261018_ampliacion/miteco_exportacion_actual.csv',
    'metadatos_recarga': 'resultados/investigacion_datum_miteco/20261016/catalogo.json',
    'metadatos_capacidad': 'resultados/investigacion_ide/20261019/metadatos.json',
    'decisiones_recarga': 'resultados/revision_recarga/20261018_revision_manual/decisiones.json',
}


# %% 2. Recorrido completo: las transformaciones se definen por fuente más abajo

def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('--comprobar-entradas', action='store_true',
                        help='Comprueba entorno, originales y metadatos sin transformar tablas.')
    args = parser.parse_args()
    bronze, archivos, metadatos, adquisicion, origenes = comprobar_entradas()
    if args.comprobar_entradas:
        print(f'Entradas comprobadas: {len(bronze["archivos"])} archivos Bronze y evidencias complementarias.')
        return

    version = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    salida = RAIZ / 'resultados/entrega02' / version
    salida.mkdir(parents=True, exist_ok=False)
    guardar_json(salida / 'bronze/manifiesto.json', bronze)
    tablas, controles = {}, {}

    @contextmanager
    def etapa(bloque):
        print(f'\nPreparando {bloque}...', flush=True)
        with tempfile.TemporaryDirectory(dir=salida) as temporal:
            temporal = Path(temporal)
            with conectar(temporal) as c:
                yield c, temporal
        print(f'Completado {bloque}.', flush=True)

    def cargar(c, bloque, nombres=None, prefijo='', alias=None):
        for nombre in nombres if nombres is not None else tablas[bloque]:
            entrada = tablas[bloque][nombre]
            path = RAIZ / entrada['ruta']
            comprobar_archivo(RAIZ, path, entrada)
            destino = (alias or {}).get(nombre, prefijo + nombre)
            c.execute(f'CREATE TEMP VIEW {identificador_sql(destino)} AS '
                      f'SELECT * FROM read_parquet({texto_sql(str(path))})')

    def territorio(c):
        cargar(c, 'silver/base_territorial', alias={'recintos_municipales': 'recintos'})
        c.execute('CREATE TEMP TABLE ambito_cv AS SELECT ST_Union_Agg(geometria) geometria FROM recintos')

    def exportar(c, bloque, claves, calidad):
        tablas[bloque] = exportar_tablas(c, salida / bloque, claves)
        controles[bloque] = calidad

    # Silver: conservar entidades, periodos, estados y valores originales por fuente.
    with etapa('silver/base_territorial') as (c, temporal):
        municipios, total_nacional = referencia_municipal_reference(RAIZ / TERRITORIO_INE)
        if Counter(codigo[:2] for codigo in municipios) != {'03': 141, '12': 135, '46': 266}:
            raise ValueError('La referencia no contiene los 542 municipios previstos')
        calidad = territorio_build_base(c, municipios, RAIZ / TERRITORIO_IGN,
                                      RAIZ / TERRITORIO_HITS, origenes[TERRITORIO_IGN])
        calidad['municipios_nacionales'] = total_nacional
        exportar(c, 'silver/base_territorial',
                 dict.fromkeys(('municipios', 'recintos_municipales'), ['codigo_municipio']), calidad)

    with etapa('silver/demografia_parque') as (c, temporal):
        cargar(c, 'silver/base_territorial', ['municipios'])
        demografia_prepare_municipal_park(c, RAIZ / DEMOGRAFIA_XLSX, municipios)
        demografia_prepare_ine(c, RAIZ / DEMOGRAFIA_POP, RAIZ / DEMOGRAFIA_INCOME,
                              municipios, temporal, origenes)
        print('Leyendo el TXT nacional DGT secuencialmente, con trabajo en disco...', flush=True)
        demografia_prepare_detail(c, RAIZ / DEMOGRAFIA_TXT, origenes[DEMOGRAFIA_TXT])
        exportar(c, 'silver/demografia_parque', {**DEMOGRAFIA_KEYS, 'incidencias': None},
                 demografia_validate(c, municipios))

    with etapa('silver/recarga') as (c, temporal):
        recarga_prepare(c, RAIZ / RECARGA_SOURCE, temporal)
        exportar(c, 'silver/recarga', {**RECARGA_KEYS, 'incidencias': None}, recarga_validate(c))

    with etapa('silver/trafico_red') as (c, temporal):
        trafico_prepare_traffic(c, RAIZ / TRAFICO_GVA, temporal, origenes[TRAFICO_GVA])
        for provincia, carpeta in TRAFICO_ROADS.items():
            ruta = carpeta + '/red_viaria.gpkg'
            trafico_prepare_road(c, RAIZ / ruta, provincia, temporal, origenes[ruta])
        exportar(c, 'silver/trafico_red', {**TRAFICO_KEYS, 'incidencias': None}, trafico_validate(c))

    with etapa('silver/poi_osm') as (c, temporal):
        territorio(c)
        recuentos = osm_prepare(c, RAIZ / OSM_SOURCE, temporal, origenes[OSM_SOURCE])
        osm_classify(c)
        osm_correspondences(c)
        exportar(c, 'silver/poi_osm', {**OSM_KEYS, 'incidencias': None}, osm_validate(c, recuentos))

    with etapa('silver/capacidad_demanda') as (c, temporal):
        capacidad_prepare(c, RAIZ / CAPACIDAD_SOURCE, temporal, origenes[CAPACIDAD_SOURCE])
        exportar(c, 'silver/capacidad_demanda',
                 {'capacidad_demanda': ['punto_conexion_id'], 'incidencias': None}, capacidad_validate(c))

    with etapa('silver/recarga_espacial') as (c, temporal):
        territorio(c)
        cargar(c, 'silver/recarga', RECARGA_KEYS)
        recarga_geo_prepare(c, temporal)
        exportar(c, 'silver/recarga_espacial', {**RECARGA_GEO_KEYS, 'incidencias_espaciales': None},
                 recarga_geo_validate(c))

    with etapa('silver/capacidad_demanda_espacial') as (c, temporal):
        territorio(c)
        cargar(c, 'silver/capacidad_demanda', ['capacidad_demanda'])
        capacidad_geo_prepare(c, metadatos)
        exportar(c, 'silver/capacidad_demanda_espacial', CAPACIDAD_GEO_KEYS, capacidad_geo_validate(c))

    with etapa('silver/validacion_conjunta') as (c, temporal):
        for bloque in list(tablas):
            cargar(c, bloque, prefijo=bloque.split('/')[-1] + '__')
        calidad = integracion_integrity(c)
        decisiones, evidencias = integracion_provider_evidence(RAIZ, RAIZ / ENTRADAS['decisiones_recarga'], c)
        archivos.extend(evidencias)
        calidad['recarga'] = integracion_review_charging(c, decisiones)
        calidad['relaciones'] = integrar_relaciones(c)
        calidad['cobertura'] = integracion_coverage(c)
        crear_estado_fuentes(c)
        exportar(c, 'silver/validacion_conjunta', INTEGRACION_OUTPUT_KEYS, calidad)

    # La versión principal hereda geometría solo si los atributos coinciden exactamente.
    # Las ausencias históricas y las decisiones documentales se conservan por separado.
    with etapa('silver/recarga_versionada') as (c, temporal):
        territorio(c)
        cargar(c, 'silver/recarga', RECARGA_KEYS, prefijo='anterior_')
        cargar(c, 'silver/recarga_espacial', ['recarga_instalaciones_geo', 'recarga_municipios_candidatos'],
               alias={'recarga_instalaciones_geo': 'geometria_anterior',
                      'recarga_municipios_candidatos': 'candidatos_anteriores'})
        cargar(c, 'silver/validacion_conjunta', ['recarga_municipio_validado'],
               alias={'recarga_municipio_validado': 'admisiones_anteriores'})
        recarga_prepare(c, RAIZ / ENTRADAS['recarga_principal'], temporal)
        calidad = recarga_validate(c)
        if calidad['incidencias']:
            raise ValueError('Revisar incidencias tabulares de la recarga principal')
        recarga_versionada_inherit_geometry(c)
        calidad['espacial'] = recarga_geo_validate(c)
        crear_vistas_revision(c)
        decisiones, evidencias = integracion_provider_evidence(RAIZ, RAIZ / ENTRADAS['decisiones_recarga'], c)
        archivos.extend(evidencias)
        calidad['revision'] = integracion_review_charging(c, decisiones)
        exigir_cero(c, '''SELECT count(*) FROM (
            (SELECT * FROM recarga_municipio_validado EXCEPT ALL SELECT a.* FROM admisiones_anteriores a
             JOIN recarga_instalaciones n USING(instalacion_id)) UNION ALL
            (SELECT a.* FROM admisiones_anteriores a JOIN recarga_instalaciones n USING(instalacion_id)
             EXCEPT ALL SELECT * FROM recarga_municipio_validado))''', 'Admisiones incompatibles entre versiones')
        recarga_versionada_preserve_absences(c)
        exportar(c, 'silver/recarga_versionada', RECARGA_VERSIONADA_KEYS, calidad)

    # Gold: conjunto relacional final, sin pesos, ranking o entrenamiento.
    with etapa('gold/contexto_municipal') as (c, temporal):
        for bloque, nombres in GOLD_MUNICIPAL_BLOCKS.items():
            cargar(c, 'silver/' + bloque, nombres)
        gold_municipal_prepare(c, version, adquisicion, origenes['recarga_principal'], origenes[RECARGA_SOURCE])
        exportar(c, 'gold/contexto_municipal', GOLD_MUNICIPAL_KEYS, gold_municipal_validate(c, version))

    with etapa('gold/candidatos') as (c, temporal):
        for bloque, nombres in GOLD_CANDIDATOS_SOURCES.items():
            cargar(c, bloque, nombres)
        gold_candidatos_prepare(c, version)
        gold_candidatos_proximity(c, version)
        exportar(c, 'gold/candidatos', GOLD_CANDIDATOS_KEYS, gold_candidatos_validate(c, version))

    with etapa('gold/controles_espaciales') as (c, temporal):
        cargar(c, 'gold/contexto_municipal')
        cargar(c, 'gold/candidatos')
        gold_controles_prepare(c, version)
        exportar(c, 'gold/controles_espaciales', GOLD_CONTROLES_KEYS, gold_controles_validate(c, version))

    with etapa('gold/variables') as (c, temporal):
        for bloque in ('gold/contexto_municipal', 'gold/candidatos', 'gold/controles_espaciales'):
            cargar(c, bloque)
        gold_variables_prepare(c, version)
        exportar(c, 'gold/variables', GOLD_VARIABLES_KEYS, gold_variables_validate(c, version))

    # El manifiesto marca una ejecución terminada; una interrupción deja su carpeta sin validar.
    for entrada in archivos:
        comprobar_archivo(RAIZ, RAIZ / entrada['ruta'], entrada)
    manifiesto = {'version_script': 1, 'ejecucion': version, 'estado': 'preparada_con_limitaciones',
                  'entradas_verificadas': list({x['ruta']: x for x in archivos}.values()),
                  'metadatos_capacidad': metadatos, 'tablas': tablas, 'controles': controles,
                  'entorno': {'duckdb': duckdb.__version__, 'memoria_motor': '1GB', 'hilos': 2},
                  'limitaciones': ['Distintivo 0 aproxima demanda municipal; no demanda local observada.',
                      'Asignaciones pendientes y ausencias no equivalen a inexistencia o cierre.',
                      'POI requieren revisión local; distancias planas no son rutas ni suficiencia.',
                      'Capacidad i-DE orientativa, no aditiva ni garantía de conexión.']}
    guardar_json(salida / 'manifiesto.json', manifiesto)
    print(f'\nPreparación completada: {salida.relative_to(RAIZ).as_posix()}/manifiesto.json', flush=True)



# %% Utilidades de SQL

def identificador_sql(name):
    return '"' + name.replace('"', '""') + '"'


def texto_sql(value):
    return "'" + str(value).replace("'", "''") + "'"


def consultar(connection, sql):
    cursor = connection.execute(sql)
    return [dict(zip([c[0] for c in cursor.description], row)) for row in cursor.fetchall()]


# %% Bronze: integridad de los originales

BRONZE_RAW_ROOTS = ('Fuentes', 'datos/bronze')


def bronze_raw_path(root, relative):
    if not isinstance(relative, str) or '\\' in relative or ':' in relative:
        raise ValueError(f'Ruta relativa no válida: {relative!r}')
    parts = PurePosixPath(relative).parts
    if not parts or '..' in parts or PurePosixPath(relative).is_absolute():
        raise ValueError(f'Ruta fuera de Bronze: {relative!r}')
    path = root / relative
    if not any(path.resolve().is_relative_to((root / folder).resolve()) for folder in BRONZE_RAW_ROOTS):
        raise ValueError(f'Ruta fuera de Bronze: {relative!r}')
    if not path.resolve().is_relative_to(root.resolve()):
        raise ValueError(f'Ruta fuera del proyecto: {relative!r}')
    return path


def huella(path):
    before = path.stat()
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        while block := stream.read(4 * 1024 * 1024):
            digest.update(block)
    after = path.stat()
    if (before.st_size, before.st_mtime_ns) != (after.st_size, after.st_mtime_ns):
        raise RuntimeError(f'El archivo cambió durante la lectura: {path.name}')
    return {'bytes': before.st_size, 'sha256': digest.hexdigest(),
            'mtime_ns': before.st_mtime_ns,
            'modificacion_local_utc': datetime.fromtimestamp(before.st_mtime, timezone.utc).isoformat()}


def bronze_registered_files(root, catalog):
    resources = catalog['recursos']
    ids = [resource['id'] for resource in resources]
    if not ids or len(set(ids)) != len(ids):
        raise ValueError('El catálogo está vacío o repite identificadores')
    registered = {}
    for resource in resources:
        if not resource['entradas']:
            raise ValueError(f"Recurso sin entradas: {resource['id']}")
        for entry in resource['entradas']:
            pattern = entry['patron']
            bronze_raw_path(root, pattern)
            if entry['rol'] not in ('datos', 'documentacion'):
                raise ValueError(f'Rol no válido: {entry["rol"]}')
            matches = sorted(root.glob(pattern))
            if len(matches) != entry.get('cantidad', 1):
                raise ValueError(f'Número de archivos inesperado para {pattern}: {len(matches)}')
            for path in matches:
                relative = path.relative_to(root).as_posix()
                bronze_raw_path(root, relative)
                if not path.is_file() or relative in registered:
                    raise ValueError(f'Entrada repetida o no regular: {relative}')
                registered[relative] = {'recurso_id': resource['id'], 'rol': entry['rol']}
    actual = {path.relative_to(root).as_posix() for folder in BRONZE_RAW_ROOTS
              for path in (root / folder).rglob('*') if path.is_file()}
    if actual != set(registered):
        raise ValueError(f'Archivos sin catalogar: {sorted(actual - set(registered))}')
    return registered


def bronze_build_manifest(root, catalog):
    files = bronze_registered_files(root, catalog)
    records = []
    for relative, metadata in sorted(files.items()):
        print(f'Registrando {relative}', flush=True)
        records.append({'ruta': relative, **metadata, **huella(bronze_raw_path(root, relative))})
    # Comprobar también archivos añadidos mientras se calculaban las huellas.
    if bronze_registered_files(root, catalog) != files:
        raise RuntimeError('Cambió el conjunto de entradas durante el registro')
    return {'version_manifiesto': 1, 'fecha_registro_utc': datetime.now(timezone.utc).isoformat(),
            'catalogo': catalog, 'archivos': records}


# %% Referencia municipal INE y recintos IGN

REFERENCIA_PROVINCES = ('03', '12', '46')


def referencia_municipal_reference(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    codes = {}
    national_count = 0
    try:
        for sheet in workbook:
            rows = sheet.iter_rows(values_only=True)
            next(rows)
            next(rows)
            if tuple(next(rows)) != ('CPRO', 'CMUN', 'DC', 'NOMBRE'):
                raise ValueError(f'Esquema municipal INE inesperado: {sheet.title}')
            for province, municipality, check_digit, name in rows:
                if province is None and municipality is None:
                    continue
                # La edición oficial mezcla texto y un código numérico de tres cifras.
                if type(province) is int:
                    province = str(province)
                if type(municipality) is int:
                    municipality = str(municipality)
                if (not isinstance(province, str) or len(province) != 2 or not province.isdigit()
                        or not isinstance(municipality, str) or len(municipality) != 3 or not municipality.isdigit()):
                    raise ValueError('Códigos INE sin longitud completa')
                national_count += 1
                if province not in REFERENCIA_PROVINCES:
                    continue
                code = province + municipality
                if code in codes or municipality == '000' or not isinstance(name, str) or not name.strip():
                    raise ValueError(f'Municipio repetido o no válido: {code}')
                codes[code] = name
    finally:
        workbook.close()
    if not codes:
        raise ValueError('La referencia INE no contiene municipios valencianos')
    return codes, national_count


def referencia_geometry_reference(path, hits_path, expected):
    collection = json.loads(path.read_text(encoding='utf-8'))
    matched = int(ElementTree.parse(hits_path).getroot().attrib['numberMatched'])
    features = collection['features']
    if collection['type'] != 'FeatureCollection' or len(features) != matched:
        raise ValueError('Respuesta WFS incompleta respecto al recuento del proveedor')
    codes, levels, geometries = [], Counter(), Counter()
    for feature in features:
        properties = feature['properties']
        national_code = properties['nationalCode']
        if not isinstance(national_code, str) or len(national_code) != 11 or not national_code.isdigit() or not national_code.startswith('3410'):
            raise ValueError('Unidad fuera del ámbito territorial solicitado')
        level = properties['nationalLevel']['href'].rsplit('/', 1)[-1]
        levels[level] += 1
        geometry = feature['geometry']
        if not geometry or geometry['type'] not in ('Polygon', 'MultiPolygon') or not geometry['coordinates']:
            raise ValueError('Unidad sin geometría poligonal informada')
        geometries[geometry['type']] += 1
        if level == '4thOrder':
            code = national_code[-5:]
            if code[:2] not in REFERENCIA_PROVINCES or national_code[4:6] != code[:2]:
                raise ValueError('Código nacional IGN incoherente con su municipio')
            codes.append(code)
    if len(set(codes)) != len(codes):
        raise ValueError('Códigos municipales repetidos en la cartografía')
    if set(codes) != set(expected):
        raise ValueError(f'Cobertura IGN/INE distinta: faltan {sorted(set(expected) - set(codes))}; adicionales {sorted(set(codes) - set(expected))}')
    return {'unidades': len(features), 'recuento_proveedor': matched, 'niveles': dict(levels),
            'geometrias': dict(geometries), 'municipios': len(codes),
            'faltan': [], 'adicionales': [], 'validacion_topologica': False}


# %% Silver: base territorial

TERRITORIO_INE = 'datos/bronze/ine_municipios/2025/25codmun.xlsx'


TERRITORIO_IGN = 'datos/bronze/ign_unidades_cv/2026-10-15/unidades_cv.geojson'


TERRITORIO_HITS = 'datos/bronze/ign_unidades_cv/2026-10-15/recuento.xml'


def territorio_build_base(connection, codes, geojson, hits, origin_id):
    profile = referencia_geometry_reference(geojson, hits, codes)
    provinces = Counter(code[:2] for code in codes)
    expected_levels = {'2ndOrder': 1, '3rdOrder': len(provinces), '4thOrder': len(codes)}
    if profile['niveles'] != expected_levels:
        raise ValueError('Niveles administrativos inesperados en la respuesta IGN')
    connection.execute('CREATE TEMP TABLE municipios (codigo_municipio VARCHAR PRIMARY KEY, codigo_provincia VARCHAR, nombre_municipio VARCHAR, fecha_referencia DATE)')
    connection.executemany('INSERT INTO municipios VALUES (?,?,?,?)',
                           [(code, code[:2], name, date(2025, 1, 1)) for code, name in sorted(codes.items())])
    connection.execute("""
        CREATE TEMP TABLE unidades_ign AS SELECT id registro_origen,
            nationalCode codigo_nacional_ign,
            json_extract_string(nationalLevel,'$.href') nivel,
            json_extract_string(name,'$.GeographicalName.spelling.SpellingOfName.text') nombre_municipio_ign,
            geom geometria_origen
        FROM ST_Read(?)
    """, [str(geojson)])
    if connection.execute('SELECT count(*) FROM unidades_ign').fetchone()[0] != profile['unidades']:
        raise ValueError('El lector espacial no conserva el recuento Bronze')
    invalid = consultar(connection, """
        SELECT codigo_nacional_ign FROM unidades_ign
        WHERE geometria_origen IS NULL OR ST_IsEmpty(geometria_origen)
           OR NOT ST_IsValid(geometria_origen)
           OR ST_GeometryType(geometria_origen) NOT IN ('POLYGON','MULTIPOLYGON')
           OR ST_HasZ(geometria_origen) OR ST_HasM(geometria_origen)
           OR ST_XMin(geometria_origen) NOT BETWEEN -2 AND 1
           OR ST_XMax(geometria_origen) NOT BETWEEN -2 AND 1
           OR ST_YMin(geometria_origen) NOT BETWEEN 37 AND 42
           OR ST_YMax(geometria_origen) NOT BETWEEN 37 AND 42
    """)
    if invalid:
        raise ValueError(f'Geometrías IGN no admitidas, sin reparación automática: {invalid}')
    missing = connection.execute("""
        SELECT count(*) FROM unidades_ign WHERE registro_origen IS NULL
            OR trim(registro_origen)='' OR codigo_nacional_ign IS NULL
            OR nombre_municipio_ign IS NULL OR trim(nombre_municipio_ign)=''
    """).fetchone()[0]
    unique = connection.execute('SELECT count(DISTINCT registro_origen),count(DISTINCT codigo_nacional_ign) FROM unidades_ign').fetchone()
    if missing or unique != (profile['unidades'], profile['unidades']):
        raise ValueError('Identificadores o nombres IGN ausentes o repetidos')
    connection.execute("""
        ALTER TABLE unidades_ign ADD COLUMN geometria GEOMETRY('EPSG:25830')
    """)
    connection.execute("""
        UPDATE unidades_ign SET geometria=ST_Transform(geometria_origen,
            'OGC:CRS84','EPSG:25830',always_xy:=true)::GEOMETRY('EPSG:25830')
    """)
    bad = consultar(connection, """
        SELECT codigo_nacional_ign FROM unidades_ign WHERE geometria IS NULL
            OR NOT ST_IsValid(geometria) OR ST_IsEmpty(geometria)
            OR NOT isfinite(ST_Area(geometria)) OR ST_Area(geometria)<=0
    """)
    if bad:
        raise ValueError(f'Geometrías proyectadas no admitidas: {bad}')
    connection.execute("""
        CREATE TEMP TABLE recintos_municipales AS SELECT
            right(codigo_nacional_ign,5) codigo_municipio,
            codigo_nacional_ign,nombre_municipio_ign,geometria,
            ?::VARCHAR origen_id,registro_origen
        FROM unidades_ign WHERE ends_with(nivel,'/4thOrder')
    """, [origin_id])
    observed = connection.execute('SELECT codigo_municipio FROM recintos_municipales').fetchall()
    if len(observed) != len(codes) or {row[0] for row in observed} != set(codes):
        raise ValueError('Correspondencia entre municipios y recintos distinta de uno a uno')
    edges = connection.execute('SELECT ST_CoverageInvalidEdges(list(geometria)) IS NULL FROM recintos_municipales').fetchone()[0]
    if not edges:
        raise ValueError('Los recintos municipales tienen aristas de cobertura incoherentes')
    connection.execute('CREATE TEMP TABLE union_municipal AS SELECT ST_Union_Agg(geometria) geometria FROM recintos_municipales')
    area_sum = connection.execute('SELECT sum(ST_Area(geometria)) FROM recintos_municipales').fetchone()[0]
    area_union = connection.execute('SELECT ST_Area(geometria) FROM union_municipal').fetchone()[0]
    if abs(area_sum - area_union) > 1:
        raise ValueError('Área de solapamiento municipal superior a la tolerancia numérica de 1 m²')
    connection.execute("""
        CREATE TEMP TABLE uniones_provinciales AS SELECT substr(codigo_municipio,1,2) codigo_provincia,
            ST_Union_Agg(geometria) geometria FROM recintos_municipales GROUP BY 1
    """)
    coverage = consultar(connection, """
        SELECT 'comunidad' nivel,b.codigo_nacional_ign,
            ST_Area(ST_Difference(a.geometria,b.geometria)) fuera_del_padre_m2,
            ST_Area(ST_Difference(b.geometria,a.geometria)) no_cubierto_m2
        FROM union_municipal a,unidades_ign b WHERE ends_with(b.nivel,'/2ndOrder')
        UNION ALL
        SELECT 'provincia',b.codigo_nacional_ign,
            ST_Area(ST_Difference(a.geometria,b.geometria)),
            ST_Area(ST_Difference(b.geometria,a.geometria))
        FROM uniones_provinciales a JOIN unidades_ign b
          ON a.codigo_provincia=substr(b.codigo_nacional_ign,5,2)
        WHERE ends_with(b.nivel,'/3rdOrder') ORDER BY 1,2
    """)
    if len(coverage) != len(provinces) + 1 or any(row['fuera_del_padre_m2'] > 1 or row['no_cubierto_m2'] > 1 for row in coverage):
        raise ValueError(f'Cobertura municipal distinta de sus niveles superiores: {coverage}')
    differences = consultar(connection, """
        SELECT codigo_municipio,nombre_municipio,nombre_municipio_ign
        FROM municipios JOIN recintos_municipales USING(codigo_municipio)
        WHERE nombre_municipio<>nombre_municipio_ign ORDER BY codigo_municipio
    """)
    return {'unidades_bronze': profile['unidades'], 'niveles': profile['niveles'],
            'municipios': len(codes), 'por_provincia': dict(sorted(provinces.items())),
            'claves_y_correspondencia_uno_a_uno': True,
            'geometrias_validas_origen': len(codes), 'geometrias_validas_salida': len(codes),
            'geometrias_nulas': 0, 'geometrias_vacias': 0, 'reparaciones': 0,
            'aristas_cobertura_invalidas': False, 'tolerancia_aristas': 0,
            'area_suma_municipios_m2': area_sum, 'area_union_municipios_m2': area_union,
            'diferencia_suma_union_m2': area_sum - area_union,
            'tolerancia_numerica_areas_m2': 1, 'contraste_niveles_superiores': coverage,
            'nombres_distintos': differences}


# %% Silver: parque DGT, población y renta

DEMOGRAFIA_TXT = 'Fuentes/parque_vehiculos_202512.txt'


DEMOGRAFIA_XLSX = 'datos/bronze/dgt_municipal/2025/DatosMunicipalesGeneral_2025.xlsx'


DEMOGRAFIA_POP = 'Fuentes/Poblacion.csv'


DEMOGRAFIA_INCOME = 'Fuentes/Renta.csv'


DEMOGRAFIA_PROVINCES = ('03', '12', '46')


DEMOGRAFIA_SCHEMAS = {
    'poblacion_municipal': {'codigo_municipio': 'VARCHAR', 'sexo': 'VARCHAR', 'anio': 'SMALLINT',
        'poblacion': 'BIGINT', 'estado_valor': 'VARCHAR', 'valor_origen': 'VARCHAR', 'registro_origen': 'BIGINT'},
    'renta_territorial': {'nivel_territorial': 'VARCHAR', 'codigo_territorio': 'VARCHAR',
        'codigo_municipio': 'VARCHAR', 'codigo_distrito': 'VARCHAR', 'codigo_seccion': 'VARCHAR',
        'indicador': 'VARCHAR', 'anio': 'SMALLINT', 'valor_eur': 'DECIMAL(18,2)',
        'estado_valor': 'VARCHAR', 'valor_origen': 'VARCHAR', 'registro_origen': 'BIGINT'},
    'incidencias': {'tabla': 'VARCHAR', 'origen_id': 'VARCHAR', 'registro_origen': 'BIGINT',
        'codigo_municipio': 'VARCHAR', 'campo': 'VARCHAR', 'valor_origen': 'VARCHAR', 'motivo': 'VARCHAR'},
}


DEMOGRAFIA_SCHEMAS['renta_codigos_historicos'] = {**DEMOGRAFIA_SCHEMAS['renta_territorial'], 'motivo_separacion': 'VARCHAR'}


DEMOGRAFIA_KEYS = {'parque_municipal': ['codigo_municipio', 'fecha_referencia'],
        'parque_sin_municipio': ['codigo_agrupacion_origen', 'fecha_referencia'],
        'parque_detalle_observado': ['origen_id', 'registro_origen'],
        'poblacion_municipal': ['codigo_municipio', 'sexo', 'anio'],
        'renta_territorial': ['nivel_territorial', 'codigo_territorio', 'indicador', 'anio'],
        'renta_codigos_historicos': ['nivel_territorial', 'codigo_territorio', 'indicador', 'anio']}


def demografia_ine_number(raw, integer=False):
    text = raw.strip()
    if text in ('', '.', '..', '...'):
        return None, 'no_informado'
    if not re.fullmatch(r'-?(?:[0-9]+|[0-9]{1,3}(?:\.[0-9]{3})+)(?:,[0-9]+)?', text):
        return None, 'error_conversion'
    value = Decimal(text.replace('.', '').replace(',', '.'))
    if integer:
        if value != value.to_integral_value() or not 0 <= value <= 2**63 - 1:
            return None, 'error_conversion'
        return int(value), 'informado'
    if (',' in text and len(text.split(',')[1]) > 2) or abs(value) >= Decimal('10000000000000000'):
        return None, 'error_conversion'
    return value, 'informado'


def demografia_territory_code(raw, width):
    match = re.match(r'^([0-9]{' + str(width) + r'})(?:\s|$)', raw)
    if not match:
        raise ValueError(f'Código territorial de longitud {width} no válido: {raw!r}')
    return match[1]


def demografia_csv_rows(path, headers):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.reader(stream, delimiter=';', strict=True)
        if next(reader) != headers:
            raise ValueError(f'Cabecera inesperada: {path.name}')
        for ordinal, values in enumerate(reader, 1):
            if len(values) != len(headers):
                raise ValueError(f'{path.name}: ancho inesperado en registro {ordinal}')
            yield ordinal, dict(zip(headers, values))


def demografia_prepare_ine(connection, population, income, codes, temporary, origins):
    counts = Counter()
    with (temporary / 'incidencias.csv').open('w', encoding='utf-8', newline='') as inc_stream:
        incidents = csv.writer(inc_stream)
        incidents.writerow(DEMOGRAFIA_SCHEMAS['incidencias'])
        with (temporary / 'poblacion_municipal.csv').open('w', encoding='utf-8', newline='') as target:
            writer = csv.writer(target)
            writer.writerow(DEMOGRAFIA_SCHEMAS['poblacion_municipal'])
            for ordinal, row in demografia_csv_rows(population, ['Total Nacional', 'Municipios', 'Sexo', 'Periodo', 'Total']):
                counts['poblacion_nacional'] += 1
                if row['Municipios'][:2] not in DEMOGRAFIA_PROVINCES:
                    continue
                code = demografia_territory_code(row['Municipios'], 5)
                if code not in codes or row['Sexo'] not in ('Total', 'Hombres', 'Mujeres'):
                    raise ValueError(f'Población: territorio o sexo no admitido en registro {ordinal}')
                if not re.fullmatch(r'[0-9]{4}', row['Periodo']):
                    raise ValueError(f'Periodo de población no válido: registro {ordinal}')
                value, state = demografia_ine_number(row['Total'], integer=True)
                writer.writerow([code, row['Sexo'], int(row['Periodo']), value, state, row['Total'], ordinal])
                counts['poblacion_cv'] += 1
                if state == 'error_conversion':
                    incidents.writerow(['poblacion_municipal', origins[DEMOGRAFIA_POP], ordinal, code, 'poblacion', row['Total'], state])
        with (temporary / 'renta_territorial.csv').open('w', encoding='utf-8', newline='') as target, (temporary / 'renta_codigos_historicos.csv').open('w', encoding='utf-8', newline='') as historic:
            writer, historical_writer = csv.writer(target), csv.writer(historic)
            writer.writerow(DEMOGRAFIA_SCHEMAS['renta_territorial'])
            historical_writer.writerow(DEMOGRAFIA_SCHEMAS['renta_codigos_historicos'])
            headers = ['Municipios', 'Distritos', 'Secciones', 'Indicadores de renta media', 'Periodo', 'Total']
            for ordinal, row in demografia_csv_rows(income, headers):
                counts['renta_nacional'] += 1
                if not any(row[field][:2] in DEMOGRAFIA_PROVINCES for field in ('Municipios', 'Distritos', 'Secciones')):
                    continue
                code = demografia_territory_code(row['Municipios'], 5)
                if code not in codes and code != '12066':
                    raise ValueError(f'Renta: municipio ajeno a la referencia en registro {ordinal}: {code}')
                district = demografia_territory_code(row['Distritos'], 7) if row['Distritos'] else None
                section = demografia_territory_code(row['Secciones'], 10) if row['Secciones'] else None
                if (district and not district.startswith(code)) or (section and (not district or not section.startswith(district))):
                    raise ValueError(f'Jerarquía territorial de renta incoherente: registro {ordinal}')
                if not re.fullmatch(r'[0-9]{4}', row['Periodo']) or not row['Indicadores de renta media'].strip():
                    raise ValueError(f'Clave de renta incompleta: registro {ordinal}')
                level, territory = ('seccion', section) if section else ('distrito', district) if district else ('municipio', code)
                value, state = demografia_ine_number(row['Total'])
                values = [level, territory, code, district, section, row['Indicadores de renta media'],
                          int(row['Periodo']), value, state, row['Total'], ordinal]
                if code == '12066':
                    historical_writer.writerow(values + ['codigo_historico_fuera_de_referencia_2025'])
                    counts['renta_historica'] += 1
                else:
                    writer.writerow(values)
                    counts['renta_cv_vigente'] += 1
                if state == 'error_conversion':
                    incidents.writerow(['renta_codigos_historicos' if code == '12066' else 'renta_territorial',
                                        origins[DEMOGRAFIA_INCOME], ordinal, code, 'valor_eur', row['Total'], state])
    for table, schema in DEMOGRAFIA_SCHEMAS.items():
        # El CSV temporal usa vacío como NULL; el texto original vacío se restaura expresamente.
        connection.execute(f"CREATE TEMP TABLE {identificador_sql(table)} AS SELECT * REPLACE(coalesce(valor_origen,'') AS valor_origen) FROM read_csv(?,columns=?,header=true,delim=',',strict_mode=true,ignore_errors=false)", [str(temporary / f'{table}.csv'), schema])
    for table, counter in [('poblacion_municipal', 'poblacion_cv'),
                           ('renta_territorial', 'renta_cv_vigente'),
                           ('renta_codigos_historicos', 'renta_historica')]:
        if connection.execute(f'SELECT count(*) FROM {identificador_sql(table)}').fetchone()[0] != counts[counter]:
            raise ValueError(f'Pérdida de registros INE: {table}')
    return dict(counts)


def demografia_prepare_municipal_park(connection, path, codes):
    connection.execute('CREATE TEMP TABLE parque_municipal (codigo_municipio VARCHAR,fecha_referencia DATE,parque_total BIGINT,vehiculos_distintivo_0 BIGINT,electrificado_origen BIGINT)')
    connection.execute('CREATE TEMP TABLE parque_sin_municipio (codigo_agrupacion_origen VARCHAR,codigo_provincia VARCHAR,fecha_referencia DATE,parque_total BIGINT,vehiculos_distintivo_0 BIGINT,electrificado_origen BIGINT)')
    workbook = load_workbook(path, read_only=True, data_only=True)
    municipalities, unspecified, seen = [], [], set()
    try:
        rows = workbook['DatosMunicipalesGeneral_2025'].iter_rows(values_only=True)
        headers = next(rows)
        fields = ('Código INE', 'Parque Total', 'Distintivo 0', 'Electrificado')
        if len(set(headers)) != len(headers) or not set(fields).issubset(headers):
            raise ValueError('Cabecera municipal DGT inesperada')
        for ordinal, values in enumerate(rows, 1):
            if all(value is None for value in values):
                continue
            row = dict(zip(headers, values))
            code = row['Código INE']
            if not isinstance(code, str) or not re.fullmatch('[0-9]{5}', code) or code in seen:
                raise ValueError(f'Código municipal DGT no válido o repetido: registro {ordinal}')
            seen.add(code)
            if code[:2] not in DEMOGRAFIA_PROVINCES:
                continue
            measures = [row[field] for field in fields[1:]]
            if any(type(value) not in (int, float) or not math.isfinite(value) or not 0 <= value <= 2**63 - 1 or value != int(value) for value in measures):
                raise ValueError(f'Recuentos DGT no admitidos: {code}')
            measures = [int(value) for value in measures]
            if code.endswith('000'):
                unspecified.append([code, code[:2], date(2025, 12, 31), *measures])
            elif code in codes:
                municipalities.append([code, date(2025, 12, 31), *measures])
            else:
                raise ValueError(f'Municipio DGT no vigente en la referencia: {code}')
    finally:
        workbook.close()
    if {row[0] for row in municipalities} != set(codes):
        raise ValueError('Cobertura municipal DGT incompleta')
    if {row[0] for row in unspecified} != {province + '000' for province in {code[:2] for code in codes}}:
        raise ValueError('Agrupaciones provinciales sin municipio DGT inesperadas')
    connection.executemany('INSERT INTO parque_municipal VALUES (?,?,?,?,?)', municipalities)
    connection.executemany('INSERT INTO parque_sin_municipio VALUES (?,?,?,?,?,?)', unspecified)
    return {'filas_nacionales': len(seen), 'municipios': len(municipalities), 'sin_municipio': len(unspecified)}


def demografia_prepare_detail(connection, path, origin):
    # Tabla de trabajo en disco: solo atributos necesarios, sin cargar el TXT en Python.
    connection.execute("""
        CREATE TABLE dgt_entrada AS SELECT PROVINCIA,MUNICIPIO,TIPO_DGT,PROPULSION,CATELECT,TIPO_DISTINTIVO,ordinality
        FROM read_csv(?,header=true,delim='|',encoding='latin-1',all_varchar=true,
            parallel=false,strict_mode=true,ignore_errors=false) WITH ORDINALITY
    """, [str(path)])
    national = connection.execute('SELECT count(*),min(ordinality),max(ordinality) FROM dgt_entrada').fetchone()
    if national[1:] != (1, national[0]):
        raise ValueError('Ordinales nacionales DGT incoherentes')
    connection.execute("""
        CREATE TABLE parque_detalle_observado AS SELECT ?::VARCHAR origen_id,
            ordinality::BIGINT registro_origen,PROVINCIA codigo_provincia,
            nullif(trim(MUNICIPIO),'') codigo_municipio,'2025-12'::VARCHAR periodo,
            TIPO_DGT tipo_dgt_codigo,PROPULSION propulsion_codigo,
            CATELECT categoria_electrica_codigo,TIPO_DISTINTIVO distintivo_codigo
        FROM dgt_entrada WHERE PROVINCIA IN ('03','12','46')
    """, [origin])
    bad = consultar(connection, """
        SELECT registro_origen,codigo_provincia,codigo_municipio FROM parque_detalle_observado
        WHERE codigo_municipio IS NOT NULL AND (substr(codigo_municipio,1,2)<>codigo_provincia
           OR codigo_municipio NOT IN (SELECT codigo_municipio FROM municipios)) LIMIT 10
    """)
    if bad:
        raise ValueError(f'Municipios publicados por DGT no admitidos: {bad}')
    return {'filas_nacionales': national[0], 'filas_fuera_ambito': connection.execute("SELECT count(*) FROM dgt_entrada WHERE PROVINCIA NOT IN ('03','12','46') OR PROVINCIA IS NULL").fetchone()[0]}


def comprobar_clave(connection, table, columns):
    key = ','.join(identificador_sql(column) for column in columns)
    nulls = ' OR '.join(f'{identificador_sql(column)} IS NULL' for column in columns)
    if connection.execute(f'SELECT count(*) FROM {identificador_sql(table)} WHERE {nulls}').fetchone()[0]:
        raise ValueError(f'Clave incompleta: {table}')
    if connection.execute(f'SELECT count(*) FROM (SELECT {key} FROM {identificador_sql(table)} GROUP BY {key} HAVING count(*)>1)').fetchone()[0]:
        raise ValueError(f'Clave repetida: {table}')


def demografia_validate(connection, codes):
    for table, key in DEMOGRAFIA_KEYS.items():
        comprobar_clave(connection, table, key)
    park = consultar(connection, 'SELECT sum(parque_total) parque_total,sum(vehiculos_distintivo_0) distintivo_0,sum(electrificado_origen) electrificado FROM parque_municipal')[0]
    unspecified = consultar(connection, 'SELECT codigo_provincia,parque_total,vehiculos_distintivo_0,electrificado_origen FROM parque_sin_municipio ORDER BY 1')
    detail = consultar(connection, 'SELECT codigo_provincia,count(*) registros,count(*) FILTER(WHERE codigo_municipio IS NULL) sin_municipio,count(DISTINCT codigo_municipio) municipios_observados FROM parque_detalle_observado GROUP BY 1 ORDER BY 1')
    if sum(row['registros'] for row in detail) != park['parque_total'] + sum(row['parque_total'] for row in unspecified):
        raise ValueError('El detalle DGT no reconcilia con el total municipal y las agrupaciones PP000')
    disagreements = consultar(connection, """
        SELECT codigo_municipio,observados,parque_total FROM
        (SELECT codigo_municipio,count(*) observados FROM parque_detalle_observado
         WHERE codigo_municipio IS NOT NULL GROUP BY 1) d JOIN parque_municipal USING(codigo_municipio)
        WHERE observados<>parque_total
    """)
    if disagreements:
        raise ValueError(f'El detalle observado DGT no coincide con el parque municipal: {disagreements}')
    population = consultar(connection, """
        SELECT anio,sexo,count(*) filas,count(DISTINCT codigo_municipio) municipios,
            count(*) FILTER(WHERE estado_valor<>'informado') sin_valor
        FROM poblacion_municipal GROUP BY 1,2 ORDER BY 1,2
    """)
    if any(row['municipios'] != len(codes) for row in population):
        raise ValueError('Cobertura de población incompleta por año y sexo')
    sex_conflicts = consultar(connection, """
        SELECT codigo_municipio,anio FROM poblacion_municipal GROUP BY 1,2
        HAVING count(*)<>3 OR (count(poblacion)=3 AND
            max(poblacion) FILTER(WHERE sexo='Total')<>
            sum(poblacion) FILTER(WHERE sexo IN ('Hombres','Mujeres')))
    """)
    if sex_conflicts:
        raise ValueError(f'Población: Total no reconcilia con hombres y mujeres: {sex_conflicts[:10]}')
    income = consultar(connection, """
        SELECT nivel_territorial,anio,count(*) filas,count(DISTINCT codigo_municipio) municipios,
            count(*) FILTER(WHERE estado_valor='no_informado') sin_valor,
            count(*) FILTER(WHERE estado_valor='error_conversion') errores
        FROM renta_territorial GROUP BY 1,2 ORDER BY 1,2
    """)
    if any(row['municipios'] != len(codes) for row in income if row['nivel_territorial'] == 'municipio'):
        raise ValueError('Cobertura municipal de renta incompleta')
    return {'parque_municipal_totales': park, 'parque_sin_municipio': unspecified,
            'detalle_por_provincia': detail, 'municipios_detalle_reconciliados': True,
            'poblacion_por_anio_sexo': population, 'poblacion_total_reconciliada': True,
            'renta_por_nivel_anio': income,
            'categorias_electricas_detalle': consultar(connection, 'SELECT categoria_electrica_codigo,count(*) registros FROM parque_detalle_observado GROUP BY 1 ORDER BY 1'),
            'distintivos_detalle': consultar(connection, 'SELECT distintivo_codigo,count(*) registros FROM parque_detalle_observado GROUP BY 1 ORDER BY 1'),
            'incidencias_conversion': connection.execute('SELECT count(*) FROM incidencias').fetchone()[0]}


# %% Silver: instalaciones, puntos y conectores MITECO

RECARGA_SOURCE = 'Fuentes/instalaciones_actuales.csv'


RECARGA_HEADERS = ['COMUNIDAD AUTONOMA', 'PROVINCIA', 'MUNICIPIO', 'LATITUD', 'LONGITUD',
           'NOMBRE INSTALACION', 'DIRECCIÓN', 'CODIGO POSTAL', 'LOCALIZACION',
           'TIPO HORARIO APERTURA', 'HORARIO APERTURA', 'NOMBRE OPERADOR', 'COD.OPERADOR',
           'TIPOS DE SERVICIOS', 'COD.INSTALACION', 'ID. PUNTO DE RECARGA', 'ACCESIBILIDAD',
           'METODOS DE PAGOS', 'COD. PUNTO DE RECARGA', 'ID. CONECTOR', 'TIPO CONECTOR',
           'TIPO DE CARGA', 'FORMATO', 'POTENCIA MAXIMA', 'VOLTAJE', 'INTENSIDAD',
           'FECHA DE ULTIMA MODIFICACION']


RECARGA_INSTALL = {'instalacion_id': 14, 'nombre_instalacion': 5, 'operador_id': 12,
           'nombre_operador': 11, 'provincia_nombre_origen': 1, 'municipio_nombre_origen': 2,
           'direccion': 6, 'codigo_postal': 7, 'latitud_origen': 3, 'longitud_origen': 4,
           'localizacion_origen': 8, 'tipo_horario_origen': 9}


RECARGA_POINT = {'instalacion_id': 14, 'punto_id': 15, 'codigo_punto_origen': 18, 'accesibilidad_origen': 16}


RECARGA_CONNECTOR = {'instalacion_id': 14, 'punto_id': 15, 'conector_id': 19, 'tipo_conector': 20,
             'tipo_carga': 21, 'formato_conector': 22, 'potencia_maxima_origen': 23,
             'fecha_ultima_modificacion_origen': 26, 'horario_apertura_origen': 10,
             'servicios_origen': 13, 'metodos_pago_origen': 17,
             'voltaje_origen': 24, 'intensidad_origen': 25}


RECARGA_KEYS = {'recarga_instalaciones': ['instalacion_id'],
        'recarga_puntos': ['instalacion_id', 'punto_id'],
        'recarga_conectores': ['instalacion_id', 'punto_id', 'conector_id']}


RECARGA_PROVINCES = {'Alicante/Alacant', 'Castellón/Castelló', 'Valencia/València'}


RECARGA_UNKNOWN = {'', 'DESCONOCIDO', 'NODISPONIBLE'}


RECARGA_FIELD = re.compile(r'(=?)"(.*?)"(?=;|\r?\n|\Z)', re.S)


def recarga_export_rows(path):
    # shortcut: dialecto de esta exportación de hasta 64 MiB; revisar si cambia el exportador.
    with path.open('rb') as stream:
        raw = stream.read(64 * 1024**2 + 1)
    if len(raw) > 64 * 1024**2:
        raise ValueError('Exportación mayor que el límite de lectura de 64 MiB')
    text = raw.decode('utf-16-le')
    header, separator, _ = text.partition('\n')
    if not separator or header.rstrip('\r').split(';') != RECARGA_HEADERS:
        raise ValueError('Cabecera MITECO inesperada')
    pos, ordinal = len(header) + 1, 0
    while pos < len(text):
        start, values, wrappers = pos, [], []
        for column in range(len(RECARGA_HEADERS)):
            match = RECARGA_FIELD.match(text, pos)
            if not match or (match[1] and column not in (3, 4, 7, 8)):
                raise ValueError(f'Dialecto MITECO no admitido: registro {ordinal + 1}, columna {column + 1}')
            wrappers.append(match[1])
            values.append(match[2])
            pos = match.end()
            if column < len(RECARGA_HEADERS) - 1:
                if text[pos:pos + 1] != ';':
                    raise ValueError(f'Fila incompleta MITECO: registro {ordinal + 1}')
                pos += 1
            elif pos != len(text):
                if text[pos:pos + 2] == '\r\n':
                    pos += 2
                elif text[pos:pos + 1] == '\n':
                    pos += 1
                else:
                    raise ValueError(f'Fila demasiado ancha MITECO: registro {ordinal + 1}')
        ordinal += 1
        # Solo se retiran envoltorios: los caracteres interiores, incluidos CR/LF, son literales.
        reconstructed = ';'.join(w + '"' + v + '"' for w, v in zip(wrappers, values))
        if reconstructed != text[start:match.end()]:
            raise ValueError(f'Lectura no reversible: registro {ordinal}')
        standard = next(csv.reader(io.StringIO(text[start:pos]), delimiter=';'))
        standard = [v[2:-1] if v.startswith('="') and v.endswith('"') else v for v in standard]
        yield ordinal, values, values != standard


def recarga_coordinate(raw, latitude):
    if raw.strip() in RECARGA_UNKNOWN:
        return None, 'no_informado'
    if not re.fullmatch(r'-?[0-9]+(?:\.[0-9]+)?', raw):
        return None, 'error_conversion'
    value = float(raw)
    limit = 90 if latitude else 180
    return (value, 'rango_valido_datum_pendiente') if -limit <= value <= limit else (None, 'fuera_rango')


def recarga_power(raw):
    if raw.strip() in RECARGA_UNKNOWN:
        return None, 'no_informado'
    if not re.fullmatch(r'[0-9]+(?:,[0-9]{1,2})? kW', raw):
        return None, 'error_conversion'
    value = Decimal(raw[:-3].replace(',', '.'))
    if value >= Decimal('10000000000000000'):
        return None, 'error_conversion'
    return value, 'informado'


def recarga_modified(raw):
    if raw.strip() in RECARGA_UNKNOWN:
        return None, 'no_informado'
    if not re.fullmatch(r'[0-9]{2}/[0-9]{2}/[0-9]{4} [0-9]{1,2}:[0-9]{2}:[0-9]{2}', raw):
        return None, 'error_conversion'
    try:
        value = datetime.strptime(raw, '%d/%m/%Y %H:%M:%S')
    except ValueError:
        return None, 'error_conversion'
    return value, 'informado_sin_zona'


def recarga_parent(parents, key, values, entity):
    if key in parents and parents[key] != values:
        raise ValueError(f'Atributos contradictorios en {entity}: {key}')
    parents[key] = values


def recarga_prepare(connection, path, temporary):
    installs, points, national = {}, {}, defaultdict(list)
    details, incidents, counts = [], [], Counter()
    for ordinal, row, differs in recarga_export_rows(path):
        counts['filas_nacionales'] += 1
        counts['textos_preservados_frente_csv_permisivo_nacional'] += differs
        key = tuple(row[i] for i in (14, 15, 19))
        if any(not value.strip() for value in key):
            raise ValueError(f'Clave de conector incompleta: registro {ordinal}')
        national[key].append((ordinal, row))
        if (row[0] == 'Comunitat Valenciana') != (row[1] in RECARGA_PROVINCES):
            raise ValueError(f'Comunidad y provincia incoherentes: registro {ordinal}')
        if row[0] != 'Comunitat Valenciana':
            counts['filas_fuera_ambito'] += 1
            continue
        counts['filas_cv'] += 1
        counts['textos_preservados_frente_csv_permisivo_cv'] += differs
        recarga_parent(installs, key[0], tuple(row[i] for i in RECARGA_INSTALL.values()), 'instalación')
        recarga_parent(points, key[:2], tuple(row[i] for i in RECARGA_POINT.values()), 'punto')
        kw, kw_state = recarga_power(row[23])
        timestamp, date_state = recarga_modified(row[26])
        details.append([*(row[i] for i in RECARGA_CONNECTOR.values()), ordinal, kw, kw_state, timestamp, date_state])
        for field, raw, state in [('potencia_maxima_kw', row[23], kw_state),
                                  ('fecha_ultima_modificacion', row[26], date_state)]:
            if state == 'error_conversion':
                incidents.append(['recarga_conectores', ordinal, field, raw, state])
    if not details:
        raise ValueError('No hay conectores de la Comunitat Valenciana')
    for name, fields in [('recarga_instalaciones', RECARGA_INSTALL), ('recarga_puntos', RECARGA_POINT), ('recarga_conectores', RECARGA_CONNECTOR)]:
        columns = ','.join(identificador_sql(field) + ' VARCHAR' for field in fields)
        if name == 'recarga_instalaciones':
            columns += ',latitud DOUBLE,longitud DOUBLE,estado_coordenadas VARCHAR'
        elif name == 'recarga_conectores':
            columns += ',registro_origen BIGINT,potencia_maxima_kw DECIMAL(18,2),estado_potencia VARCHAR,fecha_ultima_modificacion TIMESTAMP,estado_fecha VARCHAR'
        connection.execute(f'CREATE TABLE {name} ({columns})')
    installation_rows = []
    for key, values in installs.items():
        lat, lat_state = recarga_coordinate(values[8], True)
        lon, lon_state = recarga_coordinate(values[9], False)
        state = 'rango_valido_datum_pendiente'
        if lat is None or lon is None:
            state = 'coordenada_no_admitida'
        elif lat == lon == 0:
            state = 'coordenadas_0_0'
        installation_rows.append([*values, lat, lon, state])
        if state != 'rango_valido_datum_pendiente':
            ordinal = next(r[-5] for r in details if r[0] == key)
            incidents.append(['recarga_instalaciones', ordinal, 'coordenadas',
                              json.dumps(values[8:10]), state + ':' + lat_state + ':' + lon_state])
    connection.execute('CREATE TABLE incidencias (tabla VARCHAR,registro_origen BIGINT,campo VARCHAR,valor_origen VARCHAR,motivo VARCHAR)')
    for table, records in [('recarga_instalaciones', installation_rows),
                           ('recarga_puntos', points.values()), ('recarga_conectores', details),
                           ('incidencias', incidents)]:
        schema = {r[0]: r[1] for r in connection.execute(f'DESCRIBE {table}').fetchall()}
        staging = temporary / f'{table}.csv'
        with staging.open('w', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(schema)
            writer.writerows(records)
        # Una carga por tabla evita miles de inserciones; restaurar los textos vacíos, no los nulos numéricos.
        columns = ','.join(f"coalesce({identificador_sql(k)},'')" if v == 'VARCHAR' else identificador_sql(k) for k, v in schema.items())
        connection.execute(f'''INSERT INTO {table} SELECT {columns} FROM read_csv(?,columns=?,
            header=true,delim=',',quote='"',escape='"',auto_detect=false,
            timestampformat='%Y-%m-%d %H:%M:%S',strict_mode=true,ignore_errors=false)''', [str(staging), schema])
    duplicates = []
    for key, rows in national.items():
        if len(rows) > 1:
            duplicates.append({'clave': list(key), 'registros_origen': [n for n, _ in rows],
                'comunidades': sorted({r[0] for _, r in rows}),
                'campos_variables': [RECARGA_HEADERS[i] for i in range(len(RECARGA_HEADERS)) if len({r[i] for _, r in rows}) > 1],
                'filas_identicas': len({tuple(r) for _, r in rows}) == 1})
    return {'recuentos': dict(counts), 'claves_repetidas_nacionales': duplicates,
            'filas_adicionales_por_clave_nacional': sum(len(d['registros_origen']) - 1 for d in duplicates),
            'filas_eliminadas_en_cv': 0, 'atributos_padre_contrastados': True,
            'dialecto': 'UTF-16LE; campos entre comillas exteriores, comillas interiores literales, envoltorios ="texto" sin evaluación; 27 campos; reconstrucción exacta de cada registro'}


def recarga_validate(connection):
    for table, keys in RECARGA_KEYS.items():
        comprobar_clave(connection, table, keys)
        for key in keys:
            if connection.execute(f'SELECT count(*) FROM {table} WHERE trim({identificador_sql(key)})=\'\'').fetchone()[0]:
                raise ValueError(f'Clave vacía: {table}')
    comprobar_clave(connection, 'recarga_conectores', ['registro_origen'])
    for child, target, keys in [('recarga_conectores', 'recarga_puntos', ['instalacion_id', 'punto_id']),
                                 ('recarga_puntos', 'recarga_instalaciones', ['instalacion_id'])]:
        columns = ','.join(keys)
        if connection.execute(f'SELECT count(*) FROM {child} c ANTI JOIN {target} p USING({columns})').fetchone()[0]:
            raise ValueError(f'Relación incompleta: {child}')
        if connection.execute(f'SELECT count(*) FROM {target} p ANTI JOIN {child} c USING({columns})').fetchone()[0]:
            raise ValueError(f'Entidad sin detalle: {target}')
    variation = {}
    for name, keys, field in [('horarios_por_instalacion', 'instalacion_id', 'horario_apertura_origen'),
                              ('servicios_por_instalacion', 'instalacion_id', 'servicios_origen'),
                              ('pagos_por_punto', 'instalacion_id,punto_id', 'metodos_pago_origen')]:
        variation[name] = connection.execute(f'SELECT count(*) FROM (SELECT {keys} FROM recarga_conectores GROUP BY {keys} HAVING count(DISTINCT {field})>1)').fetchone()[0]
    return {'claves_y_relaciones_validas': True, 'variantes_textuales_conservadas': variation,
        'por_provincia': consultar(connection, '''SELECT provincia_nombre_origen,count(DISTINCT i.instalacion_id) instalaciones,
            count(DISTINCT (p.instalacion_id,p.punto_id)) puntos,count(*) conectores
            FROM recarga_instalaciones i JOIN recarga_puntos p USING(instalacion_id)
            JOIN recarga_conectores c USING(instalacion_id,punto_id) GROUP BY 1 ORDER BY 1'''),
        'coordenadas': consultar(connection, 'SELECT estado_coordenadas,count(*) instalaciones FROM recarga_instalaciones GROUP BY 1'),
        'potencia': consultar(connection, 'SELECT estado_potencia,count(*) conectores,min(potencia_maxima_kw)::VARCHAR minimo_kw,max(potencia_maxima_kw)::VARCHAR maximo_kw FROM recarga_conectores GROUP BY 1'),
        'fechas': consultar(connection, 'SELECT estado_fecha,count(*) conectores,min(fecha_ultima_modificacion)::VARCHAR minima,max(fecha_ultima_modificacion)::VARCHAR maxima FROM recarga_conectores GROUP BY 1'),
        'accesibilidad': consultar(connection, 'SELECT accesibilidad_origen,count(*) puntos FROM recarga_puntos GROUP BY 1 ORDER BY 1'),
        'incidencias': connection.execute('SELECT count(*) FROM incidencias').fetchone()[0]}


# %% Lectura de geometrías GeoPackage

def gpkg_gpkg_wkb(blob, expected_srs):
    """Extrae el WKB estándar sin perder el fid que se lee con SQLite."""
    if not isinstance(blob, bytes) or len(blob) < 8 or blob[:3] != b'GP\x00':
        raise ValueError('Cabecera GeoPackage no admitida')
    flags = blob[3]
    envelope = (flags >> 1) & 7
    if flags & 0xE0 or envelope > 4:
        raise ValueError('Geometría extendida o indicadores GeoPackage no admitidos')
    order = '<' if flags & 1 else '>'
    if struct.unpack(order + 'i', blob[4:8])[0] != expected_srs:
        raise ValueError('SRS de geometría distinto del declarado en su capa')
    offset = 8 + (0, 32, 48, 48, 64)[envelope]
    if len(blob) < offset + 5 or blob[offset] not in (0, 1):
        raise ValueError('Contenido WKB ausente o truncado')
    return blob[offset:]


# %% Silver: aforos y red viaria

TRAFICO_GVA = 'Fuentes/Trafico.gpkg'


TRAFICO_ROADS = {'03': 'Fuentes/RT_ALICANTE_ALACANT_gpkg',
         '12': 'Fuentes/RT_CASTELLON_CASTELLO_gpkg', '46': 'Fuentes/RT_VALENCIA_gpkg'}


TRAFICO_TRAFFIC_FIELDS = ['id', 'clau_tramo', 'carretera', 'titular', 'tipo_calzada', 'tipo_medida_trafico',
                  'pk_inicio', 'pk_estacion', 'pk_fin', 'imd', 'vehiculo_pesados_zp', 'label2_planos']


TRAFICO_ROAD_FIELDS = ['id_tramo', 'id_vial', 'tipo_tramo', 'tipo_tramd', 'codigo', 'nombre', 'tipo_vial', 'tipo_viald']


TRAFICO_MEASURES = {'1a': 'Primaria', '2a': 'Secundaria', 'C': 'Cobertura', 'Est': 'Estimada',
            'P': 'Permanente', 'R': 'Reforzada', 'Eix': 'Salida', 'Dip': 'Diputación'}


TRAFICO_KEYS = {'aforos_tramo_anio': ['tramo_aforo_id', 'anio'],
        'red_tramos': ['provincia_origen', 'tramo_id'],
        'red_tramo_vial': ['provincia_origen', 'tramo_id', 'vial_id']}


TRAFICO_NULL = '\x1eNULL_SILVER\x1e'


def trafico_read_layer(connection, path, layer, srs, fields, temporary):
    """SQLite conserva fid y WKB de la misma fila; CSV temporal evita inserciones individuales."""
    staging = temporary / 'entrada.csv'
    schema = {'fid_origen': 'BIGINT', **dict.fromkeys(fields, 'VARCHAR'), 'wkb_hex': 'VARCHAR'}
    with closing(sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)) as source:
        geo = source.execute('SELECT column_name,srs_id,m FROM gpkg_geometry_columns WHERE table_name=?', [layer]).fetchall()
        if len(geo) != 1 or geo[0][1:] != (srs, 0):
            raise ValueError(f'Geometría/CRS/M no admitidos: {layer}')
        reference = source.execute('SELECT organization,organization_coordsys_id FROM gpkg_spatial_ref_sys WHERE srs_id=?', [srs]).fetchone()
        if reference != ('EPSG', srs):
            raise ValueError(f'Autoridad CRS no acreditada: {layer}')
        columns = {r[1] for r in source.execute(f'PRAGMA table_info({identificador_sql(layer)})')}
        if not {'fid', geo[0][0], *fields}.issubset(columns):
            raise ValueError(f'Campos de entrada ausentes: {layer}')
        with staging.open('w', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(schema)
            rows = source.execute(f'SELECT fid,{",".join(identificador_sql(f) for f in fields)},{identificador_sql(geo[0][0])} FROM {identificador_sql(layer)} ORDER BY fid')
            count = 0
            for fid, *values in rows:
                blob = values.pop()
                if type(fid) is not int or any(v is not None and str(v) == TRAFICO_NULL for v in values):
                    raise ValueError(f'Identificador o marcador reservado no admitido: {layer}/{fid}')
                wkb = gpkg_gpkg_wkb(blob, srs).hex() if blob is not None else None
                writer.writerow([fid, *(TRAFICO_NULL if v is None else str(v) for v in values), TRAFICO_NULL if wkb is None else wkb])
                count += 1
    connection.execute('DROP TABLE IF EXISTS entrada')
    connection.execute('''CREATE TEMP TABLE entrada AS SELECT * FROM read_csv(?,columns=?,nullstr=?,
        header=true,delim=',',quote='"',escape='"',auto_detect=false,strict_mode=true,ignore_errors=false)''',
        [str(staging), schema, TRAFICO_NULL])
    if connection.execute('SELECT count(*) FROM entrada').fetchone()[0] != count:
        raise ValueError(f'Pérdida al cargar: {layer}')
    comprobar_clave(connection, 'entrada', ['fid_origen'])
    staging.unlink()
    return count


def trafico_numeric_sql(field):
    name = identificador_sql(field)
    return f"CASE WHEN regexp_full_match(trim({name}),'[0-9]+([.][0-9]{{1,2}})?') THEN try_cast({name} AS DECIMAL(18,2)) END"


def trafico_prepare_traffic(connection, path, temporary, origin, years=range(2009, 2026)):
    summaries = []
    cases = ' '.join(f'WHEN {texto_sql(k)} THEN {texto_sql(v)}' for k, v in TRAFICO_MEASURES.items())
    for year in years:
        layer = f'Aforos{year}'
        count = trafico_read_layer(connection, path, layer, 25830, TRAFICO_TRAFFIC_FIELDS, temporary)
        if not count:
            raise ValueError(f'Capa de aforos vacía: {layer}')
        connection.execute(f'''CREATE OR REPLACE TEMP TABLE aforo_preparado AS
            SELECT clau_tramo tramo_aforo_id,?::SMALLINT anio,carretera,titular titular_origen,
                tipo_calzada tipo_calzada_origen,tipo_medida_trafico,
                CASE tipo_medida_trafico {cases} END tipo_medida_descripcion,
                pk_inicio,pk_estacion,pk_fin,imd imd_origen,vehiculo_pesados_zp pesados_origen,
                {trafico_numeric_sql('imd')} imd_vehiculos_dia,{trafico_numeric_sql('vehiculo_pesados_zp')} porcentaje_pesados,
                label2_planos etiqueta_plano_origen,id identificador_origen,
                ?::VARCHAR origen_id,?::VARCHAR capa_origen,fid_origen,
                ST_GeomFromWKB(from_hex(wkb_hex))::GEOMETRY('EPSG:25830') geometria
            FROM entrada''', [year, origin, layer])
        connection.execute('''CREATE OR REPLACE TEMP TABLE aforo_final AS SELECT *,
            CASE WHEN imd_origen IS NULL OR trim(imd_origen)='' THEN 'no_informado'
                 WHEN imd_vehiculos_dia IS NULL THEN 'error_conversion'
                 WHEN imd_vehiculos_dia=0 THEN 'cero_publicado_revisar_uso' ELSE 'informado' END estado_imd,
            CASE WHEN pesados_origen IS NULL OR trim(pesados_origen)='' THEN 'no_informado'
                 WHEN porcentaje_pesados IS NULL THEN 'error_conversion'
                 WHEN porcentaje_pesados NOT BETWEEN 0 AND 100 THEN 'fuera_rango' ELSE 'informado' END estado_pesados,
            CASE WHEN tipo_medida_descripcion IS NOT NULL THEN 'documentado'
                 WHEN tipo_medida_trafico IS NULL OR trim(tipo_medida_trafico)='' THEN 'no_informado'
                 ELSE 'codigo_no_documentado' END estado_tipo_medida
            FROM aforo_preparado''')
        if connection.execute('SELECT count(*) FROM aforo_final WHERE estado_pesados=\'fuera_rango\'').fetchone()[0]:
            connection.execute("UPDATE aforo_final SET porcentaje_pesados=NULL WHERE estado_pesados='fuera_rango'")
        if not summaries:
            connection.execute('CREATE TABLE aforos_tramo_anio AS SELECT * FROM aforo_final WHERE false')
        connection.execute('INSERT INTO aforos_tramo_anio SELECT * FROM aforo_final')
        summaries.append({'anio': year, 'filas_bronze': count})
    return summaries


def trafico_prepare_road(connection, path, province, temporary, origin):
    count = trafico_read_layer(connection, path, 'rt_tramo_vial', 4258, TRAFICO_ROAD_FIELDS, temporary)
    connection.execute('''CREATE OR REPLACE TEMP TABLE tramos_candidatos AS
        SELECT id_tramo tramo_id,tipo_tramo tipo_tramo_codigo,tipo_tramd tipo_tramo_descripcion,
            wkb_hex,min(fid_origen)::BIGINT fid_referencia
        FROM entrada GROUP BY id_tramo,tipo_tramo,tipo_tramd,wkb_hex''')
    comprobar_clave(connection, 'tramos_candidatos', ['tramo_id'])
    for field in ('id_tramo', 'id_vial'):
        if connection.execute(f'SELECT count(*) FROM entrada WHERE {identificador_sql(field)} IS NULL OR trim({identificador_sql(field)})=\'\'').fetchone()[0]:
            raise ValueError(f'Clave viaria incompleta: {province}/{field}')
    connection.execute('''CREATE OR REPLACE TEMP TABLE tramos_preparados AS SELECT
        ?::VARCHAR provincia_origen,tramo_id,tipo_tramo_codigo,tipo_tramo_descripcion,
        ?::VARCHAR origen_id,'rt_tramo_vial'::VARCHAR capa_origen,fid_referencia,
        ST_GeomFromWKB(from_hex(wkb_hex)) geometria_origen
        FROM tramos_candidatos''', [province, origin])
    source_invalid = connection.execute('''SELECT count(*) FROM tramos_preparados
        WHERE geometria_origen IS NULL OR ST_IsEmpty(geometria_origen) OR NOT ST_IsValid(geometria_origen)
            OR ST_GeometryType(geometria_origen) NOT IN ('LINESTRING','MULTILINESTRING')''').fetchone()[0]
    if source_invalid:
        raise ValueError(f'Geometrías IGN de origen no admitidas: {province}, {source_invalid}')
    connection.execute('''CREATE OR REPLACE TEMP TABLE tramos_final AS SELECT * EXCLUDE(geometria_origen),
        ST_Transform(geometria_origen,'EPSG:4258','EPSG:25830',always_xy:=true)
            ::GEOMETRY('EPSG:25830') geometria
        FROM tramos_preparados''')
    altered = connection.execute('''SELECT count(*) FROM tramos_preparados a JOIN tramos_final b USING(tramo_id)
        WHERE ST_HasZ(a.geometria_origen) IS DISTINCT FROM ST_HasZ(b.geometria)
        OR ST_NPoints(a.geometria_origen) IS DISTINCT FROM ST_NPoints(b.geometria)
        OR ST_ZMin(a.geometria_origen) IS DISTINCT FROM ST_ZMin(b.geometria)
        OR ST_ZMax(a.geometria_origen) IS DISTINCT FROM ST_ZMax(b.geometria)''').fetchone()[0]
    if altered:
        raise ValueError(f'Cambios de dimensión, vértices o extremos Z: {province}')
    connection.execute('''CREATE OR REPLACE TEMP TABLE relaciones_final AS SELECT
        ?::VARCHAR provincia_origen,id_tramo tramo_id,id_vial vial_id,codigo codigo_vial,
        nombre nombre_vial,tipo_vial tipo_vial_codigo,tipo_viald tipo_vial_descripcion,
        ?::VARCHAR origen_id,'rt_tramo_vial'::VARCHAR capa_origen,fid_origen
        FROM entrada''', [province, origin])
    if not connection.execute("SELECT count(*) FROM information_schema.tables WHERE table_name='red_tramos'").fetchone()[0]:
        connection.execute('CREATE TABLE red_tramos AS SELECT * FROM tramos_final WHERE false')
        connection.execute('CREATE TABLE red_tramo_vial AS SELECT * FROM relaciones_final WHERE false')
        connection.execute('CREATE TABLE control_tramos_origen (provincia_origen VARCHAR,tramo_id VARCHAR,wkb_hex VARCHAR)')
    connection.execute('INSERT INTO red_tramos SELECT * FROM tramos_final')
    connection.execute('INSERT INTO red_tramo_vial SELECT * FROM relaciones_final')
    connection.execute('INSERT INTO control_tramos_origen SELECT ?,tramo_id,wkb_hex FROM tramos_candidatos', [province])
    return {'provincia_origen': province, 'asociaciones_bronze': count,
            'tramos': connection.execute('SELECT count(*) FROM tramos_candidatos').fetchone()[0],
            'geometrias_origen_validas': True, 'dimension_vertices_extremos_z_conservados': True}


def trafico_validate(connection):
    for table, keys in TRAFICO_KEYS.items():
        comprobar_clave(connection, table, keys)
        for key in keys:
            if key != 'anio' and connection.execute(f"SELECT count(*) FROM {table} WHERE trim({identificador_sql(key)})=''").fetchone()[0]:
                raise ValueError(f'Clave vacía: {table}/{key}')
    for child, parent in [('red_tramo_vial', 'red_tramos'), ('red_tramos', 'red_tramo_vial')]:
        if connection.execute(f'SELECT count(*) FROM {child} ANTI JOIN {parent} USING(provincia_origen,tramo_id)').fetchone()[0]:
            raise ValueError('Relación tramo-vial incompleta')
    connection.execute('''CREATE TABLE incidencias (
        tabla VARCHAR,provincia_origen VARCHAR,capa_origen VARCHAR,fid_origen BIGINT,
        entidad_id VARCHAR,campo VARCHAR,valor_origen VARCHAR,motivo VARCHAR)''')
    for field, raw, state, accepted in [('imd_vehiculos_dia', 'imd_origen', 'estado_imd', 'informado'),
                                        ('porcentaje_pesados', 'pesados_origen', 'estado_pesados', 'informado'),
                                        ('tipo_medida_trafico', 'tipo_medida_trafico', 'estado_tipo_medida', 'documentado')]:
        connection.execute(f'''
            INSERT INTO incidencias SELECT 'aforos_tramo_anio',NULL,capa_origen,fid_origen,
            tramo_aforo_id,?,{identificador_sql(raw)},{identificador_sql(state)} FROM aforos_tramo_anio WHERE {identificador_sql(state)}<>?''', [field, accepted])
    geometry = []
    for table, entity, fid, province in [('aforos_tramo_anio', 'tramo_aforo_id', 'fid_origen', 'NULL::VARCHAR'),
                                         ('red_tramos', 'tramo_id', 'fid_referencia', 'provincia_origen')]:
        connection.execute(f'''ALTER TABLE {table} ADD COLUMN estado_geometria VARCHAR''')
        connection.execute(f'''UPDATE {table} SET estado_geometria=CASE
            WHEN geometria IS NULL THEN 'nula' WHEN ST_IsEmpty(geometria) THEN 'vacia'
            WHEN NOT ST_IsValid(geometria) THEN 'invalida'
            WHEN ST_GeometryType(geometria) NOT IN ('LINESTRING','MULTILINESTRING') THEN 'tipo_no_lineal'
            WHEN NOT isfinite(ST_Length(geometria)) OR ST_Length(geometria)<=0 THEN 'longitud_no_positiva'
            ELSE 'admitida' END''')
        connection.execute(f'''INSERT INTO incidencias SELECT ?,{province},capa_origen,{fid},
            {entity},'geometria',NULL,estado_geometria FROM {table} WHERE estado_geometria<>'admitida' ''', [table])
        geometry.extend(consultar(connection, f'''SELECT {texto_sql(table)} tabla,estado_geometria,count(*) filas,
            sum(ST_Length(geometria)) longitud_plana_m,count(*) FILTER(WHERE ST_HasZ(geometria)) con_z
            FROM {table} GROUP BY 1,2 ORDER BY 2'''))
    temporal = consultar(connection, '''SELECT anio,count(*) filas,count(DISTINCT carretera) carreteras,
        count(*) FILTER(WHERE tipo_medida_trafico='Est') estimados,
        count(*) FILTER(WHERE estado_imd='cero_publicado_revisar_uso') ceros,
        min(imd_vehiculos_dia)::VARCHAR minimo_imd,max(imd_vehiculos_dia)::VARCHAR maximo_imd
        FROM aforos_tramo_anio GROUP BY 1 ORDER BY 1''')
    changes = consultar(connection, '''SELECT a.anio anio_anterior,b.anio anio_siguiente,count(*) claves_comunes,
        count(*) FILTER(WHERE a.carretera IS DISTINCT FROM b.carretera OR a.pk_inicio IS DISTINCT FROM b.pk_inicio
            OR a.pk_fin IS DISTINCT FROM b.pk_fin) atributos_trazado_distintos,
        count(*) FILTER(WHERE ST_AsWKB(a.geometria) IS DISTINCT FROM ST_AsWKB(b.geometria)) wkb_distinto
        FROM aforos_tramo_anio a JOIN aforos_tramo_anio b
        ON a.tramo_aforo_id=b.tramo_aforo_id AND b.anio=a.anio+1 GROUP BY 1,2 ORDER BY 1''')
    cross = consultar(connection, '''SELECT count(*) ids_compartidos,
        count(*) FILTER(WHERE representaciones>1) con_wkb_distinto FROM
        (SELECT tramo_id,count(DISTINCT wkb_hex) representaciones FROM control_tramos_origen
        GROUP BY tramo_id HAVING count(DISTINCT provincia_origen)>1)''')[0]
    return {'claves_y_relaciones_validas': True, 'aforos_por_anio': temporal,
            'comparacion_entre_anios': changes, 'geometrias': geometry,
            'tipos_medida': consultar(connection, 'SELECT tipo_medida_trafico,tipo_medida_descripcion,count(*) filas FROM aforos_tramo_anio GROUP BY 1,2 ORDER BY 1'),
            'tramos_compartidos_entre_paquetes': cross,
            'incidencias_por_motivo': consultar(connection, 'SELECT tabla,campo,motivo,count(*) filas FROM incidencias GROUP BY 1,2,3 ORDER BY 1,2,3')}


# %% Silver: elementos OSM y correspondencias

OSM_SOURCE = 'Fuentes/ubicaciones_candidatasgpkg/valencia.gpkg'


OSM_LAYERS = [f'gis_osm_{family}{suffix}_free' for family in ('pois', 'traffic', 'transport') for suffix in ('', '_a')]


OSM_CRS84_WKT = 'GEOGCS["WGS 84 (CRS84)",DATUM["WGS_1984",SPHEROID["WGS 84",6378137,298.257223563]],PRIMEM["Greenwich",0],UNIT["degree",0.0174532925199433,AUTHORITY["EPSG","9122"]],AXIS["Longitude",EAST],AXIS["Latitude",NORTH],AUTHORITY["OGC","CRS84"]]'


OSM_CATEGORIES = {'parking': ('5260', 'aparcamiento'), 'fuel': ('5250', 'parada_viaria'),
    'service': ('5251', 'parada_viaria'), 'supermarket': ('2501', 'comercio'),
    'mall': ('2504', 'comercio'), 'department_store': ('2505', 'comercio'),
    'hotel': ('2401', 'alojamiento'), 'motel': ('2402', 'alojamiento'),
    'guesthouse': ('2404', 'alojamiento'), 'hostel': ('2405', 'alojamiento'),
    'camp_site': ('2422', 'alojamiento'), 'caravan_site': ('2424', 'alojamiento')}


OSM_KEYS = {'poi_elementos': ['origen_id', 'fid_origen'],
        'poi_correspondencias': ['origen_punto', 'fid_punto', 'origen_area', 'fid_area']}


OSM_NULL = '\x1eNULL_SILVER\x1e'


def osm_read_osm(connection, path, layer, temporary):
    schema = {'fid_origen': 'BIGINT', 'osm_id': 'VARCHAR', 'code': 'VARCHAR',
              'fclass': 'VARCHAR', 'name': 'VARCHAR', 'wkb_hex': 'VARCHAR'}
    staging = temporary / 'entrada_osm.csv'
    with closing(sqlite3.connect(path.resolve().as_uri() + '?mode=ro', uri=True)) as source:
        geo = source.execute('SELECT column_name,srs_id,z,m FROM gpkg_geometry_columns WHERE table_name=?', [layer]).fetchall()
        if geo != [('geom', 100000, 0, 0)]:
            raise ValueError(f'Metadatos geométricos OSM no admitidos: {layer}')
        reference = source.execute('SELECT organization,organization_coordsys_id,definition FROM gpkg_spatial_ref_sys WHERE srs_id=100000').fetchone()
        # El identificador 100000 es local; se exige el WKT CRS84 comprobado, no un EPSG inventado.
        if reference != ('NONE', 100000, OSM_CRS84_WKT):
            raise ValueError('Referencia OSM distinta de la definición CRS84 comprobada')
        with staging.open('w', encoding='utf-8', newline='') as stream:
            writer = csv.writer(stream)
            writer.writerow(schema)
            count = 0
            for fid, osm, code, category, name, blob in source.execute(f'SELECT fid,osm_id,code,fclass,name,geom FROM {identificador_sql(layer)} ORDER BY fid'):
                if type(fid) is not int:
                    raise ValueError(f'fid no entero: {layer}')
                values = [osm, code, category, name, gpkg_gpkg_wkb(blob, 100000).hex() if blob is not None else None]
                if any(v is not None and str(v) == OSM_NULL for v in values):
                    raise ValueError('Marcador de nulo reservado en el original')
                writer.writerow([fid, *(OSM_NULL if v is None else str(v) for v in values)])
                count += 1
    connection.execute('DROP TABLE IF EXISTS entrada_osm')
    connection.execute('''CREATE TEMP TABLE entrada_osm AS SELECT * FROM read_csv(?,columns=?,nullstr=?,
        header=true,delim=',',quote='"',escape='"',auto_detect=false,strict_mode=true,ignore_errors=false)''', [str(staging), schema, OSM_NULL])
    comprobar_clave(connection, 'entrada_osm', ['fid_origen'])
    if not count or connection.execute('SELECT count(*) FROM entrada_osm').fetchone()[0] != count:
        raise ValueError(f'Capa vacía o lectura no reconciliada: {layer}')
    staging.unlink()
    return count


def osm_prepare(connection, path, temporary, origin, layers=OSM_LAYERS):
    connection.execute('CREATE TEMP TABLE categorias (fclass VARCHAR PRIMARY KEY,code VARCHAR,grupo VARCHAR)')
    connection.executemany('INSERT INTO categorias VALUES (?,?,?)', [(k, *v) for k, v in OSM_CATEGORIES.items()])
    counts = []
    for layer in layers:
        count = osm_read_osm(connection, path, layer, temporary)
        area = '_a_free' in layer
        allowed = "('POLYGON','MULTIPOLYGON')" if area else "('POINT')"
        connection.execute(f'''CREATE OR REPLACE TEMP TABLE osm_geometria AS SELECT *,
            CASE WHEN g IS NULL THEN 'origen_nulo' WHEN ST_IsEmpty(g) THEN 'origen_vacio'
                 WHEN NOT ST_IsValid(g) THEN 'origen_invalido'
                 WHEN ST_GeometryType(g) NOT IN {allowed} THEN 'tipo_origen_no_admitido'
                 ELSE 'admitida' END estado_origen
            FROM (SELECT *,ST_GeomFromWKB(from_hex(wkb_hex)) g FROM entrada_osm)''')
        connection.execute('''CREATE OR REPLACE TEMP TABLE osm_preparado AS SELECT
            ?::VARCHAR origen_id,?::VARCHAR capa_origen,fid_origen,osm_id,
            o.code categoria_codigo_origen,o.fclass categoria_origen,name nombre_origen,
            ?::VARCHAR representacion,grupo grupo_uso,
            CASE WHEN c.fclass IS NULL THEN 'sin_seleccion'
                 WHEN o.code IS DISTINCT FROM c.code THEN 'codigo_no_concordante'
                 ELSE 'seleccionada' END estado_categoria,
            CASE WHEN estado_origen<>'admitida' THEN from_hex(wkb_hex) ELSE NULL END geometria_origen_no_admitida_wkb,
            CASE WHEN estado_origen='admitida' THEN
                ST_Transform(g,'OGC:CRS84','EPSG:25830',always_xy:=true) ELSE NULL END
                ::GEOMETRY('EPSG:25830') geometria,estado_origen
            FROM osm_geometria o LEFT JOIN categorias c ON o.fclass=c.fclass''',
            [origin + ':' + layer, layer, 'area' if area else 'punto'])
        connection.execute('''CREATE OR REPLACE TEMP TABLE osm_final AS SELECT * EXCLUDE(estado_origen),
            CASE WHEN estado_origen<>'admitida' THEN estado_origen
                 WHEN geometria IS NULL OR ST_IsEmpty(geometria) OR NOT ST_IsValid(geometria) THEN 'proyeccion_no_admitida'
                 ELSE 'admitida' END estado_geometria
            FROM osm_preparado''')
        if not counts:
            connection.execute('CREATE TABLE poi_elementos AS SELECT * FROM osm_final WHERE false')
        connection.execute('INSERT INTO poi_elementos SELECT * FROM osm_final')
        counts.append({'capa': layer, 'filas_bronze': count})
        print(f'OSM: {layer}, {count} registros conservados.', flush=True)
    comprobar_clave(connection, 'poi_elementos', OSM_KEYS['poi_elementos'])
    return counts


def osm_classify(connection):
    connection.execute('ALTER TABLE poi_elementos ADD COLUMN estado_ambito VARCHAR')
    connection.execute('''UPDATE poi_elementos SET estado_ambito=CASE
        WHEN estado_geometria<>'admitida' THEN 'no_evaluable'
        WHEN ST_Intersects(poi_elementos.geometria,a.geometria) THEN 'intersecta_cv' ELSE 'fuera_cv' END
        FROM ambito_cv a''')
    connection.execute('ALTER TABLE poi_elementos ADD COLUMN distancia_ambito_m DOUBLE')
    connection.execute('''UPDATE poi_elementos SET distancia_ambito_m=CASE
        WHEN estado_ambito='intersecta_cv' THEN 0.0
        WHEN estado_ambito='fuera_cv' THEN ST_Distance(poi_elementos.geometria,a.geometria) ELSE NULL END
        FROM ambito_cv a''')
    connection.execute('ALTER TABLE poi_elementos ADD COLUMN interes_para_revision BOOLEAN')
    connection.execute("UPDATE poi_elementos SET interes_para_revision=(estado_categoria='seleccionada' AND estado_geometria='admitida' AND estado_ambito='intersecta_cv')")


def osm_correspondences(connection):
    for representation in ('punto', 'area'):
        connection.execute(f'''CREATE TEMP TABLE {representation}s_revision AS SELECT * FROM poi_elementos
            WHERE estado_categoria='seleccionada' AND estado_geometria='admitida' AND representacion=?''', [representation])
    # shortcut: cruce separado por fallo de poda en spatial 1.5.0; revisar al actualizar el motor.
    connection.execute('''CREATE TEMP TABLE intersecciones AS
        SELECT p.origen_id origen_punto,p.fid_origen fid_punto,a.origen_id origen_area,a.fid_origen fid_area,
            p.categoria_origen categoria_punto,a.categoria_origen categoria_area
        FROM puntos_revision p JOIN areas_revision a ON ST_Intersects(p.geometria,a.geometria)''')
    connection.execute('''CREATE TEMP TABLE indicios AS
        SELECT origen_punto,fid_punto,origen_area,fid_area FROM intersecciones WHERE categoria_punto=categoria_area
        UNION
        SELECT p.origen_id origen_punto,p.fid_origen fid_punto,a.origen_id origen_area,a.fid_origen fid_area
        FROM puntos_revision p JOIN areas_revision a
        ON p.osm_id=a.osm_id AND p.categoria_origen=a.categoria_origen
        WHERE p.osm_id IS NOT NULL AND trim(p.osm_id)<>'' ''')
    connection.execute('''CREATE TABLE poi_correspondencias AS
        SELECT i.*,p.osm_id osm_id_punto,a.osm_id osm_id_area,p.categoria_origen,
            ST_Intersects(p.geometria,a.geometria) punto_intersecta_area,
            (coalesce(p.osm_id=a.osm_id,false) AND coalesce(trim(p.osm_id),'')<>'') coincide_id_publicado,
            (coalesce(trim(p.nombre_origen),'')<>'' AND
                lower(coalesce(trim(p.nombre_origen),''))=lower(coalesce(trim(a.nombre_origen),''))) coincide_nombre_normalizado,
            ST_Distance(p.geometria,a.geometria) distancia_plana_m,
            'posible_correspondencia_sin_fusion'::VARCHAR estado
        FROM indicios i JOIN puntos_revision p ON i.origen_punto=p.origen_id AND i.fid_punto=p.fid_origen
        JOIN areas_revision a ON i.origen_area=a.origen_id AND i.fid_area=a.fid_origen''')
    comprobar_clave(connection, 'poi_correspondencias', OSM_KEYS['poi_correspondencias'])
    for rep in ('punto', 'area'):
        if connection.execute(f'''SELECT count(*) FROM poi_correspondencias r ANTI JOIN poi_elementos p
            ON r.origen_{rep}=p.origen_id AND r.fid_{rep}=p.fid_origen''').fetchone()[0]:
            raise ValueError('Correspondencias sin elemento de origen')


def osm_validate(connection, counts):
    expected = sum(r['filas_bronze'] for r in counts)
    if connection.execute('SELECT count(*) FROM poi_elementos').fetchone()[0] != expected:
        raise ValueError('No se conservan todos los elementos OSM leídos')
    connection.execute('''CREATE TABLE incidencias AS
        SELECT origen_id,capa_origen,fid_origen,osm_id,'geometria'::VARCHAR campo,estado_geometria motivo
        FROM poi_elementos WHERE estado_geometria<>'admitida'
        UNION ALL SELECT origen_id,capa_origen,fid_origen,osm_id,'categoria',estado_categoria
        FROM poi_elementos WHERE estado_categoria='codigo_no_concordante'
        UNION ALL SELECT origen_id,capa_origen,fid_origen,osm_id,'osm_id','identificador_no_informado'
        FROM poi_elementos WHERE osm_id IS NULL OR trim(osm_id)='' ''')
    return {'filas_conservadas': expected, 'filas_eliminadas': 0,
        'por_capa': consultar(connection, '''SELECT capa_origen,count(*) filas,count(DISTINCT osm_id) ids_distintos,
            count(*) FILTER(WHERE interes_para_revision) para_revision FROM poi_elementos GROUP BY 1 ORDER BY 1'''),
        'por_categoria_seleccionada': consultar(connection, '''SELECT categoria_origen,grupo_uso,representacion,
            count(*) elementos,count(*) FILTER(WHERE interes_para_revision) para_revision
            FROM poi_elementos WHERE estado_categoria='seleccionada' GROUP BY 1,2,3 ORDER BY 2,1,3'''),
        'por_estado': consultar(connection, 'SELECT estado_geometria,estado_ambito,count(*) filas FROM poi_elementos GROUP BY 1,2 ORDER BY 1,2'),
        'fuera_ambito_mas_distantes': consultar(connection, '''SELECT capa_origen,fid_origen,osm_id,categoria_origen,nombre_origen,
            distancia_ambito_m FROM poi_elementos WHERE estado_ambito='fuera_cv' ORDER BY distancia_ambito_m DESC LIMIT 20'''),
        'identificadores_repetidos_por_capa': consultar(connection, '''SELECT capa_origen,count(*) ids_repetidos,sum(n-1) filas_adicionales
            FROM (SELECT capa_origen,osm_id,count(*) n FROM poi_elementos GROUP BY 1,2 HAVING count(*)>1) GROUP BY 1 ORDER BY 1'''),
        'correspondencias': consultar(connection, '''SELECT punto_intersecta_area,coincide_id_publicado,
            coincide_nombre_normalizado,count(*) pares FROM poi_correspondencias GROUP BY 1,2,3 ORDER BY 1,2,3'''),
        'incidencias': consultar(connection, 'SELECT campo,motivo,count(*) filas FROM incidencias GROUP BY 1,2 ORDER BY 1,2')}


# %% Silver: capacidad de demanda i-DE

CAPACIDAD_SOURCE = 'Fuentes/Demanda_energetica.csv'


CAPACIDAD_PROVINCES = {'Alicante/Alacant': '03', 'Castellón/Castelló': '12', 'Valencia/València': '46'}


CAPACIDAD_NULL = '\x1eNULL_SILVER\x1e'


CAPACIDAD_RAW = {
    'gestor_red': 'Gestor de red', 'provincia_nombre_origen': 'Provincia',
    'municipio_nombre_origen': 'Municipio', 'x_utm_origen': 'Coordenada UTM X',
    'y_utm_origen': 'Coordenada UTM Y', 'subestacion_origen': 'Subestación',
    'nivel_tension_kv_origen': 'Nivel de Tensión (kV)',
    'capacidad_firme_disponible_mw_origen': 'Capacidad firme disponible (MW)',
    'compromiso_regulatorio_origen': 'Capacidad comprometida por cuestiones regulatorias',
    'capacidad_ocupada_mw_origen': 'Capacidad de acceso firme de demanda ocupada (MW)',
    'capacidad_admitida_no_evaluada_mw_origen': 'Capacidad de acceso firme admitida y no evaluada (MW)',
    'posiciones_ocupadas_origen': 'Posiciones ocupadas', 'posiciones_libres_origen': 'Posiciones libres',
    'nudo_0_asterisco_origen': 'Nudo 0*', 'comentarios_origen': 'Comentarios',
    'nombre_punto_conexion': 'Denominación del Punto de Conexión',
    'punto_conexion_id': 'Identificador del Punto de Conexión',
}


CAPACIDAD_TYPED = {**dict.fromkeys(('nivel_tension_kv', 'capacidad_firme_disponible_mw',
    'capacidad_ocupada_mw', 'capacidad_admitida_no_evaluada_mw'), 'DECIMAL(18,2)'),
    'posiciones_ocupadas': 'BIGINT', 'posiciones_libres': 'BIGINT',
    'compromiso_regulatorio': 'BOOLEAN', 'nudo_0_asterisco': 'BOOLEAN'}


CAPACIDAD_SCHEMA = {**dict.fromkeys(CAPACIDAD_RAW, 'VARCHAR'), 'origen_id': 'VARCHAR', 'registro_origen': 'BIGINT',
    'codigo_provincia': 'VARCHAR', 'fecha_referencia': 'DATE', 'estado_fecha': 'VARCHAR',
    'estado_coordenadas': 'VARCHAR'}


for field, datatype in CAPACIDAD_TYPED.items():
    CAPACIDAD_SCHEMA.update({field: datatype, 'estado_' + field: 'VARCHAR'})


CAPACIDAD_INC_SCHEMA = dict.fromkeys(('origen_id', 'punto_conexion_id', 'campo', 'valor_origen', 'motivo'), 'VARCHAR')


CAPACIDAD_INC_SCHEMA['registro_origen'] = 'BIGINT'


def capacidad_typed_value(raw, datatype, positive=False):
    text = raw.strip()
    if not text:
        return None, 'no_informado'
    if datatype == 'BOOLEAN':
        return (text == '1', 'informado') if text in ('0', '1') else (None, 'error_conversion')
    pattern = r'[0-9]+' if datatype == 'BIGINT' else r'[0-9]+(?:,[0-9]{1,2})?'
    if len(text) > 64 or not re.fullmatch(pattern, text):
        return None, 'error_conversion'
    value = Decimal(text.replace(',', '.'))
    limit = Decimal(2**63) if datatype == 'BIGINT' else Decimal('10000000000000000')
    if value >= limit or (positive and value <= 0):
        return None, 'error_conversion'
    return (int(value) if datatype == 'BIGINT' else value), 'informado'


def capacidad_coordinate_state(x, y):
    if not x.strip() or not y.strip():
        return 'no_informadas'
    if any(len(v.strip()) > 64 or not re.fullmatch(r'[0-9]+(?:,[0-9]+)?', v.strip()) for v in (x, y)):
        return 'error_formato'
    # shortcut: no crear geometría hasta acreditar huso y datum de esta copia.
    return 'pendiente_crs'


def capacidad_prepare(connection, path, temporary, origin):
    national, selected, seen = Counter(), Counter(), set()
    count = 0
    with (temporary / 'capacidad.csv').open('w', encoding='utf-8', newline='') as out, \
            (temporary / 'incidencias.csv').open('w', encoding='utf-8', newline='') as incidents:
        writer, incident_writer = csv.writer(out), csv.writer(incidents)
        writer.writerow(CAPACIDAD_SCHEMA)
        incident_writer.writerow(CAPACIDAD_INC_SCHEMA)
        for ordinal, row in demografia_csv_rows(path, list(CAPACIDAD_RAW.values())):
            if any(CAPACIDAD_NULL in v for v in row.values()):
                raise ValueError('El texto de entrada contiene el marcador reservado de NULL')
            key = row[CAPACIDAD_RAW['punto_conexion_id']]
            if not key.strip() or key != key.strip() or key in seen:
                raise ValueError(f'Identificador nacional vacío, con espacios exteriores o repetido: registro {ordinal}')
            seen.add(key)
            province = row['Provincia']
            national[province] += 1
            if province not in CAPACIDAD_PROVINCES:
                continue
            selected[province] += 1
            count += 1
            result = {field: row[header] for field, header in CAPACIDAD_RAW.items()}
            result.update(origen_id=origin, registro_origen=ordinal, codigo_provincia=CAPACIDAD_PROVINCES[province],
                fecha_referencia=None, estado_fecha='no_acreditada',
                estado_coordenadas=capacidad_coordinate_state(result['x_utm_origen'], result['y_utm_origen']))

            def incidence(field, value, reason):
                entry = {'origen_id': origin, 'punto_conexion_id': key, 'campo': field,
                         'valor_origen': value, 'motivo': reason, 'registro_origen': ordinal}
                incident_writer.writerow([entry[k] for k in CAPACIDAD_INC_SCHEMA])

            for field, datatype in CAPACIDAD_TYPED.items():
                raw = result[field + '_origen']
                value, state = capacidad_typed_value(raw, datatype, positive=(field == 'nivel_tension_kv'))
                result[field], result['estado_' + field] = value, state
                if state == 'error_conversion':
                    incidence(field, raw, state)
            if result['estado_coordenadas'] != 'pendiente_crs':
                incidence('coordenadas', json.dumps([result['x_utm_origen'], result['y_utm_origen']]), result['estado_coordenadas'])
            for field in ('gestor_red', 'municipio_nombre_origen', 'subestacion_origen', 'nombre_punto_conexion'):
                if not result[field].strip():
                    incidence(field, result[field], 'no_informado')
            if result['nudo_0_asterisco'] is True and result['capacidad_firme_disponible_mw'] not in (None, Decimal(0)):
                incidence('nudo_0_asterisco', result['nudo_0_asterisco_origen'], 'capacidad_disponible_no_cero')
            writer.writerow([CAPACIDAD_NULL if result[k] is None else result[k] for k in CAPACIDAD_SCHEMA])
    if not count:
        raise ValueError('No hay puntos de conexión de las tres provincias previstas')
    for table, filename, schema in [('capacidad_demanda', 'capacidad.csv', CAPACIDAD_SCHEMA),
                                    ('incidencias', 'incidencias.csv', CAPACIDAD_INC_SCHEMA)]:
        connection.execute(f'''CREATE TABLE {table} AS SELECT * FROM read_csv(?,columns=?,
            header=true,delim=',',quote='"',escape='"',auto_detect=false,nullstr=?,
            strict_mode=true,ignore_errors=false)''', [str(temporary / filename), schema, CAPACIDAD_NULL])
    comprobar_clave(connection, 'capacidad_demanda', ['punto_conexion_id'])
    comprobar_clave(connection, 'capacidad_demanda', ['origen_id', 'registro_origen'])
    if connection.execute('SELECT count(*) FROM capacidad_demanda').fetchone()[0] != count:
        raise ValueError('Pérdida de registros CV al preparar i-DE')
    observed = dict(connection.execute('SELECT provincia_nombre_origen,count(*) FROM capacidad_demanda GROUP BY 1').fetchall())
    if observed != dict(selected):
        raise ValueError('Recuentos provinciales no reconciliados')
    if connection.execute('''SELECT count(*) FROM incidencias i ANTI JOIN capacidad_demanda c
            USING(origen_id,registro_origen,punto_conexion_id)''').fetchone()[0]:
        raise ValueError('Incidencia sin registro de origen')
    return {'filas_nacionales': sum(national.values()), 'filas_cv': count,
        'fuera_cv_en_bronze': sum(national.values()) - count, 'filas_cv_eliminadas': 0,
        'claves_nacionales_unicas': True, 'provincias_nacionales': dict(sorted(national.items()))}


def capacidad_validate(connection):
    return {'por_provincia': consultar(connection, '''SELECT codigo_provincia,provincia_nombre_origen,
        count(*) puntos,count(DISTINCT subestacion_origen) codigos_subestacion,
        count(DISTINCT municipio_nombre_origen) etiquetas_municipales
        FROM capacidad_demanda GROUP BY 1,2 ORDER BY 1'''),
        'medidas': {field: consultar(connection, f'''SELECT estado_{field} estado,count(*) filas,
            min({field})::VARCHAR minimo,max({field})::VARCHAR maximo
            FROM capacidad_demanda GROUP BY 1 ORDER BY 1''') for field in CAPACIDAD_TYPED},
        'banderas': consultar(connection, '''SELECT compromiso_regulatorio,nudo_0_asterisco,
            count(*) puntos FROM capacidad_demanda GROUP BY 1,2 ORDER BY 1,2'''),
        'capacidad_disponible': consultar(connection, '''SELECT CASE WHEN capacidad_firme_disponible_mw IS NULL
            THEN 'no_evaluable' WHEN capacidad_firme_disponible_mw=0 THEN 'cero' ELSE 'positiva' END estado,
            count(*) puntos FROM capacidad_demanda GROUP BY 1 ORDER BY 1'''),
        'coordenadas': consultar(connection, 'SELECT estado_coordenadas,count(*) puntos FROM capacidad_demanda GROUP BY 1'),
        'fechas': consultar(connection, 'SELECT estado_fecha,count(*) puntos FROM capacidad_demanda GROUP BY 1'),
        'incidencias': consultar(connection, 'SELECT campo,motivo,count(*) filas FROM incidencias GROUP BY 1,2 ORDER BY 1,2'),
        'numero_incidencias': connection.execute('SELECT count(*) FROM incidencias').fetchone()[0]}


# %% Silver: geometría y contraste municipal de recarga

RECARGA_GEO_RESOURCE = '047c1f64-7552-44da-a639-2291292e650d'


RECARGA_GEO_ENDPOINT = 'https://energia.serviciosmin.gob.es/Ripree/ExportarInstalaciones/Export'


RECARGA_GEO_KEYS = {'recarga_instalaciones_geo': ['instalacion_id'],
        'recarga_municipios_candidatos': ['instalacion_id', 'codigo_municipio'],
        'entidad_municipio': ['tabla_entidad', 'clave_entidad', 'codigo_municipio']}


def comprobar_archivo(root, path, expected=None):
    path = path.resolve()
    if not path.is_relative_to(root.resolve()):
        raise ValueError('Entrada fuera del proyecto')
    stamp = huella(path)
    if expected and any(stamp[k] != expected[k] for k in ('bytes', 'sha256')):
        raise ValueError(f'Entrada distinta de su informe o recibo: {path.name}')
    return {'ruta': path.relative_to(root.resolve()).as_posix(), **stamp}


def recarga_geo_evidence(root, path):
    stamp = comprobar_archivo(root, path)
    receipt_path = path.with_suffix(path.suffix + '.adquisicion.json')
    receipt_stamp = comprobar_archivo(root, receipt_path)
    receipt = json.loads(receipt_path.read_text(encoding='utf-8'))
    if any(stamp[k] != receipt[k] for k in ('bytes', 'sha256')):
        raise ValueError('El catálogo MITECO difiere del recibo conservado')
    catalog = json.loads(path.read_text(encoding='utf-8'))
    resources = [r for r in catalog['result']['resources'] if r['id'] == RECARGA_GEO_RESOURCE]
    if not catalog.get('success') or len(resources) != 1:
        raise ValueError('Recurso CSV MITECO no identificado de forma única')
    resource = resources[0]
    if (resource['format'] != 'CSV' or resource['url'] != RECARGA_GEO_ENDPOINT or
            resource['reference_system'] != 'http://www.opengis.net/def/crs/EPSG/0/4258'):
        raise ValueError('Declaración espacial o endpoint MITECO distintos de los investigados')
    return {'catalogo': stamp, 'recibo': receipt_stamp, 'recurso': RECARGA_GEO_RESOURCE,
            'url': receipt['url_solicitada'], 'crs_declarado': 'EPSG:4258'}


def recarga_geo_name_key(value):
    # Solo variantes de escritura; nunca buscar el municipio más parecido.
    return '/'.join(sorted(''.join(c for c in unicodedata.normalize('NFKD', part.casefold())
                                  if c.isalnum()) for part in value.split('/')))


def recarga_geo_prepare(connection, temporary):
    names = connection.execute('''SELECT nombre_municipio FROM municipios UNION SELECT nombre_municipio_ign
        FROM recintos UNION SELECT municipio_nombre_origen FROM recarga_instalaciones ORDER BY 1''').fetchall()
    path = temporary / 'nombres.csv'
    marker = '\x1eNULL_SILVER\x1e'
    with path.open('w', encoding='utf-8', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(['texto_origen', 'clave'])
        for (text,) in names:
            if text is not None:
                if marker in text:
                    raise ValueError('Nombre contiene marcador reservado de NULL')
                writer.writerow([text, recarga_geo_name_key(text)])
    connection.execute('''CREATE TEMP TABLE nombres_normalizados AS SELECT * FROM read_csv(?,
        columns={'texto_origen':'VARCHAR','clave':'VARCHAR'},header=true,delim=',',
        auto_detect=false,nullstr=?,strict_mode=true,ignore_errors=false)''', [str(path), marker])
    connection.execute('''CREATE TEMP TABLE etiquetas AS
        SELECT codigo_municipio,left(codigo_municipio,2) codigo_provincia,n.clave nombre
        FROM municipios m JOIN nombres_normalizados n ON m.nombre_municipio=n.texto_origen
        UNION SELECT codigo_municipio,left(codigo_municipio,2),n.clave FROM recintos r
            JOIN nombres_normalizados n ON r.nombre_municipio_ign=n.texto_origen''')
    connection.execute('''CREATE TEMP TABLE etiquetas_unicas AS SELECT codigo_provincia,nombre,
        min(codigo_municipio) codigo_municipio FROM etiquetas WHERE nombre<>''
        GROUP BY 1,2 HAVING count(DISTINCT codigo_municipio)=1''')
    connection.execute('''CREATE TABLE recarga_instalaciones_geo AS SELECT *,
        CASE provincia_nombre_origen WHEN 'Alicante/Alacant' THEN '03'
            WHEN 'Castellón/Castelló' THEN '12' WHEN 'Valencia/València' THEN '46' END codigo_provincia_declarada,
        CASE WHEN latitud IS NULL OR longitud IS NULL OR NOT isfinite(latitud) OR NOT isfinite(longitud)
            OR abs(latitud)>90 OR abs(longitud)>180 OR (latitud=0 AND longitud=0) THEN NULL
        ELSE TRY(ST_Transform(ST_Point(longitud,latitud),'EPSG:4258','EPSG:25830',always_xy:=true)) END
            ::GEOMETRY('EPSG:25830') geometria
        FROM recarga_instalaciones''')
    connection.execute('''ALTER TABLE recarga_instalaciones_geo ADD COLUMN estado_geometria VARCHAR''')
    connection.execute('''UPDATE recarga_instalaciones_geo SET estado_geometria=CASE
        WHEN geometria IS NULL THEN 'coordenadas_o_proyeccion_no_admitidas'
        WHEN ST_IsEmpty(geometria) OR NOT ST_IsValid(geometria) OR
            NOT isfinite(ST_X(geometria)) OR NOT isfinite(ST_Y(geometria)) THEN 'proyeccion_no_admitida'
        ELSE 'admitida' END''')
    # shortcut: materializar el cruce antes de filtros de atributos por fallo spatial 1.5.0; revisar al actualizar.
    connection.execute('''CREATE TABLE recarga_municipios_candidatos AS
        SELECT p.instalacion_id,r.codigo_municipio FROM recarga_instalaciones_geo p JOIN recintos r
        ON ST_Intersects(p.geometria,r.geometria)''')
    connection.execute('''DELETE FROM recarga_municipios_candidatos WHERE instalacion_id IN
        (SELECT instalacion_id FROM recarga_instalaciones_geo WHERE estado_geometria<>'admitida')''')
    connection.execute('''CREATE TEMP TABLE contraste AS SELECT p.instalacion_id,
        coalesce(c.n,0)::BIGINT numero_municipios_intersectados,
        CASE WHEN c.n=1 THEN c.codigo END codigo_municipio_espacial,
        e.codigo_municipio codigo_municipio_declarado
        FROM recarga_instalaciones_geo p
        LEFT JOIN (SELECT instalacion_id,count(*) n,min(codigo_municipio) codigo
            FROM recarga_municipios_candidatos GROUP BY 1) c USING(instalacion_id)
        LEFT JOIN nombres_normalizados n ON p.municipio_nombre_origen=n.texto_origen
        LEFT JOIN etiquetas_unicas e ON p.codigo_provincia_declarada=e.codigo_provincia AND n.clave=e.nombre''')
    connection.execute('''CREATE TEMP TABLE enriquecida AS SELECT p.*,c.* EXCLUDE(instalacion_id),
        m.nombre_municipio nombre_municipio_espacial,
        CASE WHEN p.estado_geometria<>'admitida' THEN 'geometria_no_admitida'
            WHEN c.numero_municipios_intersectados=0 THEN 'fuera_ambito'
            WHEN c.numero_municipios_intersectados>1 THEN 'limite_ambiguo'
            WHEN p.codigo_provincia_declarada IS NULL THEN 'provincia_no_reconocida'
            WHEN left(c.codigo_municipio_espacial,2)<>p.codigo_provincia_declarada THEN 'provincia_discrepante'
            WHEN c.codigo_municipio_declarado IS NULL THEN 'nombre_no_resuelto'
            WHEN c.codigo_municipio_declarado<>c.codigo_municipio_espacial THEN 'municipio_discrepante'
            ELSE 'concordante' END estado_contraste_territorial,
        CASE WHEN p.estado_geometria='admitida' AND rd.codigo_municipio IS NOT NULL
            THEN ST_Distance(p.geometria,rd.geometria) ELSE NULL END distancia_municipio_declarado_m
        FROM recarga_instalaciones_geo p JOIN contraste c USING(instalacion_id)
        LEFT JOIN municipios m ON c.codigo_municipio_espacial=m.codigo_municipio
        LEFT JOIN recintos rd ON c.codigo_municipio_declarado=rd.codigo_municipio''')
    connection.execute('''CREATE OR REPLACE TABLE recarga_instalaciones_geo AS SELECT *,
        (estado_contraste_territorial='concordante') admitida_asignacion_municipal FROM enriquecida''')
    connection.execute('''CREATE TABLE entidad_municipio AS SELECT 'recarga_instalaciones_geo'::VARCHAR tabla_entidad,
        instalacion_id clave_entidad,codigo_municipio_espacial codigo_municipio,
        'interseccion_unica_y_nombre_oficial'::VARCHAR metodo_asignacion,'contrastada'::VARCHAR estado_asignacion
        FROM recarga_instalaciones_geo WHERE admitida_asignacion_municipal''')
    connection.execute('''CREATE TABLE incidencias_espaciales AS SELECT instalacion_id,
        'contraste_territorial'::VARCHAR campo,estado_contraste_territorial motivo
        FROM recarga_instalaciones_geo WHERE NOT admitida_asignacion_municipal''')


def recarga_geo_validate(connection):
    for table, keys in RECARGA_GEO_KEYS.items():
        comprobar_clave(connection, table, keys)
    fields = ','.join(identificador_sql(r[0]) for r in connection.execute('DESCRIBE recarga_instalaciones').fetchall())
    if connection.execute(f'''SELECT count(*) FROM (
        (SELECT {fields} FROM recarga_instalaciones EXCEPT ALL SELECT {fields} FROM recarga_instalaciones_geo)
        UNION ALL (SELECT {fields} FROM recarga_instalaciones_geo EXCEPT ALL SELECT {fields} FROM recarga_instalaciones))''').fetchone()[0]:
        raise ValueError('Los atributos originales de instalaciones han cambiado')
    for table, field in [('recarga_municipios_candidatos', 'instalacion_id'), ('entidad_municipio', 'clave_entidad'),
                         ('incidencias_espaciales', 'instalacion_id')]:
        if connection.execute(f'''SELECT count(*) FROM {table} x ANTI JOIN recarga_instalaciones_geo p
            ON x.{field}=p.instalacion_id''').fetchone()[0]:
            raise ValueError(f'Referencia huérfana en {table}')
    if connection.execute('''SELECT count(*) FROM entidad_municipio x ANTI JOIN recintos r USING(codigo_municipio)''').fetchone()[0]:
        raise ValueError('Relación con municipio inexistente')
    error = connection.execute('''SELECT max(greatest(abs(ST_X(v)-longitud),abs(ST_Y(v)-latitud))) FROM
        (SELECT *,ST_Transform(geometria,'EPSG:25830','EPSG:4258',always_xy:=true) v
        FROM recarga_instalaciones_geo WHERE estado_geometria='admitida')''').fetchone()[0]
    if error is not None and error > 1e-8:
        raise ValueError('La vuelta de proyección altera las coordenadas')
    controls = {}
    for table in RECARGA_KEYS:
        original = connection.execute(f'SELECT count(*) FROM {table}').fetchone()[0]
        joined, admitted = connection.execute(f'''SELECT count(*),count(*) FILTER(WHERE g.admitida_asignacion_municipal)
            FROM {table} p JOIN recarga_instalaciones_geo g USING(instalacion_id)''').fetchone()
        if original != joined:
            raise ValueError('Se pierden o multiplican entidades al cruzar instalaciones')
        controls[table] = {'originales': original, 'tras_relacion': joined, 'con_asignacion_admitida': admitted}
    return {'entidades': controls, 'error_vuelta_grados_max': error,
        'por_estado': consultar(connection, 'SELECT estado_geometria,estado_contraste_territorial,count(*) instalaciones FROM recarga_instalaciones_geo GROUP BY 1,2 ORDER BY 1,2'),
        'municipios_con_instalaciones_admitidas': connection.execute('SELECT count(DISTINCT codigo_municipio) FROM entidad_municipio').fetchone()[0],
        'discrepancias_por_etiqueta': consultar(connection, '''SELECT provincia_nombre_origen,municipio_nombre_origen,
            codigo_municipio_declarado,codigo_municipio_espacial,nombre_municipio_espacial,count(*) instalaciones,
            max(distancia_municipio_declarado_m) distancia_maxima_m FROM recarga_instalaciones_geo
            WHERE NOT admitida_asignacion_municipal GROUP BY 1,2,3,4,5 ORDER BY instalaciones DESC,1,2,4'''),
        'incidencias': connection.execute('SELECT count(*) FROM incidencias_espaciales').fetchone()[0]}


# %% Verificación de metadatos i-DE

METADATOS_IDE_PAGE = 'https://www.i-de.es/es/conexion-red-electrica/suministro-electrico/mapa-capacidad-consumo'


METADATOS_IDE_BOE = 'https://www.boe.es/buscar/doc.php?id=BOE-A-2025-9913'


def metadatos_ide_normalized(text):
    return ' '.join(unicodedata.normalize('NFKC', text).split())


class metadatos_ide_Document(HTMLParser):
    def __init__(self):
        super().__init__()
        self.text = []
        self.links = set()
        self.hidden = 0

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style'):
            self.hidden += 1
        if tag == 'a':
            self.links.update(v for k, v in attrs if k == 'href')

    def handle_endtag(self, tag):
        if tag in ('script', 'style'):
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden:
            self.text.append(data)


def metadatos_ide_verify(root, evidence):
    if evidence['version'] != 1 or evidence['recurso_id'] != 'ide_capacidad_demanda':
        raise ValueError('Contrato de evidencia i-DE no admitido')

    def checked(entry):
        path = (root / entry['ruta']).resolve()
        if not path.is_relative_to(root.resolve()):
            raise ValueError('Evidencia fuera del proyecto')
        stamp = huella(path)
        if any(stamp[k] != entry[k] for k in ('bytes', 'sha256')):
            raise ValueError(f"Huella distinta: {entry['ruta']}")
        return path

    paths = {name: checked(entry) for name, entry in evidence['archivos'].items()}
    source, official = paths['original'], paths['csv_oficial']
    if evidence['archivos']['original']['sha256'] != evidence['archivos']['csv_oficial']['sha256']:
        raise ValueError('El CSV oficial no coincide con el original')
    with source.open('rb') as a, official.open('rb') as b:
        while True:
            left, right = a.read(1024 * 1024), b.read(1024 * 1024)
            if left != right:
                raise ValueError('Los CSV no son idénticos byte por byte')
            if not left:
                break
    with source.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.reader(stream, delimiter=';')
        if next(reader) != list(CAPACIDAD_RAW.values()):
            raise ValueError('Cabecera ajena al formato comprobado')
        rows = list(reader)
    if any(len(row) != len(CAPACIDAD_RAW) for row in rows):
        raise ValueError('Ancho de registro i-DE incorrecto')
    if len({r[-1] for r in rows}) != len(rows) or any(not r[-1].strip() for r in rows):
        raise ValueError('Identificadores vacíos o repetidos')
    provinces = Counter(r[1] for r in rows if r[1] in CAPACIDAD_PROVINCES)
    if len(rows) != evidence['filas_nacionales'] or dict(provinces) != evidence['filas_cv_provincia']:
        raise ValueError('Recuentos diferentes de la edición acreditada')

    published = date.fromisoformat(evidence['fecha_publicacion'])
    url = evidence['url_csv_oficial']
    if not url.startswith('https://www.i-de.es/documents/') or published.strftime('%Y_%m_%d') not in url:
        raise ValueError('Fecha y enlace de edición incoherentes')
    page = json.loads(paths['pagina_ide'].read_text(encoding='utf-8'))
    regulation = metadatos_ide_Document()
    regulation.feed(paths['norma'].read_text(encoding='utf-8'))
    if url not in {entry['url'] for entry in page['enlaces']}:
        raise ValueError('El histórico conservado no enlaza el CSV')
    text = metadatos_ide_normalized(page['texto'])
    if 'resolución de 30 de abril de 2025' not in text:
        raise ValueError('No consta la declaración de formato del proveedor')
    text = metadatos_ide_normalized(' '.join(regulation.text))
    start = '1. Formato para la publicación de las capacidades de acceso para instalaciones de demanda'
    end = '2. Formato para la publicación de las capacidades de acceso para instalaciones de generación'
    if start not in text or end not in text:
        raise ValueError('No se identifica el apartado de demanda de la norma')
    demand = text.split(start, 1)[1].split(end, 1)[0]
    for field in ('X', 'Y'):
        fragment = f'Coordenada UTM {field}. Coordenadas de la ubicación de la subestación, en el sistema ETRS89 UTM referidas al huso 30 para la península'
        if fragment not in demand:
            raise ValueError('La norma conservada no acredita CRS en el apartado de demanda')
    rule = 'dentro de los cinco primeros días del mes siguiente al que se refiere la información'
    if rule not in text or not 1 <= published.day <= 5:
        raise ValueError('No se acredita la regla del periodo mensual')
    month = date(published.year - (published.month == 1), 12 if published.month == 1 else published.month - 1, 1)
    if evidence['periodo_referencia'] != month.strftime('%Y-%m'):
        raise ValueError('Periodo mensual incoherente con la publicación')
    if evidence['crs_cv'] != 'EPSG:25830' or evidence['fecha_corte_exacta'] is not None:
        raise ValueError('Metadatos incompatibles con el alcance acreditado')

    # Comprobar de nuevo las evidencias tras la lectura; no se modifica ningún insumo.
    for entry in evidence['archivos'].values():
        checked(entry)
    return {'estado': 'fecha_y_crs_acreditados_documentalmente',
            'comprobado_utc': datetime.now(timezone.utc).isoformat(),
            'identidad_csv': 'igualdad_binaria_completa', 'sha256_original': evidence['archivos']['original']['sha256'],
            'filas_nacionales': len(rows), 'filas_cv': sum(provinces.values()),
            'fecha_publicacion': published.isoformat(), 'periodo_referencia': evidence['periodo_referencia'],
            'criterio_periodo': 'mes_anterior_segun_formato_publicacion_CNMC',
            'fecha_corte_exacta': None, 'crs_cv': 'EPSG:25830',
            'referencia_crs': 'declaracion_de_formato_ide_y_anexo_demanda_BOE_A_2025_9913',
            'pagina': METADATOS_IDE_PAGE, 'norma': METADATOS_IDE_BOE,
            'alcance': 'Acreditación de metadatos; no genera capas, asignaciones municipales o una nueva ejecución Silver.'}


# %% Silver: geometría y contraste municipal de i-DE

CAPACIDAD_GEO_KEYS = {'capacidad_demanda_geo': ['punto_conexion_id'],
        'capacidad_municipios_candidatos': ['punto_conexion_id', 'codigo_municipio'],
        'entidad_municipio': ['punto_conexion_id'],
        'incidencias_espaciales': ['punto_conexion_id']}


CAPACIDAD_GEO_HISTORICAL = {'estado_fecha': 'estado_fecha_preparacion',
              'estado_coordenadas': 'estado_coordenadas_preparacion'}


def capacidad_geo_prepare(c, metadata):
    comprobar_clave(c, 'capacidad_demanda', ['punto_conexion_id'])
    # Variantes oficiales INE/IGN dentro de la provincia. Sin similitud aproximada.
    labels = c.execute('''SELECT codigo_municipio,nombre_municipio FROM municipios
        UNION SELECT codigo_municipio,nombre_municipio_ign FROM recintos''').fetchall()
    aliases = sorted({(code[:2], recarga_geo_name_key(variant), code) for code, label in labels
                      for variant in [label, *label.split('/')] if variant.strip()})
    c.execute('CREATE TEMP TABLE etiquetas(provincia VARCHAR,nombre VARCHAR,codigo_municipio VARCHAR)')
    c.executemany('INSERT INTO etiquetas VALUES (?,?,?)', aliases)
    names = c.execute('SELECT DISTINCT municipio_nombre_origen FROM capacidad_demanda').fetchall()
    c.execute('CREATE TEMP TABLE nombres(texto VARCHAR,clave VARCHAR)')
    c.executemany('INSERT INTO nombres VALUES (?,?)', [(t, recarga_geo_name_key(t or '')) for (t,) in names])
    c.execute('''CREATE TEMP TABLE etiquetas_unicas AS SELECT provincia,nombre,min(codigo_municipio) codigo
        FROM etiquetas WHERE nombre<>'' GROUP BY 1,2 HAVING count(DISTINCT codigo_municipio)=1''')
    c.execute(f'''CREATE TABLE capacidad_demanda_geo AS SELECT * RENAME(
        estado_fecha AS estado_fecha_preparacion,estado_coordenadas AS estado_coordenadas_preparacion),
        {texto_sql(metadata['fecha_publicacion'])}::DATE fecha_publicacion,
        {texto_sql(metadata['periodo_referencia'])}::VARCHAR periodo_referencia,
        'periodo_mensual_acreditado'::VARCHAR estado_fecha,
        'crs_acreditado'::VARCHAR estado_coordenadas,'EPSG:25830'::VARCHAR crs_origen,
        CASE WHEN regexp_full_match(trim(x_utm_origen),'[0-9]+(,[0-9]+)?') AND length(trim(x_utm_origen))<=64
            THEN try_cast(replace(trim(x_utm_origen),',','.') AS DOUBLE) END x_utm,
        CASE WHEN regexp_full_match(trim(y_utm_origen),'[0-9]+(,[0-9]+)?') AND length(trim(y_utm_origen))<=64
            THEN try_cast(replace(trim(y_utm_origen),',','.') AS DOUBLE) END y_utm
        FROM capacidad_demanda''')
    c.execute('''ALTER TABLE capacidad_demanda_geo ADD COLUMN geometria GEOMETRY('EPSG:25830')''')
    c.execute('''UPDATE capacidad_demanda_geo SET geometria=ST_Point(x_utm,y_utm)::GEOMETRY('EPSG:25830')
        WHERE x_utm IS NOT NULL AND y_utm IS NOT NULL AND isfinite(x_utm) AND isfinite(y_utm)
            AND NOT (x_utm=0 AND y_utm=0)''')
    c.execute('ALTER TABLE capacidad_demanda_geo ADD COLUMN estado_geometria VARCHAR')
    c.execute('''UPDATE capacidad_demanda_geo SET estado_geometria=CASE
        WHEN geometria IS NULL THEN 'coordenadas_no_admitidas'
        WHEN ST_IsEmpty(geometria) OR NOT ST_IsValid(geometria) THEN 'geometria_no_admitida'
        ELSE 'admitida' END''')
    # Materializar antes de filtrar atributos: requisito observado con spatial 1.5.0.
    c.execute('''CREATE TABLE capacidad_municipios_candidatos AS SELECT p.punto_conexion_id,r.codigo_municipio
        FROM capacidad_demanda_geo p JOIN recintos r ON ST_Intersects(p.geometria,r.geometria)''')
    c.execute('''DELETE FROM capacidad_municipios_candidatos WHERE punto_conexion_id IN
        (SELECT punto_conexion_id FROM capacidad_demanda_geo WHERE estado_geometria<>'admitida')''')
    c.execute('''CREATE TEMP TABLE contraste AS SELECT p.punto_conexion_id,
        coalesce(x.n,0)::BIGINT numero_municipios_intersectados,
        CASE WHEN x.n=1 THEN x.codigo END codigo_municipio_espacial,e.codigo codigo_municipio_declarado
        FROM capacidad_demanda_geo p LEFT JOIN
        (SELECT punto_conexion_id,count(*) n,min(codigo_municipio) codigo
         FROM capacidad_municipios_candidatos GROUP BY 1) x USING(punto_conexion_id)
        LEFT JOIN nombres n ON p.municipio_nombre_origen=n.texto
        LEFT JOIN etiquetas_unicas e ON p.codigo_provincia=e.provincia AND n.clave=e.nombre''')
    c.execute('''CREATE TEMP TABLE enriquecida AS SELECT p.*,x.* EXCLUDE(punto_conexion_id),
        m.nombre_municipio nombre_municipio_espacial,
        CASE WHEN p.estado_geometria<>'admitida' THEN 'geometria_no_admitida'
        WHEN x.numero_municipios_intersectados=0 THEN 'fuera_ambito'
        WHEN x.numero_municipios_intersectados>1 THEN 'limite_ambiguo'
        WHEN left(x.codigo_municipio_espacial,2)<>p.codigo_provincia THEN 'provincia_discrepante'
        WHEN x.codigo_municipio_declarado IS NULL THEN 'nombre_no_resuelto'
        WHEN x.codigo_municipio_declarado<>x.codigo_municipio_espacial THEN 'municipio_discrepante'
        ELSE 'concordante' END estado_contraste_territorial,
        CASE WHEN p.estado_geometria='admitida' AND rd.codigo_municipio IS NOT NULL
        THEN ST_Distance(p.geometria,rd.geometria) END distancia_municipio_declarado_m
        FROM capacidad_demanda_geo p JOIN contraste x USING(punto_conexion_id)
        LEFT JOIN municipios m ON x.codigo_municipio_espacial=m.codigo_municipio
        LEFT JOIN recintos rd ON x.codigo_municipio_declarado=rd.codigo_municipio''')
    c.execute('''CREATE OR REPLACE TABLE capacidad_demanda_geo AS SELECT *,
        (estado_contraste_territorial='concordante') admitida_asignacion_municipal FROM enriquecida''')
    c.execute('''CREATE TABLE entidad_municipio AS SELECT punto_conexion_id,codigo_municipio_espacial codigo_municipio,
        'interseccion_unica_y_nombre_oficial'::VARCHAR metodo_asignacion FROM capacidad_demanda_geo
        WHERE admitida_asignacion_municipal''')
    c.execute('''CREATE TABLE incidencias_espaciales AS SELECT punto_conexion_id,
        estado_contraste_territorial motivo FROM capacidad_demanda_geo WHERE NOT admitida_asignacion_municipal''')


def capacidad_geo_validate(c):
    for name, keys in CAPACIDAD_GEO_KEYS.items():
        comprobar_clave(c, name, keys)
    original = [r[0] for r in c.execute('DESCRIBE capacidad_demanda').fetchall()]
    fields = ','.join(identificador_sql(k) for k in original)
    restored = ','.join(identificador_sql(CAPACIDAD_GEO_HISTORICAL.get(k, k)) + ' AS ' + identificador_sql(k) for k in original)
    if c.execute(f'''SELECT count(*) FROM (
        (SELECT {fields} FROM capacidad_demanda EXCEPT ALL SELECT {restored} FROM capacidad_demanda_geo)
        UNION ALL (SELECT {restored} FROM capacidad_demanda_geo EXCEPT ALL SELECT {fields} FROM capacidad_demanda))''').fetchone()[0]:
        raise ValueError('Se han modificado o perdido atributos de la tabla original')
    if c.execute('''SELECT count(*) FROM capacidad_demanda_geo WHERE estado_geometria='admitida'
        AND (NOT ST_IsValid(geometria) OR ST_IsEmpty(geometria) OR ST_X(geometria)<>x_utm OR ST_Y(geometria)<>y_utm)''').fetchone()[0]:
        raise ValueError('Geometría distinta de las coordenadas publicadas')
    if c.execute('SELECT count(*) FROM entidad_municipio ANTI JOIN municipios USING(codigo_municipio)').fetchone()[0]:
        raise ValueError('Relación municipal sin referencia')
    total, admitted = c.execute('SELECT count(*),count(*) FILTER(WHERE admitida_asignacion_municipal) FROM capacidad_demanda_geo').fetchone()
    if total != c.execute('SELECT count(*) FROM capacidad_demanda').fetchone()[0]:
        raise ValueError('Se pierden puntos al generar la capa')
    for table in ('entidad_municipio', 'incidencias_espaciales', 'capacidad_municipios_candidatos'):
        if c.execute(f'SELECT count(*) FROM {table} ANTI JOIN capacidad_demanda_geo USING(punto_conexion_id)').fetchone()[0]:
            raise ValueError('Referencia espacial i-DE huérfana')
    if admitted != c.execute('SELECT count(*) FROM entidad_municipio').fetchone()[0] or total-admitted != c.execute('SELECT count(*) FROM incidencias_espaciales').fetchone()[0]:
        raise ValueError('Admisiones e incidencias no reconciliadas')
    return {'puntos_conservados': total, 'asignaciones_admitidas': admitted, 'pendientes': total-admitted,
        'municipios_admitidos': c.execute('SELECT count(DISTINCT codigo_municipio) FROM entidad_municipio').fetchone()[0],
        'por_estado': consultar(c, '''SELECT estado_geometria,estado_contraste_territorial,count(*) puntos
            FROM capacidad_demanda_geo GROUP BY 1,2 ORDER BY 1,2'''),
        'por_provincia': consultar(c, '''SELECT codigo_provincia,count(*) puntos,
            count(*) FILTER(WHERE admitida_asignacion_municipal) admitidos FROM capacidad_demanda_geo GROUP BY 1 ORDER BY 1'''),
        'atributos_originales': 'conservados_integramente', 'coordenadas': 'sin_desplazamiento_ni_reproyeccion'}


# %% Silver: revisión documental y relaciones entre fuentes

INTEGRACION_SPATIAL = (
    ('poi_osm', 'poi_elementos', ['origen_id', 'fid_origen'], 'POI'),
    ('trafico_red', 'aforos_tramo_anio', ['tramo_aforo_id', 'anio'], 'LINEA'),
    ('trafico_red', 'red_tramos', ['provincia_origen', 'tramo_id'], 'LINEA'),
    ('capacidad_demanda_espacial', 'capacidad_demanda_geo', ['punto_conexion_id'], 'PUNTO'),
)


INTEGRACION_OUTPUT_KEYS = {
    'revision_recarga': ['instalacion_id'],
    'recarga_municipio_validado': ['instalacion_id'],
    'relaciones_espaciales': ['tabla_entidad', 'clave_entidad', 'codigo_municipio'],
    'estado_relaciones_espaciales': ['tabla_entidad', 'clave_entidad'],
    'cobertura_estadistica': ['tabla_entidad', 'periodo', 'segmento', 'codigo_municipio'],
    'estado_fuentes': ['bloque'],
}


def integracion_table(block, name):
    return identificador_sql(block + '__' + name)


def integracion_normalized(value):
    return ' '.join(''.join(c if c.isalnum() else ' ' for c in
        unicodedata.normalize('NFKD', (value or '').casefold()) if not unicodedata.combining(c)).split())


def integracion_within(directory, relative):
    path = (directory / relative).resolve()
    if not path.is_relative_to(directory.resolve()):
        raise ValueError('Ruta fuera de la carpeta esperada')
    return path


def integracion_scalar(c, sql):
    return c.execute(sql).fetchone()[0]


def exigir_cero(c, sql, message):
    if integracion_scalar(c, sql):
        raise ValueError(message)


def integracion_integrity(c):
    m = integracion_table('base_territorial', 'municipios')
    r = integracion_table('base_territorial', 'recintos_municipales')
    if integracion_scalar(c, f'SELECT count(*) FROM {m}') != 542:
        raise ValueError('Referencia municipal incompleta')
    exigir_cero(c, f'SELECT count(*) FROM {m} ANTI JOIN {r} USING(codigo_municipio)', 'Municipio sin recinto')
    for name in ('parque_municipal','parque_detalle_observado','poblacion_municipal','renta_territorial'):
        source = integracion_table('demografia_parque', name)
        exigir_cero(c, f'SELECT count(*) FROM {source} s ANTI JOIN {m} m USING(codigo_municipio) WHERE s.codigo_municipio IS NOT NULL', 'Código municipal sin referencia')
        before = integracion_scalar(c, f'SELECT count(*) FROM {source}')
        after = integracion_scalar(c, f'SELECT count(*) FROM {source} LEFT JOIN {m} USING(codigo_municipio)')
        if before != after:
            raise ValueError('Multiplicación en relación estadística')
    hist = integracion_table('demografia_parque', 'renta_codigos_historicos')
    exigir_cero(c, f'SELECT count(*) FROM {hist} JOIN {m} USING(codigo_municipio)', 'Código histórico incorporado al ámbito vigente')
    i, p, k = [integracion_table('recarga', t) for t in ('recarga_instalaciones','recarga_puntos','recarga_conectores')]
    exigir_cero(c, f'SELECT count(*) FROM {p} ANTI JOIN {i} USING(instalacion_id)', 'Punto sin instalación')
    exigir_cero(c, f'SELECT count(*) FROM {k} ANTI JOIN {p} USING(instalacion_id,punto_id)', 'Conector sin punto')
    g=integracion_table('recarga_espacial','recarga_instalaciones_geo')
    exigir_cero(c, f'SELECT count(*) FROM {i} FULL JOIN {g} USING(instalacion_id) WHERE {i}.instalacion_id IS NULL OR {g}.instalacion_id IS NULL', 'Instalaciones espaciales sin correspondencia')
    fields = ','.join(identificador_sql(r[0]) for r in c.execute(f'DESCRIBE {i}').fetchall())
    exigir_cero(c, f'SELECT count(*) FROM ((SELECT {fields} FROM {i} EXCEPT ALL SELECT {fields} FROM {g}) UNION ALL (SELECT {fields} FROM {g} EXCEPT ALL SELECT {fields} FROM {i}))', 'Atributos de recarga cambiados')
    tr, tv = integracion_table('trafico_red','red_tramos'), integracion_table('trafico_red','red_tramo_vial')
    exigir_cero(c, f'SELECT count(*) FROM {tv} ANTI JOIN {tr} USING(provincia_origen,tramo_id)', 'Asociación viaria huérfana')
    o, pairs = integracion_table('poi_osm','poi_elementos'), integracion_table('poi_osm','poi_correspondencias')
    for side in ('punto','area'):
        exigir_cero(c, f'SELECT count(*) FROM {pairs} x ANTI JOIN {o} o ON x.origen_{side}=o.origen_id AND x.fid_{side}=o.fid_origen', 'Correspondencia OSM huérfana')
    capacity = integracion_capacity_integrity(c)
    return {'integridad_referencial':'correcta', 'conservacion_atributos_recarga':'correcta',
            'capacidad_demanda':capacity,
            'detalle_dgt_sin_municipio':integracion_scalar(c, f"SELECT count(*) FROM {integracion_table('demografia_parque','parque_detalle_observado')} WHERE codigo_municipio IS NULL"),
            'renta_historica_separada':integracion_scalar(c, f'SELECT count(*) FROM {hist}')}


def integracion_capacity_integrity(c):
    source = integracion_table('capacidad_demanda', 'capacidad_demanda')
    geo = integracion_table('capacidad_demanda_espacial', 'capacidad_demanda_geo')
    relation = integracion_table('capacidad_demanda_espacial', 'entidad_municipio')
    fields = [r[0] for r in c.execute(f'DESCRIBE {source}').fetchall()]
    original = ','.join(identificador_sql(k) for k in fields)
    restored = ','.join(identificador_sql(CAPACIDAD_GEO_HISTORICAL.get(k,k))+' AS '+identificador_sql(k) for k in fields)
    exigir_cero(c, f'''SELECT count(*) FROM ((SELECT {original} FROM {source}
        EXCEPT ALL SELECT {restored} FROM {geo}) UNION ALL (SELECT {restored} FROM {geo}
        EXCEPT ALL SELECT {original} FROM {source}))''', 'Atributos i-DE originales cambiados')
    exigir_cero(c, f'''SELECT count(*) FROM ((SELECT punto_conexion_id,codigo_municipio FROM {relation}
        EXCEPT ALL SELECT punto_conexion_id,codigo_municipio_espacial FROM {geo} WHERE admitida_asignacion_municipal)
        UNION ALL (SELECT punto_conexion_id,codigo_municipio_espacial FROM {geo} WHERE admitida_asignacion_municipal
        EXCEPT ALL SELECT punto_conexion_id,codigo_municipio FROM {relation}))''', 'Relaciones i-DE admitidas incoherentes')
    exigir_cero(c, f'''SELECT count(*) FROM {geo} WHERE crs_origen IS DISTINCT FROM 'EPSG:25830' OR fecha_referencia IS NOT NULL
        OR estado_fecha IS DISTINCT FROM 'periodo_mensual_acreditado' OR estado_coordenadas IS DISTINCT FROM 'crs_acreditado'
        OR fecha_publicacion IS NULL OR periodo_referencia IS NULL
        OR periodo_referencia<>strftime(fecha_publicacion-INTERVAL 1 MONTH,'%Y-%m')
        OR (estado_geometria='admitida' AND (geometria IS NULL OR x_utm IS NULL OR y_utm IS NULL OR NOT ST_IsValid(geometria)
            OR ST_IsEmpty(geometria) OR ST_X(geometria)<>x_utm OR ST_Y(geometria)<>y_utm))''',
        'Metadatos o coordenadas i-DE incoherentes')
    exigir_cero(c, f'''SELECT count(*) FROM {geo} WHERE admitida_asignacion_municipal IS NULL
        OR admitida_asignacion_municipal IS DISTINCT FROM (estado_contraste_territorial='concordante')
        OR (admitida_asignacion_municipal AND (estado_geometria IS DISTINCT FROM 'admitida'
            OR numero_municipios_intersectados IS DISTINCT FROM 1
            OR codigo_municipio_espacial IS DISTINCT FROM codigo_municipio_declarado
            OR left(codigo_municipio_espacial,2) IS DISTINCT FROM codigo_provincia))''',
        'Admisión municipal i-DE incompatible con el contraste')
    for name in ('entidad_municipio', 'capacidad_municipios_candidatos'):
        current = integracion_table('capacidad_demanda_espacial', name)
        exigir_cero(c, f'SELECT count(*) FROM {current} ANTI JOIN {integracion_table("base_territorial","municipios")} USING(codigo_municipio)', 'Municipio i-DE sin referencia')
        exigir_cero(c, f'SELECT count(*) FROM {current} ANTI JOIN {geo} USING(punto_conexion_id)', 'Punto i-DE huérfano')
    return {'puntos':integracion_scalar(c,f'SELECT count(*) FROM {geo}'),
            'asignaciones_admitidas':integracion_scalar(c,f'SELECT count(*) FROM {relation}'),
            'pendientes':integracion_scalar(c,f'SELECT count(*) FROM {geo} WHERE NOT admitida_asignacion_municipal')}


def integracion_provider_evidence(root, path, c):
    """Decisiones explícitas acotadas, respaldadas por copias y recibos; sin fuzzy."""
    config_stamp = comprobar_archivo(root, path)
    config = json.loads(path.read_text(encoding='utf-8'))
    files, decisions = [config_stamp], {}
    for item in config['decisiones']:
        key = item['instalacion_id']
        if item['decision'] not in ('admitir', 'mantener_pendiente'):
            raise ValueError('Estado de decisión documental desconocido')
        if key in decisions:
            raise ValueError('Decisión documental duplicada')
        source = integracion_within(root, item['copia'])
        stamp = comprobar_archivo(root, source, item)
        receipt_path = source.with_suffix(source.suffix+'.adquisicion.json')
        files.extend([stamp,comprobar_archivo(root,receipt_path)])
        receipt = json.loads(receipt_path.read_text(encoding='utf-8'))
        if receipt['url_solicitada'] != item['url'] or any(receipt[k]!=stamp[k] for k in ('bytes','sha256')):
            raise ValueError('Evidencia documental distinta del recibo')
        kind = item.get('tipo_evidencia', 'ficha_jsonld')
        geo = None
        if kind == 'ficha_jsonld':
            scripts = re.findall(r'<script[^>]*type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',source.read_text(encoding='utf-8'),re.S)
            structured = [json.loads(s) for s in scripts]
            nodes = [node for obj in structured for node in (obj.get('@graph',[obj]) if isinstance(obj,dict) else [])]
            expected_type = item.get('tipo_jsonld', 'Store')
            if expected_type not in ('Store', 'ParkingFacility', 'GasStation'):
                raise ValueError('Tipo de ficha no admitido')
            stores = [s for s in nodes if isinstance(s,dict) and s.get('@type')==expected_type and s.get('url')==item['url']]
            if len(stores)!=1:
                raise ValueError('Ficha de establecimiento no identificada de forma única')
            address=stores[0]['address']
            if integracion_normalized(address['streetAddress'])!=integracion_normalized(item['direccion_documentada']) or integracion_normalized(address['addressLocality'])!=integracion_normalized(item['municipio_documentado']):
                raise ValueError('La ficha no respalda la dirección y municipio revisados')
            geo = stores[0].get('geo')
        elif kind in ('portal_cartociudad', 'via_cartociudad', 'registro_operador_json'):
            records = json.loads(source.read_text(encoding='utf-8'))
            selector = item['selector']
            matches = [r for r in records if isinstance(r,dict) and r.get(selector['campo'])==selector['valor']]
            if len(matches)!=1 or not item['campos_documentados']:
                raise ValueError('Registro documental no identificado de forma única')
            record = matches[0]
            if any(record.get(k)!=v for k,v in item['campos_documentados'].items()):
                raise ValueError('El registro no respalda los campos documentados')
            if kind in ('portal_cartociudad', 'via_cartociudad'):
                # Solo direcciones postales: los POI de recarga de CartoCiudad derivan de MITECO.
                params=parse_qs(urlparse(item['url']).query)
                if (params.get('cod_postal_filter')!=[item['codigo_postal_origen']] or
                        any(k in params for k in ('municipio_filter','poblacion_filter','lon','lat')) or
                        record.get('muniCode')!=item['codigo_municipio'] or record.get('state')!=0):
                    raise ValueError('Consulta postal incompatible con la revisión')
                if kind == 'portal_cartociudad':
                    if (record.get('type')!='portal' or record.get('postalCode')!=item['codigo_postal_origen'] or
                            record.get('portalNumber')!=item['portal_documentado']):
                        raise ValueError('Portal oficial incompatible con la revisión')
                    geo={'longitude':record['lng'],'latitude':record['lat']}
                elif (record.get('type')!='callejero' or
                        {r.get('muniCode') for r in records}!={item['codigo_municipio']} or
                        any(r.get('type')!='callejero' for r in records)):
                    raise ValueError('Vial ambiguo o no admitido')
                # Un vial acredita la referencia municipal de la calle, no un portal ni el cargador.
            else:
                if (record.get('direccion')!=item['direccion_documentada'] or
                        record.get('ciudad')!=item['municipio_documentado']):
                    raise ValueError('Registro del operador sin dirección/localidad concordante')
                geo={'longitude':record['lng'],'latitude':record['lat']}
        elif kind == 'texto_primario_revisado':
            # Extractos de una página primaria revisados explícitamente, nunca búsquedas fuzzy.
            content=source.read_text(encoding='utf-8')
            content=re.sub(r'<(script|style)\b[^>]*>.*?</\1>', ' ', content, flags=re.S|re.I)
            plain=integracion_normalized(unescape(re.sub(r'<[^>]*>', ' ', content)))
            fragments=item['fragmentos_documentados']
            if (len(fragments)<3 or any(not integracion_normalized(v) or integracion_normalized(v) not in plain for v in fragments) or
                    integracion_normalized(item['municipio_documentado']) not in plain):
                raise ValueError('El texto primario no respalda la revisión explícita')
        elif kind in ('publicacion_primaria_no_concluyente', 'revision_visual_y_fuente_primaria'):
            scripts=re.findall(r'<script[^>]*type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',source.read_text(encoding='utf-8'),re.S)
            nodes=[json.loads(s) for s in scripts]
            matches=[n for n in nodes if isinstance(n,dict) and n.get('@type')=='SocialMediaPosting' and n.get('@id')==item['url']]
            if ((kind=='publicacion_primaria_no_concluyente' and item['decision']!='mantener_pendiente') or len(matches)!=1 or
                    matches[0].get('author',{}).get('name')!=item['autor_documentado'] or
                    any(integracion_normalized(v) not in integracion_normalized(matches[0].get('articleBody','')) for v in item['fragmentos_documentados'])):
                raise ValueError('Publicación no concluyente no identificada o admisión indebida')
            if kind=='revision_visual_y_fuente_primaria':
                # La identificación visual es una revisión explícita, no reconocimiento automático.
                entry=item['revision_visual']
                review_path=integracion_within(root,entry['ruta'])
                files.append(comprobar_archivo(root,review_path,entry))
                review=json.loads(review_path.read_text(encoding='utf-8'))
                if (review.get('instalacion_id')!=key or review.get('codigo_municipio')!=item['codigo_municipio'] or
                        review.get('alcance')!='municipio_emplazamiento' or not review.get('motivo') or
                        not review.get('url_visual') or not review.get('evidencias')):
                    raise ValueError('Revisión visual incompleta o para otra instalación/municipio')
                for evidence in review['evidencias']:
                    files.append(comprobar_archivo(root,integracion_within(root,evidence['ruta']),evidence))
        else:
            raise ValueError('Tipo de evidencia desconocido')
        rows=c.execute(f'''SELECT nombre_instalacion,direccion,codigo_municipio_espacial,
            codigo_municipio_declarado,numero_municipios_intersectados,estado_geometria,
            codigo_provincia_declarada FROM {integracion_table('recarga_espacial','recarga_instalaciones_geo')}
            WHERE instalacion_id=?''',[key]).fetchall()
        if len(rows)!=1 or list(rows[0][:2])!=item['atributos_esperados']:
            raise ValueError('La instalación revisada ha cambiado')
        row=rows[0]
        if (row[2]!=item['codigo_municipio'] or row[2]==row[3] or row[4]!=1 or
                row[5]!='admitida' or row[6]!=row[2][:2]):
            raise ValueError('Decisión incompatible con el contraste espacial')
        if kind in ('portal_cartociudad', 'via_cartociudad'):
            postal=c.execute(f"SELECT codigo_postal FROM {integracion_table('recarga_espacial','recarga_instalaciones_geo')} WHERE instalacion_id=?",[key]).fetchone()[0]
            if postal!=item['codigo_postal_origen']:
                raise ValueError('El código postal de origen ha cambiado')
        if geo is not None:
            candidates=c.execute(f'''SELECT codigo_municipio FROM {integracion_table('base_territorial','recintos_municipales')}
                WHERE ST_Intersects(geometria,ST_Transform(ST_Point(?,?),
                'EPSG:4326','EPSG:25830',always_xy:=true)) ORDER BY 1''',
                [float(geo['longitude']),float(geo['latitude'])]).fetchall()
            # Contraste diagnóstico de la ficha comercial: solo impide una admisión contradictoria.
            if item['decision']=='admitir' and candidates!=[(item['codigo_municipio'],)]:
                raise ValueError('La propia coordenada de la ficha contradice la admisión municipal')
        # La correspondencia exacta de establecimiento/dirección se revisó explícitamente en la decisión.
        # La coordenada del comercio no certifica el punto de recarga ni se usa como umbral de admisión.
        decisions[key] = item
    return decisions,files


def integracion_review_charging(c, decisions):
    g=integracion_table('recarga_espacial','recarga_instalaciones_geo')
    rows=c.execute(f'''SELECT instalacion_id,nombre_instalacion,direccion,municipio_nombre_origen,
        nombre_municipio_espacial,codigo_municipio_declarado,codigo_municipio_espacial,
        distancia_municipio_declarado_m FROM {g} WHERE NOT admitida_asignacion_municipal ORDER BY instalacion_id''').fetchall()
    c.execute('''CREATE TABLE revision_recarga(instalacion_id VARCHAR,nombre_instalacion VARCHAR,
        direccion VARCHAR,municipio_declarado VARCHAR,municipio_espacial VARCHAR,
        codigo_municipio_declarado VARCHAR,codigo_municipio_espacial VARCHAR,distancia_recinto_declarado_m DOUBLE,
        indicio_textual VARCHAR,estado_revision VARCHAR,codigo_municipio_validado VARCHAR,
        url_evidencia VARCHAR,motivo VARCHAR)''')
    for row in rows:
        key,name,address,declared,spatial,*_=row
        # Solo indicios: un nombre municipal puede ser una calle o parte de una marca.
        text=integracion_normalized(name+' '+address)
        def mentioned(label):
            return any(len(integracion_normalized(p))>=4 and (' '+integracion_normalized(p)+' ') in (' '+text+' ') for p in label.split('/'))
        a,b=mentioned(declared),mentioned(spatial)
        hint='ambos' if a and b else 'espacial' if b else 'declarado' if a else 'sin_mencion_literal'
        decision=decisions.get(key)
        admitted=decision is not None and decision['decision']=='admitir'
        c.execute('INSERT INTO revision_recarga VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)', [*row,hint,
            'resuelta_documentalmente' if admitted else 'pendiente_evidencia_documental_no_concluyente' if decision else 'pendiente_contraste_externo',
            decision['codigo_municipio'] if admitted else None,decision['url'] if decision else None,
            decision['motivo'] if decision else 'Revisados nombre, dirección y contraste; indicios internos insuficientes para decidir qué atributo falla.'])
    if set(decisions)-{r[0] for r in rows}:
        raise ValueError('Decisión para instalación que no es un conflicto vigente')
    c.execute(f'''CREATE TABLE recarga_municipio_validado AS SELECT instalacion_id,
        codigo_municipio_espacial codigo_municipio,'interseccion_unica_y_nombre_oficial'::VARCHAR metodo,
        NULL::VARCHAR url_evidencia FROM {g} WHERE admitida_asignacion_municipal
        UNION ALL SELECT instalacion_id,codigo_municipio_validado,'revision_documental_establecimiento',
        url_evidencia FROM revision_recarga WHERE estado_revision='resuelta_documentalmente' ''')
    for key, decision in decisions.items():
        if decision['decision']=='admitir' and decision.get('tipo_evidencia') in ('portal_cartociudad','via_cartociudad','revision_visual_y_fuente_primaria'):
            c.execute('UPDATE recarga_municipio_validado SET metodo=? WHERE instalacion_id=?',
                ['revision_visual_y_documental' if decision['tipo_evidencia']=='revision_visual_y_fuente_primaria' else 'revision_cartografica_'+decision['tipo_evidencia'],key])
    for name in ('revision_recarga','recarga_municipio_validado'):
        comprobar_clave(c,name,INTEGRACION_OUTPUT_KEYS[name])
    controls={}
    for name in RECARGA_KEYS:
        source=integracion_table('recarga',name)
        total,admitted=c.execute(f'''SELECT count(*),count(v.instalacion_id) FROM {source} s
            LEFT JOIN recarga_municipio_validado v USING(instalacion_id)''').fetchone()
        if total!=integracion_scalar(c,f'SELECT count(*) FROM {source}'):
            raise ValueError('Multiplicación de entidades de recarga')
        controls[name]={'conservadas':total,'asignacion_municipal_admitida':admitted,'pendiente':total-admitted}
    return {'estados_revision':consultar(c,'SELECT estado_revision,count(*) instalaciones FROM revision_recarga GROUP BY 1 ORDER BY 1'),
            'metodos_admision':consultar(c,'SELECT metodo,count(*) instalaciones FROM recarga_municipio_validado GROUP BY 1 ORDER BY 1'),
            'indicios_textuales':consultar(c,'SELECT indicio_textual,count(*) instalaciones FROM revision_recarga GROUP BY 1 ORDER BY 1'),
            'entidades':controls}


def integracion_spatial_relations(c, source, key_fields, label, kind):
    """Materializa el cruce antes de filtrar atributos (limitación spatial 1.5.0).

    Los contactos de dimensión inferior se conservan sin asignación exclusiva.
    Medidas de intersección solo diagnósticas, sin reparto de IMD ni otros valores.
    """
    comprobar_clave(c,source,key_fields)
    key='to_json(list_value('+','.join('cast('+identificador_sql(k)+' AS VARCHAR)' for k in key_fields)+'))::VARCHAR'
    geom='ST_Force2D(geometria)'
    c.execute(f'''CREATE OR REPLACE TEMP TABLE espacial_entrada AS SELECT {key} clave_entidad,
        {geom} geometria,estado_geometria FROM {identificador_sql(source)}''')
    c.execute('''CREATE OR REPLACE TEMP TABLE espacial_admitida AS SELECT * FROM espacial_entrada
        WHERE estado_geometria='admitida' AND geometria IS NOT NULL''')
    exigir_cero(c, '''SELECT count(*) FROM espacial_admitida WHERE ST_IsEmpty(geometria)
        OR NOT ST_IsValid(geometria)''', 'Geometría marcada admitida que no es válida')
    c.execute(f'''CREATE OR REPLACE TEMP TABLE cruce_municipal AS SELECT p.clave_entidad,r.codigo_municipio,
        ST_Intersection(p.geometria,r.geometria) interseccion FROM espacial_admitida p
        JOIN {integracion_table('base_territorial','recintos_municipales')} r ON ST_Intersects(p.geometria,r.geometria)''')
    c.execute('''INSERT INTO relaciones_espaciales SELECT ?,clave_entidad,codigo_municipio,
        'interseccion'::VARCHAR,ST_Dimension(interseccion)::SMALLINT,
        CASE WHEN ST_Dimension(interseccion)=1 THEN ST_Length(interseccion) END,
        CASE WHEN ST_Dimension(interseccion)=2 THEN ST_Area(interseccion) END
        FROM cruce_municipal''',[label])
    c.execute('''INSERT INTO estado_relaciones_espaciales SELECT ?,p.clave_entidad,p.estado_geometria,
        coalesce(n,0)::BIGINT,CASE WHEN p.estado_geometria<>'admitida' OR p.geometria IS NULL THEN 'geometria_no_admitida'
        WHEN n IS NULL THEN 'sin_interseccion_cv' WHEN n=1 THEN 'interseccion_unica'
        ELSE 'interseccion_multiple' END FROM espacial_entrada p LEFT JOIN
        (SELECT clave_entidad,count(*) n FROM cruce_municipal GROUP BY 1) x USING(clave_entidad)''',[label])
    result = consultar(c,f'''SELECT estado_relacion,count(*) entidades,sum(numero_municipios) relaciones
        FROM estado_relaciones_espaciales WHERE tabla_entidad={texto_sql(label)} GROUP BY 1 ORDER BY 1''')
    # Solo se necesitan las relaciones; liberar cada lote antes del siguiente.
    for name in ('cruce_municipal','espacial_admitida','espacial_entrada'):
        c.execute(f'DROP TABLE {identificador_sql(name)}')
    return result


def integracion_coverage(c):
    c.execute('''CREATE TABLE cobertura_estadistica(tabla_entidad VARCHAR,periodo VARCHAR,segmento VARCHAR,
        codigo_municipio VARCHAR,registros BIGINT,valores_informados BIGINT,valores_ausentes BIGINT)''')
    for name,period,segment,value in (
        ('parque_municipal','fecha_referencia',"'distintivo_0'",'vehiculos_distintivo_0'),
        ('poblacion_municipal','anio','sexo','poblacion'),
        ('renta_territorial','anio',"to_json(list_value(nivel_territorial,indicador))::VARCHAR",'valor_eur')):
        source=integracion_table('demografia_parque',name)
        c.execute(f'''INSERT INTO cobertura_estadistica SELECT ?,cast({period} AS VARCHAR),{segment},
            codigo_municipio,count(*),count({value}),count(*)-count({value}) FROM {source} GROUP BY 2,3,4''',[name])
    return consultar(c,'''SELECT tabla_entidad,periodo,segmento,count(*) municipios,sum(registros) registros,
        sum(valores_ausentes) valores_ausentes FROM cobertura_estadistica GROUP BY 1,2,3 ORDER BY 1,2,3''')


# %% Silver: versión principal de recarga y ausencias históricas

RECARGA_VERSIONADA_KEYS = {**RECARGA_KEYS, 'incidencias': None, **RECARGA_GEO_KEYS, 'incidencias_espaciales': None,
        'revision_recarga': INTEGRACION_OUTPUT_KEYS['revision_recarga'],
        'recarga_municipio_validado': INTEGRACION_OUTPUT_KEYS['recarga_municipio_validado'],
        'recarga_ausentes': ['instalacion_id'],
        'recarga_ausentes_municipios_candidatos': ['instalacion_id', 'codigo_municipio']}


def recarga_versionada_inherit_geometry(c):
    fields = [row[0] for row in c.execute('DESCRIBE recarga_instalaciones').fetchall()]
    different = ' OR '.join(f'n.{identificador_sql(f)} IS DISTINCT FROM a.{identificador_sql(f)}' for f in fields)
    exigir_cero(c, f'''SELECT count(*) FROM recarga_instalaciones n
        LEFT JOIN geometria_anterior a USING(instalacion_id)
        WHERE a.instalacion_id IS NULL OR {different}''',
        'Identidad o coordenadas distintas: no se puede heredar geometría ni admisión')
    c.execute('''CREATE TABLE recarga_instalaciones_geo AS SELECT a.*
        FROM geometria_anterior a JOIN recarga_instalaciones n USING(instalacion_id)''')
    # Se repite la relación espacial para detectar cambios de recintos o candidatos.
    c.execute('''CREATE TABLE recarga_municipios_candidatos AS
        SELECT p.instalacion_id,r.codigo_municipio FROM recarga_instalaciones_geo p
        JOIN recintos r ON ST_Intersects(p.geometria,r.geometria)''')
    exigir_cero(c, '''SELECT count(*) FROM (
        (SELECT * FROM recarga_municipios_candidatos EXCEPT ALL
         SELECT a.* FROM candidatos_anteriores a JOIN recarga_instalaciones n USING(instalacion_id))
        UNION ALL
        (SELECT a.* FROM candidatos_anteriores a JOIN recarga_instalaciones n USING(instalacion_id)
         EXCEPT ALL SELECT * FROM recarga_municipios_candidatos))''',
        'Los municipios intersectados difieren del contraste histórico')
    exigir_cero(c, '''SELECT count(*) FROM recarga_instalaciones_geo WHERE
        estado_geometria IS DISTINCT FROM 'admitida' OR geometria IS NULL OR
        ST_IsEmpty(geometria) OR NOT ST_IsValid(geometria) OR
        NOT isfinite(ST_X(geometria)) OR NOT isfinite(ST_Y(geometria))''',
        'Geometría histórica no admitida para reutilización')
    if c.execute('SELECT DISTINCT typeof(geometria) FROM recarga_instalaciones_geo').fetchall() != [("GEOMETRY('EPSG:25830')",)]:
        raise ValueError('CRS histórico no admitido')
    c.execute('''CREATE TABLE entidad_municipio AS SELECT 'recarga_instalaciones_geo'::VARCHAR tabla_entidad,
        instalacion_id clave_entidad,codigo_municipio_espacial codigo_municipio,
        'interseccion_unica_y_nombre_oficial'::VARCHAR metodo_asignacion,
        'contrastada'::VARCHAR estado_asignacion FROM recarga_instalaciones_geo
        WHERE admitida_asignacion_municipal''')
    c.execute('''CREATE TABLE incidencias_espaciales AS SELECT instalacion_id,
        'contraste_territorial'::VARCHAR campo,estado_contraste_territorial motivo
        FROM recarga_instalaciones_geo WHERE NOT admitida_asignacion_municipal''')


def recarga_versionada_preserve_absences(c):
    c.execute('''CREATE TABLE recarga_ausentes AS SELECT g.*,
        p.puntos_registrados_historicos,c.conectores_registrados_historicos,
        a.codigo_municipio codigo_municipio_admitido_historico,a.metodo metodo_admision_historica,
        'ausente_exportacion_sin_cierre_acreditado'::VARCHAR estado_presencia
        FROM geometria_anterior g
        JOIN (SELECT instalacion_id,count(*)::BIGINT puntos_registrados_historicos
              FROM anterior_recarga_puntos GROUP BY 1) p USING(instalacion_id)
        JOIN (SELECT instalacion_id,count(*)::BIGINT conectores_registrados_historicos
              FROM anterior_recarga_conectores GROUP BY 1) c USING(instalacion_id)
        LEFT JOIN admisiones_anteriores a USING(instalacion_id)
        ANTI JOIN recarga_instalaciones n USING(instalacion_id)''')
    c.execute('''CREATE TABLE recarga_ausentes_municipios_candidatos AS SELECT a.*
        FROM candidatos_anteriores a JOIN recarga_ausentes h USING(instalacion_id)''')
    for name in ('recarga_ausentes', 'recarga_ausentes_municipios_candidatos'):
        comprobar_clave(c, name, RECARGA_VERSIONADA_KEYS[name])
    exigir_cero(c, '''SELECT count(*) FROM recarga_ausentes h
        JOIN recarga_instalaciones n USING(instalacion_id)''', 'Ausente incluido en oferta principal')
    exigir_cero(c, '''SELECT count(*) FROM (
        (SELECT instalacion_id FROM geometria_anterior EXCEPT
         (SELECT instalacion_id FROM recarga_instalaciones UNION ALL SELECT instalacion_id FROM recarga_ausentes))
        UNION ALL
        ((SELECT instalacion_id FROM recarga_instalaciones UNION ALL SELECT instalacion_id FROM recarga_ausentes)
         EXCEPT SELECT instalacion_id FROM geometria_anterior))''',
        'No se conserva el universo de IDs históricos')


# %% Gold: contexto municipal e incertidumbre

GOLD_MUNICIPAL_KEYS = {
    'gold_municipios': ['version_gold', 'codigo_municipio'],
    'gold_recarga_instalaciones': ['version_gold', 'instalacion_id'],
    'gold_incertidumbre_recarga_municipio': ['version_gold', 'instalacion_id', 'codigo_municipio', 'tipo_incertidumbre']}


GOLD_MUNICIPAL_BLOCKS = {
    'base_territorial': ['municipios', 'recintos_municipales'],
    'demografia_parque': ['parque_municipal', 'poblacion_municipal', 'renta_territorial'],
    'recarga_versionada': ['recarga_instalaciones_geo', 'recarga_puntos', 'recarga_conectores',
                          'recarga_municipio_validado', 'recarga_municipios_candidatos',
                          'recarga_ausentes', 'recarga_ausentes_municipios_candidatos']}


def gold_municipal_select_inputs(c):
    c.execute("CREATE TEMP TABLE parque_elegido AS SELECT * FROM parque_municipal WHERE fecha_referencia=DATE '2025-12-31'")
    c.execute("CREATE TEMP TABLE poblacion_elegida AS SELECT * FROM poblacion_municipal WHERE sexo='Total' AND anio=2025")
    c.execute("""CREATE TEMP TABLE renta_elegida AS SELECT * FROM renta_territorial
        WHERE nivel_territorial='municipio' AND indicador='Renta neta media por persona' AND anio=2023""")
    for name in ('municipios', 'recintos_municipales', 'parque_elegido', 'poblacion_elegida', 'renta_elegida'):
        comprobar_clave(c, name, ['codigo_municipio'])
        exigir_cero(c, f'''SELECT count(*) FROM (
            (SELECT codigo_municipio FROM municipios EXCEPT SELECT codigo_municipio FROM {name}) UNION ALL
            (SELECT codigo_municipio FROM {name} EXCEPT SELECT codigo_municipio FROM municipios))''',
            f'Cobertura municipal incompleta o ajena a la referencia: {name}')
    for name, keys in (
        ('recarga_instalaciones_geo', ['instalacion_id']), ('recarga_ausentes', ['instalacion_id']),
        ('recarga_puntos', ['instalacion_id', 'punto_id']),
        ('recarga_conectores', ['instalacion_id', 'punto_id', 'conector_id']),
        ('recarga_municipio_validado', ['instalacion_id']),
        ('recarga_municipios_candidatos', ['instalacion_id', 'codigo_municipio']),
        ('recarga_ausentes_municipios_candidatos', ['instalacion_id', 'codigo_municipio'])):
        comprobar_clave(c, name, keys)
    for child, parent, columns in (
        ('recarga_puntos', 'recarga_instalaciones_geo', 'instalacion_id'),
        ('recarga_conectores', 'recarga_puntos', 'instalacion_id,punto_id')):
        exigir_cero(c, f'SELECT count(*) FROM {child} ANTI JOIN {parent} USING({columns})', 'Detalle de recarga huérfano')
        exigir_cero(c, f'SELECT count(*) FROM {parent} ANTI JOIN {child} USING({columns})', 'Instalación o punto sin detalle')
    for name, parent in (('recarga_municipio_validado', 'recarga_instalaciones_geo'),
                         ('recarga_municipios_candidatos', 'recarga_instalaciones_geo'),
                         ('recarga_ausentes_municipios_candidatos', 'recarga_ausentes')):
        exigir_cero(c, f'SELECT count(*) FROM {name} ANTI JOIN {parent} USING(instalacion_id)', 'Relación de recarga huérfana')
        exigir_cero(c, f'SELECT count(*) FROM {name} ANTI JOIN municipios USING(codigo_municipio)', 'Municipio de recarga no reconocido')
    exigir_cero(c, '''SELECT count(*) FROM recarga_ausentes a
        JOIN recarga_instalaciones_geo n USING(instalacion_id)''', 'Un ID figura como presente y ausente')
    exigir_cero(c, '''SELECT count(*) FROM recarga_ausentes a ANTI JOIN municipios m
        ON a.codigo_municipio_admitido_historico=m.codigo_municipio
        WHERE a.codigo_municipio_admitido_historico IS NOT NULL''', 'Admisión histórica fuera de la referencia')


def gold_municipal_prepare(c, version, acquisition, main_origin, historical_origin):
    gold_municipal_select_inputs(c)
    c.execute('''CREATE TABLE gold_recarga_instalaciones AS WITH
        puntos AS (SELECT instalacion_id,count(*)::BIGINT n FROM recarga_puntos GROUP BY 1),
        conectores AS (SELECT instalacion_id,count(*)::BIGINT n FROM recarga_conectores GROUP BY 1)
        SELECT ?::VARCHAR version_gold,g.instalacion_id,?::VARCHAR version_recarga_observada,
            true presente_copia_principal,g.nombre_instalacion,g.operador_id,g.geometria,g.estado_geometria,
            v.codigo_municipio codigo_municipio_admitido,
            CASE WHEN v.instalacion_id IS NULL THEN 'pendiente' ELSE 'admitida_version_principal' END estado_admision,
            v.metodo metodo_admision,p.n puntos_registrados,c.n conectores_registrados,
            ?::TIMESTAMPTZ fecha_adquisicion_observacion
        FROM recarga_instalaciones_geo g LEFT JOIN recarga_municipio_validado v USING(instalacion_id)
        LEFT JOIN puntos p USING(instalacion_id) LEFT JOIN conectores c USING(instalacion_id)
        UNION ALL SELECT ?::VARCHAR,a.instalacion_id,?::VARCHAR,false,a.nombre_instalacion,a.operador_id,
            a.geometria,a.estado_geometria,a.codigo_municipio_admitido_historico,
            CASE WHEN a.codigo_municipio_admitido_historico IS NULL THEN 'pendiente' ELSE 'admitida_solo_historica' END,
            a.metodo_admision_historica,a.puntos_registrados_historicos,a.conectores_registrados_historicos,
            NULL::TIMESTAMPTZ FROM recarga_ausentes a''',
        [version, main_origin, acquisition, version, historical_origin])
    c.execute('''CREATE TEMP TABLE avisos_recarga AS
        SELECT g.instalacion_id,'asignacion_pendiente'::VARCHAR tipo_incertidumbre,
            g.codigo_municipio_declarado,NULL::VARCHAR codigo_historico
        FROM recarga_instalaciones_geo g ANTI JOIN recarga_municipio_validado v USING(instalacion_id)
        UNION ALL SELECT instalacion_id,'ausencia_exportacion',codigo_municipio_declarado,
            codigo_municipio_admitido_historico FROM recarga_ausentes''')
    c.execute('''CREATE TABLE gold_incertidumbre_recarga_municipio AS WITH roles AS (
        SELECT instalacion_id,tipo_incertidumbre,codigo_historico codigo_municipio,
            true historico,false declarado,false espacial FROM avisos_recarga WHERE codigo_historico IS NOT NULL
        UNION ALL SELECT instalacion_id,tipo_incertidumbre,codigo_municipio_declarado,
            false,true,false FROM avisos_recarga WHERE codigo_municipio_declarado IS NOT NULL
        UNION ALL SELECT a.instalacion_id,a.tipo_incertidumbre,c.codigo_municipio,false,false,true
            FROM avisos_recarga a JOIN recarga_municipios_candidatos c USING(instalacion_id)
            WHERE tipo_incertidumbre='asignacion_pendiente'
        UNION ALL SELECT a.instalacion_id,a.tipo_incertidumbre,c.codigo_municipio,false,false,true
            FROM avisos_recarga a JOIN recarga_ausentes_municipios_candidatos c USING(instalacion_id)
            WHERE tipo_incertidumbre='ausencia_exportacion')
        SELECT ?::VARCHAR version_gold,r.instalacion_id,r.codigo_municipio,r.tipo_incertidumbre,
            bool_or(historico) es_municipio_admitido_historico,bool_or(declarado) es_municipio_declarado,
            bool_or(espacial) es_candidato_espacial FROM roles r JOIN municipios m USING(codigo_municipio)
            JOIN avisos_recarga a USING(instalacion_id,tipo_incertidumbre)
        WHERE a.codigo_historico IS NULL OR r.codigo_municipio=a.codigo_historico GROUP BY 2,3,4''', [version])
    c.execute('''CREATE TABLE gold_municipios AS WITH oferta AS (
        SELECT codigo_municipio_admitido codigo_municipio,count(*)::BIGINT instalaciones,
            sum(puntos_registrados)::BIGINT puntos,sum(conectores_registrados)::BIGINT conectores
        FROM gold_recarga_instalaciones WHERE presente_copia_principal AND estado_admision='admitida_version_principal'
        GROUP BY 1), incertidumbre AS (
        SELECT codigo_municipio,count(DISTINCT instalacion_id) FILTER(WHERE tipo_incertidumbre='asignacion_pendiente') pendientes,
            count(DISTINCT instalacion_id) FILTER(WHERE tipo_incertidumbre='ausencia_exportacion') ausentes
        FROM gold_incertidumbre_recarga_municipio GROUP BY 1)
        SELECT ?::VARCHAR version_gold,m.codigo_municipio,m.codigo_provincia,m.nombre_municipio,r.geometria,
            p.fecha_referencia fecha_referencia_parque,p.parque_total,p.vehiculos_distintivo_0,
            p.vehiculos_distintivo_0::DOUBLE/nullif(p.parque_total,0) proporcion_distintivo_0,
            b.anio anio_poblacion,b.poblacion poblacion_total,b.estado_valor estado_poblacion,
            i.anio anio_renta,i.valor_eur renta_neta_media_persona_eur,i.estado_valor estado_renta,
            ?::TIMESTAMPTZ fecha_adquisicion_recarga,
            coalesce(o.instalaciones,0)::BIGINT instalaciones_con_municipio_admitido,
            coalesce(o.puntos,0)::BIGINT puntos_con_municipio_admitido,
            coalesce(o.conectores,0)::BIGINT conectores_con_municipio_admitido,
            p.vehiculos_distintivo_0::DOUBLE/nullif(o.puntos,0) distintivo_0_por_punto_admitido,
            CASE WHEN p.vehiculos_distintivo_0 IS NULL THEN 'sin_numerador_informado'
                WHEN coalesce(o.puntos,0)=0 THEN 'sin_puntos_admitidos' ELSE 'calculable' END estado_presion_descriptiva,
            coalesce(u.pendientes,0)::BIGINT instalaciones_pendientes_relacionadas,
            coalesce(u.ausentes,0)::BIGINT instalaciones_ausentes_relacionadas,
            EXISTS(SELECT 1 FROM avisos_recarga a ANTI JOIN gold_incertidumbre_recarga_municipio v
                USING(instalacion_id,tipo_incertidumbre)) hay_incertidumbre_sin_localizar
        FROM municipios m LEFT JOIN recintos_municipales r USING(codigo_municipio)
        LEFT JOIN parque_elegido p USING(codigo_municipio) LEFT JOIN poblacion_elegida b USING(codigo_municipio)
        LEFT JOIN renta_elegida i USING(codigo_municipio) LEFT JOIN oferta o USING(codigo_municipio)
        LEFT JOIN incertidumbre u USING(codigo_municipio)''', [version, acquisition])


def gold_municipal_validate(c, version):
    nullable = {
        'gold_municipios': {'parque_total', 'vehiculos_distintivo_0', 'proporcion_distintivo_0',
                           'poblacion_total', 'renta_neta_media_persona_eur', 'distintivo_0_por_punto_admitido'},
        'gold_recarga_instalaciones': {'geometria', 'codigo_municipio_admitido', 'metodo_admision', 'fecha_adquisicion_observacion'},
        'gold_incertidumbre_recarga_municipio': set()}
    for name, keys in GOLD_MUNICIPAL_KEYS.items():
        comprobar_clave(c, name, keys)
        required = [r[0] for r in c.execute('DESCRIBE ' + name).fetchall() if r[0] not in nullable[name]]
        exigir_cero(c, 'SELECT count(*) FROM ' + name + ' WHERE ' +
                     ' OR '.join(identificador_sql(f) + ' IS NULL' for f in required), 'Campo obligatorio Gold nulo')
        if c.execute(f'SELECT count(*) FROM {name} WHERE version_gold IS DISTINCT FROM ?', [version]).fetchone()[0]:
            raise ValueError('Versión Gold incoherente')
    exigir_cero(c, '''SELECT count(*) FROM (
        (SELECT codigo_municipio FROM municipios EXCEPT SELECT codigo_municipio FROM gold_municipios) UNION ALL
        (SELECT codigo_municipio FROM gold_municipios EXCEPT SELECT codigo_municipio FROM municipios))''', 'Se altera la cobertura municipal')
    exigir_cero(c, '''SELECT count(*) FROM gold_municipios WHERE geometria IS NULL OR NOT ST_IsValid(geometria)
        OR parque_total<0 OR vehiculos_distintivo_0<0 OR vehiculos_distintivo_0>parque_total
        OR poblacion_total<0 OR instalaciones_con_municipio_admitido<0 OR puntos_con_municipio_admitido<0
        OR conectores_con_municipio_admitido<0 OR instalaciones_pendientes_relacionadas<0
        OR instalaciones_ausentes_relacionadas<0''', 'Valor municipal fuera del contrato')
    exigir_cero(c, '''SELECT count(*) FROM gold_municipios WHERE
        proporcion_distintivo_0 IS DISTINCT FROM vehiculos_distintivo_0::DOUBLE/nullif(parque_total,0)
        OR distintivo_0_por_punto_admitido IS DISTINCT FROM vehiculos_distintivo_0::DOUBLE/nullif(puntos_con_municipio_admitido,0)
        OR (distintivo_0_por_punto_admitido IS NOT NULL AND NOT isfinite(distintivo_0_por_punto_admitido))
        OR estado_presion_descriptiva IS DISTINCT FROM CASE WHEN vehiculos_distintivo_0 IS NULL
            THEN 'sin_numerador_informado' WHEN puntos_con_municipio_admitido=0 THEN 'sin_puntos_admitidos'
            ELSE 'calculable' END''', 'Cociente o estado de presión incorrecto')
    exigir_cero(c, '''SELECT count(*) FROM gold_municipios m JOIN parque_elegido p USING(codigo_municipio)
        JOIN poblacion_elegida b USING(codigo_municipio) JOIN renta_elegida r USING(codigo_municipio) WHERE
        m.parque_total IS DISTINCT FROM p.parque_total OR m.vehiculos_distintivo_0 IS DISTINCT FROM p.vehiculos_distintivo_0
        OR m.fecha_referencia_parque IS DISTINCT FROM p.fecha_referencia
        OR m.poblacion_total IS DISTINCT FROM b.poblacion OR m.estado_poblacion IS DISTINCT FROM b.estado_valor
        OR m.anio_poblacion IS DISTINCT FROM b.anio OR m.renta_neta_media_persona_eur IS DISTINCT FROM r.valor_eur
        OR m.estado_renta IS DISTINCT FROM r.estado_valor OR m.anio_renta IS DISTINCT FROM r.anio''', 'Se altera una observación demográfica')
    exigir_cero(c, '''SELECT count(*) FROM gold_recarga_instalaciones WHERE puntos_registrados IS NULL
        OR conectores_registrados IS NULL OR puntos_registrados<=0 OR conectores_registrados<puntos_registrados
        OR estado_admision NOT IN ('pendiente','admitida_version_principal','admitida_solo_historica')
        OR (estado_admision='pendiente' AND (codigo_municipio_admitido IS NOT NULL OR metodo_admision IS NOT NULL))
        OR (estado_admision<>'pendiente' AND (codigo_municipio_admitido IS NULL OR metodo_admision IS NULL))
        OR (presente_copia_principal AND (fecha_adquisicion_observacion IS NULL OR estado_admision='admitida_solo_historica'))
        OR (NOT presente_copia_principal AND (fecha_adquisicion_observacion IS NOT NULL OR estado_admision='admitida_version_principal'))''',
        'Presencia, recuento o admisión de recarga incoherente')
    exigir_cero(c, '''SELECT count(*) FROM gold_recarga_instalaciones g ANTI JOIN municipios m
        ON g.codigo_municipio_admitido=m.codigo_municipio WHERE g.codigo_municipio_admitido IS NOT NULL''', 'Admisión fuera de la referencia municipal')
    exigir_cero(c, '''SELECT count(*) FROM (
        (SELECT instalacion_id,presente_copia_principal FROM gold_recarga_instalaciones EXCEPT ALL
         (SELECT instalacion_id,true FROM recarga_instalaciones_geo UNION ALL SELECT instalacion_id,false FROM recarga_ausentes))
        UNION ALL
        ((SELECT instalacion_id,true FROM recarga_instalaciones_geo UNION ALL SELECT instalacion_id,false FROM recarga_ausentes)
         EXCEPT ALL SELECT instalacion_id,presente_copia_principal FROM gold_recarga_instalaciones))''', 'Se altera el universo presente o histórico')
    exigir_cero(c, '''SELECT count(*) FROM gold_incertidumbre_recarga_municipio u
        ANTI JOIN municipios m USING(codigo_municipio)''', 'Incertidumbre para municipio inexistente')
    exigir_cero(c, '''SELECT count(*) FROM gold_incertidumbre_recarga_municipio u
        LEFT JOIN gold_recarga_instalaciones g USING(instalacion_id) WHERE g.instalacion_id IS NULL
        OR NOT (es_municipio_admitido_historico OR es_municipio_declarado OR es_candidato_espacial)
        OR tipo_incertidumbre NOT IN ('asignacion_pendiente','ausencia_exportacion')
        OR (tipo_incertidumbre='asignacion_pendiente' AND (NOT presente_copia_principal OR estado_admision<>'pendiente'))
        OR (tipo_incertidumbre='ausencia_exportacion' AND presente_copia_principal)''', 'Vínculo de incertidumbre incoherente')
    exigir_cero(c, '''SELECT count(*) FROM gold_municipios m LEFT JOIN (
        SELECT codigo_municipio,count(DISTINCT instalacion_id) FILTER(WHERE tipo_incertidumbre='asignacion_pendiente') pendientes,
            count(DISTINCT instalacion_id) FILTER(WHERE tipo_incertidumbre='ausencia_exportacion') ausentes
        FROM gold_incertidumbre_recarga_municipio GROUP BY 1) u USING(codigo_municipio)
        WHERE m.instalaciones_pendientes_relacionadas<>coalesce(u.pendientes,0)
        OR m.instalaciones_ausentes_relacionadas<>coalesce(u.ausentes,0)
        OR m.hay_incertidumbre_sin_localizar IS DISTINCT FROM EXISTS(
            SELECT 1 FROM avisos_recarga ANTI JOIN gold_incertidumbre_recarga_municipio
            USING(instalacion_id,tipo_incertidumbre))''', 'Avisos municipales no reconciliados')
    exigir_cero(c, '''SELECT count(*) FROM gold_municipios m LEFT JOIN (
        SELECT v.codigo_municipio,count(DISTINCT g.instalacion_id) instalaciones,
            count(DISTINCT (p.instalacion_id,p.punto_id)) puntos,
            count(DISTINCT (c.instalacion_id,c.punto_id,c.conector_id)) conectores
        FROM recarga_instalaciones_geo g JOIN recarga_municipio_validado v USING(instalacion_id)
        JOIN recarga_puntos p USING(instalacion_id) JOIN recarga_conectores c USING(instalacion_id,punto_id)
        GROUP BY 1) o USING(codigo_municipio) WHERE
        m.instalaciones_con_municipio_admitido<>coalesce(o.instalaciones,0)
        OR m.puntos_con_municipio_admitido<>coalesce(o.puntos,0)
        OR m.conectores_con_municipio_admitido<>coalesce(o.conectores,0)''', 'Recuento de oferta no reconciliado con Silver')
    return {'tablas': {name: c.execute('SELECT count(*) FROM ' + name).fetchone()[0] for name in GOLD_MUNICIPAL_KEYS},
        'municipios_sin_puntos_admitidos': c.execute('SELECT count(*) FROM gold_municipios WHERE puntos_con_municipio_admitido=0').fetchone()[0],
        'renta_no_informada': c.execute('SELECT count(*) FROM gold_municipios WHERE renta_neta_media_persona_eur IS NULL').fetchone()[0],
        'oferta_principal': consultar(c, '''SELECT sum(instalaciones_con_municipio_admitido) instalaciones,
            sum(puntos_con_municipio_admitido) puntos,sum(conectores_con_municipio_admitido) conectores FROM gold_municipios''')[0],
        'estados_admision': consultar(c, 'SELECT presente_copia_principal,estado_admision,count(*) instalaciones FROM gold_recarga_instalaciones GROUP BY ALL ORDER BY 1,2'),
        'incertidumbre': consultar(c, '''SELECT tipo_incertidumbre,count(*) vinculos,count(DISTINCT instalacion_id) instalaciones,
            count(DISTINCT codigo_municipio) municipios FROM gold_incertidumbre_recarga_municipio GROUP BY 1 ORDER BY 1'''),
        'incertidumbre_sin_localizar': c.execute('''SELECT count(*) FROM avisos_recarga
            ANTI JOIN gold_incertidumbre_recarga_municipio USING(instalacion_id,tipo_incertidumbre)''').fetchone()[0]}


# %% Gold: candidatos y proximidad por escenarios

GOLD_CANDIDATOS_KEYS = {
    'gold_candidatos': ['version_gold', 'candidato_id'],
    'gold_candidato_municipio': ['version_gold', 'candidato_id', 'codigo_municipio'],
    'gold_correspondencias_candidatos': ['version_gold', 'candidato_punto_id', 'candidato_area_id'],
    'gold_proximidad_candidato_recarga': ['version_gold', 'candidato_id', 'escenario']}


GOLD_CANDIDATOS_SCENARIOS = ('principal_admitida', 'incluye_pendientes', 'incluye_ausentes_historicas')


GOLD_CANDIDATOS_SOURCES = {
    'silver/poi_osm': ['poi_elementos', 'poi_correspondencias'],
    'silver/base_territorial': ['recintos_municipales'],
    'silver/validacion_conjunta': ['relaciones_espaciales'],
    'gold/contexto_municipal': ['gold_municipios', 'gold_recarga_instalaciones']}


def gold_candidatos_prepare(c, version):
    for table, keys in (
        ('poi_elementos', ['origen_id', 'fid_origen']),
        ('poi_correspondencias', ['origen_punto', 'fid_punto', 'origen_area', 'fid_area']),
        ('recintos_municipales', ['codigo_municipio']), ('gold_municipios', ['codigo_municipio']),
        ('gold_recarga_instalaciones', ['instalacion_id']),
        ('relaciones_espaciales', ['tabla_entidad', 'clave_entidad', 'codigo_municipio'])):
        comprobar_clave(c, table, keys)
    exigir_cero(c, '''SELECT count(*) FROM ((SELECT codigo_municipio FROM recintos_municipales EXCEPT
        SELECT codigo_municipio FROM gold_municipios) UNION ALL (SELECT codigo_municipio FROM gold_municipios
        EXCEPT SELECT codigo_municipio FROM recintos_municipales))''', 'Bases municipales diferentes')
    c.execute('''CREATE TEMP TABLE seleccion AS SELECT *,
        to_json(list_value(origen_id,cast(fid_origen AS VARCHAR)))::VARCHAR candidato_id
        FROM poi_elementos WHERE interes_para_revision AND grupo_uso IN ('aparcamiento','comercio','alojamiento')''')
    c.execute('CREATE TEMP TABLE reglas(categoria VARCHAR,codigo VARCHAR,grupo VARCHAR)')
    c.executemany('INSERT INTO reglas VALUES (?,?,?)', [(k, *v) for k, v in OSM_CATEGORIES.items()])
    exigir_cero(c, '''SELECT count(*) FROM seleccion s LEFT JOIN reglas r ON s.categoria_origen=r.categoria
        WHERE r.categoria IS NULL OR s.categoria_codigo_origen IS DISTINCT FROM r.codigo
        OR s.grupo_uso IS DISTINCT FROM r.grupo OR estado_geometria<>'admitida'
        OR estado_categoria<>'seleccionada' OR estado_ambito<>'intersecta_cv'
        OR geometria IS NULL OR NOT ST_IsValid(geometria) OR ST_IsEmpty(geometria)
        OR (representacion='punto' AND ST_GeometryType(geometria)<>'POINT')
        OR (representacion='area' AND ST_Dimension(geometria)<>2)
        OR representacion NOT IN ('punto','area')''', 'Preselección OSM incoherente')
    for name in ('seleccion', 'recintos_municipales', 'gold_recarga_instalaciones'):
        schema = {r[0]:r[1] for r in c.execute('DESCRIBE ' + name).fetchall()}
        if schema['geometria'] != "GEOMETRY('EPSG:25830')":
            raise ValueError('CRS de entrada no admitido')
    c.execute('CREATE TEMP TABLE ambito_cv AS SELECT ST_Union_Agg(geometria) geometria FROM recintos_municipales')
    print('Preparando puntos interiores y conservando áreas originales...', flush=True)
    c.execute('''CREATE TEMP TABLE representaciones AS SELECT s.*,
        CASE WHEN representacion='punto' THEN s.geometria ELSE ST_Intersection(s.geometria,a.geometria) END area_analisis
        FROM seleccion s CROSS JOIN ambito_cv a''')
    c.execute('''CREATE TABLE gold_candidato_municipio AS SELECT ?::VARCHAR version_gold,s.candidato_id,
        r.codigo_municipio,r.dimension_interseccion,r.superficie_interseccion_m2
        FROM seleccion s JOIN relaciones_espaciales r ON s.candidato_id=r.clave_entidad
        WHERE r.tabla_entidad='poi_elementos' ''', [version])
    c.execute('''CREATE TABLE gold_correspondencias_candidatos AS SELECT ?::VARCHAR version_gold,
        p.candidato_id candidato_punto_id,a.candidato_id candidato_area_id,
        r.punto_intersecta_area,r.coincide_id_publicado,r.coincide_nombre_normalizado,r.distancia_plana_m,r.estado
        FROM poi_correspondencias r JOIN seleccion p ON p.origen_id=r.origen_punto AND p.fid_origen=r.fid_punto
        JOIN seleccion a ON a.origen_id=r.origen_area AND a.fid_origen=r.fid_area''', [version])
    c.execute('''CREATE TABLE gold_candidatos AS WITH municipios AS (
        SELECT candidato_id,count(*)::BIGINT n,min(codigo_municipio) codigo FROM gold_candidato_municipio GROUP BY 1),
        pares AS (SELECT candidato_id,count(*)::BIGINT n FROM (
            SELECT candidato_punto_id candidato_id FROM gold_correspondencias_candidatos
            UNION ALL SELECT candidato_area_id FROM gold_correspondencias_candidatos) GROUP BY 1)
        SELECT ?::VARCHAR version_gold,s.candidato_id,s.origen_id,s.capa_origen,s.fid_origen,s.osm_id,s.nombre_origen,
        s.categoria_origen,s.grupo_uso,s.representacion,
        CASE WHEN s.grupo_uso='aparcamiento' THEN 'aparcamiento_mapeado'
            ELSE 'destino_sin_aparcamiento_acreditado' END tipo_preseleccion,
        'pendiente_comprobacion_local'::VARCHAR estado_revision,s.geometria,
        CASE WHEN s.representacion='punto' THEN s.geometria
            WHEN ST_Dimension(s.area_analisis)=2 THEN ST_PointOnSurface(s.area_analisis) ELSE NULL END punto_representativo,
        CASE WHEN s.representacion='punto' THEN 'punto_original'
            WHEN ST_Dimension(s.area_analisis)=2 THEN 'punto_interior_interseccion_cv'
            ELSE 'sin_superficie_en_cv' END metodo_representacion,
        CASE WHEN m.n=1 THEN m.codigo ELSE NULL END codigo_municipio_contexto,
        coalesce(m.n,0)::BIGINT numero_municipios_relacionados,
        CASE WHEN m.n=1 THEN 'interseccion_unica' ELSE 'interseccion_multiple' END estado_relacion_municipal,
        coalesce(p.n,0)::BIGINT numero_correspondencias FROM representaciones s
        LEFT JOIN municipios m USING(candidato_id) LEFT JOIN pares p USING(candidato_id)''', [version])
    c.execute('''CREATE TEMP TABLE oferta_escenarios AS SELECT s.escenario,r.* FROM gold_recarga_instalaciones r
        CROSS JOIN (VALUES ('principal_admitida'),('incluye_pendientes'),('incluye_ausentes_historicas')) s(escenario)
        WHERE s.escenario='incluye_ausentes_historicas' OR
            (r.presente_copia_principal AND (s.escenario='incluye_pendientes' OR r.estado_admision='admitida_version_principal'))''')
    c.execute('''CREATE TEMP TABLE oferta_admisible AS SELECT * FROM oferta_escenarios
        WHERE estado_geometria='admitida' AND geometria IS NOT NULL AND ST_IsValid(geometria) AND NOT ST_IsEmpty(geometria)''')


def gold_candidatos_proximity(c, version):
    c.execute('''CREATE TABLE gold_proximidad_candidato_recarga(version_gold VARCHAR,candidato_id VARCHAR,
        escenario VARCHAR,instalacion_id_mas_cercana VARCHAR,distancia_plana_m DOUBLE,
        numero_instalaciones_a_distancia_minima BIGINT,estado_distancia VARCHAR)''')
    c.execute('''CREATE TEMP TABLE posiciones AS SELECT row_number() OVER(ORDER BY candidato_id) indice,
        candidato_id,punto_representativo FROM gold_candidatos''')
    total = c.execute('SELECT count(*) FROM posiciones').fetchone()[0]
    # Acotar el producto de distancias evita retener millones de claves largas en memoria.
    for start in range(1, total + 1, 512):
        c.execute('CREATE OR REPLACE TEMP TABLE lote AS SELECT * FROM posiciones WHERE indice>=? AND indice<?', [start, start+512])
        c.execute('''CREATE OR REPLACE TEMP TABLE distancias AS SELECT p.indice,r.instalacion_id,
            ST_Distance(p.punto_representativo,r.geometria) d FROM lote p CROSS JOIN
            (SELECT DISTINCT instalacion_id,geometria FROM oferta_admisible) r
            WHERE p.punto_representativo IS NOT NULL''')
        for scenario in GOLD_CANDIDATOS_SCENARIOS:
            c.execute('''INSERT INTO gold_proximidad_candidato_recarga WITH
                admitidas AS (SELECT d.* FROM distancias d JOIN oferta_admisible o USING(instalacion_id) WHERE o.escenario=?),
                minimos AS (SELECT indice,min(d) d FROM admitidas GROUP BY 1),
                vecinos AS (SELECT a.indice,m.d,min(a.instalacion_id) id,count(*)::BIGINT n
                    FROM admitidas a JOIN minimos m ON a.indice=m.indice AND a.d=m.d GROUP BY 1,2)
                SELECT ?::VARCHAR,p.candidato_id,?::VARCHAR,v.id,v.d,coalesce(v.n,0)::BIGINT,
                    CASE WHEN p.punto_representativo IS NULL THEN 'sin_posicion_representativa'
                         WHEN v.id IS NULL THEN 'sin_recarga_admisible' ELSE 'calculable' END
                    FROM lote p LEFT JOIN vecinos v USING(indice)''', [scenario, version, scenario])
        if start == 1 or start % 4096 == 1:
            print(f'Proximidad: {min(start+511,total)}/{total} elementos...', flush=True)


def gold_candidatos_validate(c, version):
    nullable = {'osm_id','nombre_origen','punto_representativo','codigo_municipio_contexto',
                'superficie_interseccion_m2','instalacion_id_mas_cercana','distancia_plana_m'}
    for name, keys in GOLD_CANDIDATOS_KEYS.items():
        comprobar_clave(c, name, keys)
        columns = [r[0] for r in c.execute('DESCRIBE ' + name).fetchall()]
        required = [x for x in columns if x not in nullable or (name=='gold_correspondencias_candidatos' and x=='distancia_plana_m')]
        exigir_cero(c, f'SELECT count(*) FROM {name} WHERE ' + ' OR '.join(identificador_sql(x)+' IS NULL' for x in required), 'Campo obligatorio nulo')
        if c.execute(f'SELECT count(*) FROM {name} WHERE version_gold<>?', [version]).fetchone()[0]:
            raise ValueError('Versión de salida incorrecta')
    exigir_cero(c, '''SELECT count(*) FROM ((SELECT candidato_id FROM seleccion EXCEPT SELECT candidato_id FROM gold_candidatos)
        UNION ALL (SELECT candidato_id FROM gold_candidatos EXCEPT SELECT candidato_id FROM seleccion))''', 'Selección de candidatos alterada')
    source_fields = ('origen_id','capa_origen','fid_origen','osm_id','nombre_origen','categoria_origen','grupo_uso','representacion')
    exigir_cero(c, '''SELECT count(*) FROM gold_candidatos c JOIN seleccion s USING(candidato_id) WHERE ''' +
        ' OR '.join(f'c.{x} IS DISTINCT FROM s.{x}' for x in source_fields) + '''
        OR c.estado_revision<>'pendiente_comprobacion_local'
        OR c.tipo_preseleccion<>CASE WHEN s.grupo_uso='aparcamiento' THEN 'aparcamiento_mapeado'
            ELSE 'destino_sin_aparcamiento_acreditado' END''', 'Atributos OSM alterados')
    for name, expected in (
        ('gold_candidato_municipio', '''SELECT s.candidato_id,r.codigo_municipio,r.dimension_interseccion,r.superficie_interseccion_m2
            FROM seleccion s JOIN relaciones_espaciales r ON s.candidato_id=r.clave_entidad WHERE r.tabla_entidad='poi_elementos' '''),
        ('gold_correspondencias_candidatos', '''SELECT p.candidato_id candidato_punto_id,a.candidato_id candidato_area_id,
            r.punto_intersecta_area,r.coincide_id_publicado,r.coincide_nombre_normalizado,r.distancia_plana_m,r.estado
            FROM poi_correspondencias r JOIN seleccion p ON p.origen_id=r.origen_punto AND p.fid_origen=r.fid_punto
            JOIN seleccion a ON a.origen_id=r.origen_area AND a.fid_origen=r.fid_area''')):
        fields = ','.join(identificador_sql(r[0]) for r in c.execute('DESCRIBE ' + name).fetchall() if r[0]!='version_gold')
        exigir_cero(c, f'''SELECT count(*) FROM (({expected} EXCEPT ALL SELECT {fields} FROM {name})
            UNION ALL (SELECT {fields} FROM {name} EXCEPT ALL {expected}))''', 'Relaciones Silver alteradas')
    for name, field in (('gold_candidato_municipio','candidato_id'), ('gold_correspondencias_candidatos','candidato_punto_id'),
                        ('gold_correspondencias_candidatos','candidato_area_id'), ('gold_proximidad_candidato_recarga','candidato_id')):
        exigir_cero(c, f'SELECT count(*) FROM {name} x ANTI JOIN gold_candidatos c ON x.{field}=c.candidato_id', 'Candidato huérfano')
    exigir_cero(c, '''SELECT count(*) FROM gold_candidato_municipio ANTI JOIN gold_municipios USING(codigo_municipio)''', 'Municipio huérfano')
    exigir_cero(c, '''SELECT count(*) FROM gold_candidatos c LEFT JOIN (
        SELECT candidato_id,count(*) n,min(codigo_municipio) codigo FROM gold_candidato_municipio GROUP BY 1) m USING(candidato_id)
        WHERE m.n IS NULL OR c.numero_municipios_relacionados<>m.n
        OR c.codigo_municipio_contexto IS DISTINCT FROM CASE WHEN m.n=1 THEN m.codigo END
        OR c.estado_relacion_municipal<>CASE WHEN m.n=1 THEN 'interseccion_unica' ELSE 'interseccion_multiple' END''', 'Contexto municipal incoherente')
    exigir_cero(c, '''SELECT count(*) FROM gold_candidatos c LEFT JOIN (
        SELECT candidato_id,count(*) n FROM (SELECT candidato_punto_id candidato_id FROM gold_correspondencias_candidatos
        UNION ALL SELECT candidato_area_id FROM gold_correspondencias_candidatos) GROUP BY 1) p USING(candidato_id)
        WHERE c.numero_correspondencias<>coalesce(p.n,0)''', 'Recuento de correspondencias incoherente')
    exigir_cero(c, '''SELECT count(*) FROM gold_correspondencias_candidatos r
        JOIN gold_candidatos p ON p.candidato_id=r.candidato_punto_id JOIN gold_candidatos a ON a.candidato_id=r.candidato_area_id
        WHERE p.representacion<>'punto' OR a.representacion<>'area' OR p.categoria_origen<>a.categoria_origen
        OR r.estado<>'posible_correspondencia_sin_fusion' OR NOT isfinite(r.distancia_plana_m) OR r.distancia_plana_m<0''', 'Correspondencia incoherente')
    exigir_cero(c, '''SELECT count(*) FROM gold_candidatos c JOIN representaciones s USING(candidato_id)
        WHERE ST_AsWKB(c.geometria)<>ST_AsWKB(s.geometria)
        OR (c.punto_representativo IS NULL)<>(c.metodo_representacion='sin_superficie_en_cv')
        OR (c.punto_representativo IS NOT NULL AND (NOT ST_IsValid(c.punto_representativo)
            OR ST_IsEmpty(c.punto_representativo) OR ST_GeometryType(c.punto_representativo)<>'POINT'
            OR NOT ST_Intersects(c.punto_representativo,s.area_analisis)
            OR NOT ST_Intersects(c.punto_representativo,c.geometria)))''', 'Representación geométrica incoherente')
    exigir_cero(c, '''SELECT count(*) FROM (SELECT c.candidato_id FROM gold_candidatos c
        LEFT JOIN gold_proximidad_candidato_recarga p USING(candidato_id)
        GROUP BY c.candidato_id HAVING count(p.escenario)<>3 OR count(DISTINCT p.escenario)<>3)''', 'Deben existir tres escenarios por candidato')
    exigir_cero(c, '''SELECT count(*) FROM gold_proximidad_candidato_recarga p JOIN gold_candidatos c USING(candidato_id)
        LEFT JOIN oferta_admisible r ON p.instalacion_id_mas_cercana=r.instalacion_id AND p.escenario=r.escenario
        WHERE p.escenario NOT IN ('principal_admitida','incluye_pendientes','incluye_ausentes_historicas')
        OR p.estado_distancia<>CASE WHEN c.punto_representativo IS NULL THEN 'sin_posicion_representativa'
            WHEN NOT EXISTS(SELECT 1 FROM oferta_admisible o WHERE o.escenario=p.escenario) THEN 'sin_recarga_admisible' ELSE 'calculable' END
        OR (p.estado_distancia='calculable' AND (r.instalacion_id IS NULL OR p.distancia_plana_m IS NULL
            OR NOT isfinite(p.distancia_plana_m) OR p.distancia_plana_m<0 OR p.numero_instalaciones_a_distancia_minima<1
            OR p.distancia_plana_m<>ST_Distance(c.punto_representativo,r.geometria)))
        OR (p.estado_distancia<>'calculable' AND (p.distancia_plana_m IS NOT NULL
            OR p.instalacion_id_mas_cercana IS NOT NULL OR p.numero_instalaciones_a_distancia_minima<>0))''', 'Distancia, vecino o estado incoherente')
    exigir_cero(c, '''SELECT count(*) FROM gold_proximidad_candidato_recarga a JOIN gold_proximidad_candidato_recarga b USING(candidato_id)
        WHERE ((a.escenario='principal_admitida' AND b.escenario='incluye_pendientes') OR
            (a.escenario='incluye_pendientes' AND b.escenario='incluye_ausentes_historicas')) AND a.distancia_plana_m<b.distancia_plana_m''', 'Escenarios no monótonos')
    # Volver a contrastar cada mínimo y todos sus empates contra la oferta completa, por lotes.
    total = c.execute('SELECT count(*) FROM posiciones').fetchone()[0]
    for start in range(1,total+1,512):
        bad = c.execute('''WITH contraste AS (
            SELECT p.candidato_id,g.escenario,g.distancia_plana_m,g.instalacion_id_mas_cercana,
                g.numero_instalaciones_a_distancia_minima,r.instalacion_id,
                ST_Distance(p.punto_representativo,r.geometria) d
            FROM posiciones p JOIN gold_proximidad_candidato_recarga g USING(candidato_id)
            JOIN oferta_admisible r USING(escenario) WHERE p.indice>=? AND p.indice<? AND g.estado_distancia='calculable')
            SELECT count(*) FROM (SELECT candidato_id,escenario FROM contraste
            GROUP BY candidato_id,escenario,distancia_plana_m,instalacion_id_mas_cercana,numero_instalaciones_a_distancia_minima
            HAVING min(d)<>distancia_plana_m
                OR min(instalacion_id) FILTER(WHERE d=distancia_plana_m)<>instalacion_id_mas_cercana
                OR count(*) FILTER(WHERE d=distancia_plana_m)<>numero_instalaciones_a_distancia_minima)''',[start,start+512]).fetchone()[0]
        if bad:
            raise ValueError('Mínimo, desempate o número de empates incorrectos')
    return {'tablas': {n:c.execute('SELECT count(*) FROM '+n).fetchone()[0] for n in GOLD_CANDIDATOS_KEYS},
        'categorias': consultar(c,'SELECT grupo_uso,representacion,count(*) elementos FROM gold_candidatos GROUP BY 1,2 ORDER BY 1,2'),
        'municipios_con_elementos': c.execute('SELECT count(DISTINCT codigo_municipio) FROM gold_candidato_municipio').fetchone()[0],
        'relacion_municipal': consultar(c,'SELECT numero_municipios_relacionados,count(*) elementos FROM gold_candidatos GROUP BY 1 ORDER BY 1'),
        'representacion': consultar(c,'SELECT metodo_representacion,count(*) elementos FROM gold_candidatos GROUP BY 1 ORDER BY 1'),
        'oferta': consultar(c,'''SELECT e.escenario,count(*) instalaciones,count(a.instalacion_id) admisibles,
            count(*)-count(a.instalacion_id) excluidas FROM oferta_escenarios e LEFT JOIN oferta_admisible a USING(escenario,instalacion_id) GROUP BY 1 ORDER BY 1'''),
        'proximidad': consultar(c,'''SELECT escenario,estado_distancia,count(*) elementos,min(distancia_plana_m) minimo_m,
            median(distancia_plana_m) mediana_m,max(distancia_plana_m) maximo_m FROM gold_proximidad_candidato_recarga GROUP BY 1,2 ORDER BY 1,2''')}


# %% Gold: controles espaciales

GOLD_CONTROLES_KEYS = {
    'gold_control_area_recarga': ['version_gold','candidato_id','escenario'],
    'gold_solapamientos_candidatos': ['version_gold','candidato_a_id','candidato_b_id'],
    'gold_control_borde': ['version_gold','candidato_id','escenario']}


def gold_controles_prepare(c, version):
    for name, keys in (
        ('gold_candidatos',['candidato_id']),('gold_municipios',['codigo_municipio']),
        ('gold_recarga_instalaciones',['instalacion_id']),
        ('gold_proximidad_candidato_recarga',['candidato_id','escenario'])):
        comprobar_clave(c,name,keys)
    for name, fields in (('gold_candidatos',('geometria','punto_representativo')),
                         ('gold_municipios',('geometria',)),('gold_recarga_instalaciones',('geometria',))):
        schema = {r[0]:r[1] for r in c.execute('DESCRIBE '+name).fetchall()}
        if any(schema[f]!="GEOMETRY('EPSG:25830')" for f in fields):
            raise ValueError('CRS de entrada no admitido')
    exigir_cero(c, '''SELECT count(*) FROM gold_candidatos WHERE geometria IS NULL
        OR NOT ST_IsValid(geometria) OR ST_IsEmpty(geometria) OR representacion NOT IN ('punto','area')
        OR (representacion='area' AND ST_Dimension(geometria)<>2)
        OR (representacion='punto' AND ST_GeometryType(geometria)<>'POINT')
        OR (punto_representativo IS NOT NULL AND (NOT ST_IsValid(punto_representativo)
            OR ST_IsEmpty(punto_representativo) OR ST_GeometryType(punto_representativo)<>'POINT'))''', 'Geometría candidata incoherente')
    exigir_cero(c, '''SELECT count(*) FROM gold_municipios WHERE geometria IS NULL
        OR NOT ST_IsValid(geometria) OR ST_IsEmpty(geometria) OR ST_Dimension(geometria)<>2''', 'Recinto no admitido')
    exigir_cero(c, '''SELECT count(*) FROM gold_proximidad_candidato_recarga
        ANTI JOIN gold_candidatos USING(candidato_id)''', 'Proximidad huérfana')
    c.execute('CREATE TEMP TABLE escenarios(escenario VARCHAR)')
    c.executemany('INSERT INTO escenarios VALUES (?)',[(x,) for x in GOLD_CANDIDATOS_SCENARIOS])
    exigir_cero(c, '''SELECT count(*) FROM ((SELECT candidato_id,escenario FROM gold_candidatos CROSS JOIN escenarios
        EXCEPT SELECT candidato_id,escenario FROM gold_proximidad_candidato_recarga) UNION ALL
        (SELECT candidato_id,escenario FROM gold_proximidad_candidato_recarga EXCEPT
        SELECT candidato_id,escenario FROM gold_candidatos CROSS JOIN escenarios))''', 'Escenarios de proximidad incompletos')
    exigir_cero(c, '''SELECT count(*) FROM gold_proximidad_candidato_recarga p JOIN gold_candidatos c USING(candidato_id)
        WHERE estado_distancia IS NULL OR estado_distancia NOT IN ('calculable','sin_posicion_representativa','sin_recarga_admisible')
        OR (estado_distancia='calculable' AND (distancia_plana_m IS NULL OR NOT isfinite(distancia_plana_m)
            OR distancia_plana_m<0 OR c.punto_representativo IS NULL))
        OR (estado_distancia<>'calculable' AND distancia_plana_m IS NOT NULL)
        OR (estado_distancia='sin_posicion_representativa')<>(c.punto_representativo IS NULL)''', 'Estado de proximidad incoherente')
    c.execute('''CREATE TEMP TABLE oferta AS SELECT e.escenario,r.* FROM gold_recarga_instalaciones r CROSS JOIN escenarios e
        WHERE (e.escenario='incluye_ausentes_historicas' OR (r.presente_copia_principal AND
            (e.escenario='incluye_pendientes' OR r.estado_admision='admitida_version_principal')))
        AND r.estado_geometria='admitida' AND r.geometria IS NOT NULL
        AND ST_IsValid(r.geometria) AND NOT ST_IsEmpty(r.geometria)''')
    c.execute("CREATE TEMP TABLE areas AS SELECT * FROM gold_candidatos WHERE representacion='area'")
    print('Contrastando áreas con instalaciones por escenario...',flush=True)
    c.execute('''CREATE TEMP TABLE coincidencias AS SELECT a.candidato_id,o.escenario,o.instalacion_id
        FROM areas a JOIN oferta o ON ST_Intersects(a.geometria,o.geometria)''')
    c.execute('''CREATE TABLE gold_control_area_recarga AS SELECT ?::VARCHAR version_gold,a.candidato_id,e.escenario,
        count(x.instalacion_id)::BIGINT instalaciones_intersectan,
        coalesce(list_sort(list(x.instalacion_id) FILTER(WHERE x.instalacion_id IS NOT NULL)),[]::VARCHAR[]) instalaciones_ids
        FROM areas a CROSS JOIN escenarios e LEFT JOIN coincidencias x USING(candidato_id,escenario)
        GROUP BY a.candidato_id,e.escenario''',[version])
    print('Conservando intersecciones entre candidatos, sin fusión...',flush=True)
    # Materializar antes de filtrar claves evita mezclar el cruce espacial con su ordenación.
    c.execute('''CREATE TEMP TABLE intersecciones AS SELECT a.candidato_id candidato_a_id,b.candidato_id candidato_b_id,
        ST_Intersection(a.geometria,b.geometria) interseccion,ST_Covers(a.geometria,b.geometria) a_cubre_b,
        ST_Covers(b.geometria,a.geometria) b_cubre_a FROM gold_candidatos a
        JOIN gold_candidatos b ON ST_Intersects(a.geometria,b.geometria)''')
    c.execute('''CREATE TABLE gold_solapamientos_candidatos AS SELECT ?::VARCHAR version_gold,
        x.candidato_a_id,x.candidato_b_id,ST_Dimension(x.interseccion)::SMALLINT dimension_interseccion,
        CASE WHEN ST_Dimension(x.interseccion)=2 THEN ST_Area(x.interseccion) END superficie_interseccion_m2,
        x.a_cubre_b,x.b_cubre_a,a.categoria_origen=b.categoria_origen misma_categoria,
        a.representacion=b.representacion misma_representacion FROM intersecciones x
        JOIN gold_candidatos a ON a.candidato_id=x.candidato_a_id JOIN gold_candidatos b ON b.candidato_id=x.candidato_b_id
        WHERE x.candidato_a_id<x.candidato_b_id''',[version])
    print('Calculando distancia al borde de la referencia CV...',flush=True)
    c.execute('CREATE TEMP TABLE ambito_cv AS SELECT ST_Union_Agg(geometria) geometria FROM gold_municipios')
    c.execute('''CREATE TEMP TABLE distancias_borde AS SELECT c.candidato_id,
        CASE WHEN c.punto_representativo IS NOT NULL THEN ST_Distance(c.punto_representativo,ST_Boundary(a.geometria)) END distancia_borde_cv_m
        FROM gold_candidatos c CROSS JOIN ambito_cv a''')
    c.execute('''CREATE TABLE gold_control_borde AS SELECT ?::VARCHAR version_gold,p.candidato_id,p.escenario,d.distancia_borde_cv_m,
        CASE WHEN p.estado_distancia='calculable' THEN p.distancia_plana_m>d.distancia_borde_cv_m ELSE NULL END posible_influencia_exterior,
        p.estado_distancia estado_control FROM gold_proximidad_candidato_recarga p JOIN distancias_borde d USING(candidato_id)''',[version])


def gold_controles_validate(c, version):
    optional = {'superficie_interseccion_m2','distancia_borde_cv_m','posible_influencia_exterior'}
    for name,keys in GOLD_CONTROLES_KEYS.items():
        comprobar_clave(c,name,keys)
        required = [r[0] for r in c.execute('DESCRIBE '+name).fetchall() if r[0] not in optional]
        exigir_cero(c,f'SELECT count(*) FROM {name} WHERE '+' OR '.join(identificador_sql(x)+' IS NULL' for x in required), 'Campo de control obligatorio nulo')
        if c.execute(f'SELECT count(*) FROM {name} WHERE version_gold<>?',[version]).fetchone()[0]:
            raise ValueError('Versión de control incorrecta')
    exigir_cero(c, '''SELECT count(*) FROM ((SELECT candidato_id,escenario FROM areas CROSS JOIN escenarios EXCEPT
        SELECT candidato_id,escenario FROM gold_control_area_recarga) UNION ALL
        (SELECT candidato_id,escenario FROM gold_control_area_recarga EXCEPT SELECT candidato_id,escenario FROM areas CROSS JOIN escenarios))''', 'Cobertura de áreas incompleta')
    exigir_cero(c, '''SELECT count(*) FROM gold_control_area_recarga r WHERE
        r.instalaciones_intersectan<>len(r.instalaciones_ids)
        OR r.instalaciones_ids<>list_sort(list_distinct(r.instalaciones_ids))''', 'Coincidencias de área no reconciliadas')
    c.execute('''CREATE OR REPLACE TEMP TABLE coincidencias_recalculadas AS
        SELECT a.candidato_id,o.escenario,o.instalacion_id FROM areas a
        JOIN oferta o ON ST_Intersects(a.geometria,o.geometria)''')
    exigir_cero(c, '''SELECT count(*) FROM ((SELECT candidato_id,escenario,unnest(instalaciones_ids) instalacion_id
        FROM gold_control_area_recarga EXCEPT ALL SELECT * FROM coincidencias_recalculadas) UNION ALL
        (SELECT * FROM coincidencias_recalculadas EXCEPT ALL SELECT candidato_id,escenario,unnest(instalaciones_ids)
        FROM gold_control_area_recarga))''', 'Coincidencias de área no reconciliadas')
    exigir_cero(c, '''SELECT count(*) FROM gold_solapamientos_candidatos x
        LEFT JOIN gold_candidatos a ON a.candidato_id=x.candidato_a_id LEFT JOIN gold_candidatos b ON b.candidato_id=x.candidato_b_id
        WHERE a.candidato_id IS NULL OR b.candidato_id IS NULL OR x.candidato_a_id>=x.candidato_b_id
        OR NOT ST_Intersects(a.geometria,b.geometria)
        OR x.dimension_interseccion<>ST_Dimension(ST_Intersection(a.geometria,b.geometria))
        OR x.superficie_interseccion_m2 IS DISTINCT FROM CASE WHEN ST_Dimension(ST_Intersection(a.geometria,b.geometria))=2
            THEN ST_Area(ST_Intersection(a.geometria,b.geometria)) END
        OR x.a_cubre_b<>ST_Covers(a.geometria,b.geometria) OR x.b_cubre_a<>ST_Covers(b.geometria,a.geometria)
        OR x.misma_categoria IS DISTINCT FROM (a.categoria_origen=b.categoria_origen)
        OR x.misma_representacion IS DISTINCT FROM (a.representacion=b.representacion)''', 'Solapamiento incoherente')
    expected = c.execute('''SELECT count(*) FROM gold_candidatos a JOIN gold_candidatos b
        ON ST_Intersects(a.geometria,b.geometria) WHERE a.candidato_id<b.candidato_id''').fetchone()[0]
    if expected!=c.execute('SELECT count(*) FROM gold_solapamientos_candidatos').fetchone()[0]:
        raise ValueError('Faltan solapamientos entre candidatos')
    exigir_cero(c, '''SELECT count(*) FROM gold_correspondencias_candidatos p WHERE p.punto_intersecta_area
        AND NOT EXISTS(SELECT 1 FROM gold_solapamientos_candidatos s WHERE
        s.candidato_a_id=least(p.candidato_punto_id,p.candidato_area_id)
        AND s.candidato_b_id=greatest(p.candidato_punto_id,p.candidato_area_id))''', 'Correspondencia espacial previa omitida')
    exigir_cero(c, '''SELECT count(*) FROM ((SELECT candidato_id,escenario FROM gold_proximidad_candidato_recarga EXCEPT
        SELECT candidato_id,escenario FROM gold_control_borde) UNION ALL (SELECT candidato_id,escenario FROM gold_control_borde
        EXCEPT SELECT candidato_id,escenario FROM gold_proximidad_candidato_recarga))''', 'Cobertura de borde incompleta')
    exigir_cero(c, '''SELECT count(*) FROM gold_control_borde b JOIN gold_proximidad_candidato_recarga p USING(candidato_id,escenario)
        JOIN gold_candidatos c USING(candidato_id) CROSS JOIN ambito_cv a WHERE
        b.estado_control<>p.estado_distancia OR b.distancia_borde_cv_m IS DISTINCT FROM
            CASE WHEN c.punto_representativo IS NOT NULL THEN ST_Distance(c.punto_representativo,ST_Boundary(a.geometria)) END
        OR (b.distancia_borde_cv_m IS NOT NULL AND (NOT isfinite(b.distancia_borde_cv_m) OR b.distancia_borde_cv_m<0))
        OR b.posible_influencia_exterior IS DISTINCT FROM CASE WHEN p.estado_distancia='calculable'
            THEN p.distancia_plana_m>b.distancia_borde_cv_m END''', 'Control de borde incoherente')
    return {'tablas':{n:c.execute('SELECT count(*) FROM '+n).fetchone()[0] for n in GOLD_CONTROLES_KEYS},
        'areas':consultar(c,'''SELECT escenario,count(*) areas,count(*) FILTER(WHERE instalaciones_intersectan>0) con_coincidencia,
            sum(instalaciones_intersectan) relaciones_area_instalacion FROM gold_control_area_recarga GROUP BY 1 ORDER BY 1'''),
        'solapamientos':consultar(c,'''SELECT dimension_interseccion,misma_categoria,misma_representacion,count(*) pares
            FROM gold_solapamientos_candidatos GROUP BY 1,2,3 ORDER BY 1,2,3'''),
        'borde':consultar(c,'''SELECT escenario,estado_control,count(*) elementos,
            count(*) FILTER(WHERE posible_influencia_exterior) avisos FROM gold_control_borde GROUP BY 1,2 ORDER BY 1,2''')}


# %% Gold: variables y normalización

GOLD_VARIABLES_KEYS = {
    'gold_referencias_normalizacion': ['version_gold', 'dimension'],
    'gold_variables_municipio': ['version_gold', 'codigo_municipio'],
    'gold_variables_candidato': ['version_gold', 'candidato_id', 'escenario']}


def gold_variables_prepare(c, version):
    for name, keys in (('gold_municipios', ['codigo_municipio']), ('gold_candidatos', ['candidato_id']),
                       ('gold_recarga_instalaciones', ['instalacion_id']),
                       ('gold_proximidad_candidato_recarga', ['candidato_id', 'escenario'])):
        comprobar_clave(c, name, keys)
    for name in ('gold_municipios', 'gold_candidatos', 'gold_recarga_instalaciones'):
        schema = {r[0]:r[1] for r in c.execute('DESCRIBE '+name).fetchall()}
        if schema['geometria'] != "GEOMETRY('EPSG:25830')":
            raise ValueError('CRS de entrada no admitido')
    exigir_cero(c, '''SELECT count(*) FROM gold_municipios WHERE geometria IS NULL
        OR NOT ST_IsValid(geometria) OR ST_IsEmpty(geometria) OR ST_Dimension(geometria)<>2''', 'Recinto no admitido')
    exigir_cero(c, '''SELECT count(*) FROM gold_candidatos x LEFT JOIN gold_municipios m
        ON x.codigo_municipio_contexto=m.codigo_municipio WHERE
        (x.numero_municipios_relacionados=1 AND (x.codigo_municipio_contexto IS NULL OR m.codigo_municipio IS NULL))
        OR (x.numero_municipios_relacionados<>1 AND x.codigo_municipio_contexto IS NOT NULL)''', 'Contexto municipal incoherente')
    c.execute('CREATE TEMP TABLE escenarios(escenario VARCHAR)')
    c.executemany('INSERT INTO escenarios VALUES (?)', [(x,) for x in GOLD_CANDIDATOS_SCENARIOS])
    c.execute('''CREATE TEMP TABLE oferta AS SELECT e.escenario,r.instalacion_id,r.geometria
        FROM gold_recarga_instalaciones r CROSS JOIN escenarios e
        WHERE (e.escenario='incluye_ausentes_historicas' OR (r.presente_copia_principal AND
            (e.escenario='incluye_pendientes' OR r.estado_admision='admitida_version_principal')))
        AND r.estado_geometria='admitida' AND r.geometria IS NOT NULL
        AND ST_IsValid(r.geometria) AND NOT ST_IsEmpty(r.geometria)''')
    c.execute('''CREATE TEMP TABLE evaluables AS SELECT *,CASE
        WHEN geometria IS NULL OR NOT ST_IsValid(geometria) OR ST_IsEmpty(geometria)
            THEN 'sin_geometria_admisible'
        WHEN representacion='area' AND punto_representativo IS NULL THEN 'sin_superficie_en_cv'
        ELSE 'calculable' END estado_geometria_variable FROM gold_candidatos''')
    print('Calculando mínimos a geometría original en tres escenarios...', flush=True)
    c.execute('''CREATE TEMP TABLE minimos AS SELECT x.candidato_id,o.escenario,
        min(ST_Distance(x.geometria,o.geometria)) distancia_minima_geometria_m
        FROM evaluables x CROSS JOIN oferta o WHERE x.estado_geometria_variable='calculable' GROUP BY 1,2''')
    print('Calculando borde CV para la geometría original...', flush=True)
    c.execute('''CREATE TEMP TABLE borde_cv AS SELECT ST_Boundary(ST_Union_Agg(geometria)) geometria FROM gold_municipios''')
    comprobar_clave(c,'gold_control_borde',['candidato_id','escenario'])
    c.execute('''CREATE TEMP TABLE cotas_borde AS SELECT candidato_id,distancia_borde_cv_m cota
        FROM gold_control_borde WHERE escenario='principal_admitida' ''')
    exigir_cero(c, '''SELECT count(*) FROM evaluables x LEFT JOIN cotas_borde b USING(candidato_id)
        WHERE estado_geometria_variable='calculable' AND (b.cota IS NULL OR NOT isfinite(b.cota) OR b.cota<0
            OR x.punto_representativo IS NULL OR NOT ST_Intersects(x.geometria,x.punto_representativo))''', 'Cota de borde no admisible')
    # El representante pertenece al área: su distancia acota el mínimo; el recorte conserva el borde más cercano.
    c.execute('''CREATE TEMP TABLE bordes_recortados AS SELECT x.candidato_id,ST_Intersection(b.geometria,
        ST_MakeEnvelope(ST_XMin(x.geometria)-r.cota-1,ST_YMin(x.geometria)-r.cota-1,
            ST_XMax(x.geometria)+r.cota+1,ST_YMax(x.geometria)+r.cota+1)) geometria
        FROM evaluables x JOIN cotas_borde r USING(candidato_id) CROSS JOIN borde_cv b
        WHERE estado_geometria_variable='calculable' AND representacion='area' ''')
    exigir_cero(c, '''SELECT count(*) FROM bordes_recortados WHERE geometria IS NULL
        OR ST_IsEmpty(geometria) OR NOT ST_IsValid(geometria)''', 'Recorte de borde no admisible')
    c.execute('''CREATE TEMP TABLE distancias_borde_geometria AS SELECT x.candidato_id,
        CASE WHEN estado_geometria_variable='calculable' THEN CASE WHEN representacion='punto' THEN r.cota
            ELSE ST_Distance(x.geometria,b.geometria) END END distancia_geometria_borde_cv_m
        FROM evaluables x LEFT JOIN cotas_borde r USING(candidato_id) LEFT JOIN bordes_recortados b USING(candidato_id)''')
    exigir_cero(c, '''SELECT count(*) FROM distancias_borde_geometria d JOIN evaluables x USING(candidato_id)
        LEFT JOIN cotas_borde r USING(candidato_id) WHERE estado_geometria_variable='calculable'
        AND (distancia_geometria_borde_cv_m IS NULL OR NOT isfinite(distancia_geometria_borde_cv_m)
            OR distancia_geometria_borde_cv_m<0 OR distancia_geometria_borde_cv_m>r.cota+1e-8)''', 'Distancia a borde no admisible')
    c.execute('''CREATE TEMP TABLE valores AS SELECT 'demanda' dimension,vehiculos_distintivo_0::DOUBLE valor FROM gold_municipios
        UNION ALL SELECT 'proximidad',distancia_minima_geometria_m FROM evaluables x
        LEFT JOIN minimos m ON x.candidato_id=m.candidato_id AND m.escenario='principal_admitida' ''')
    c.execute('''CREATE TABLE gold_referencias_normalizacion AS SELECT ?::VARCHAR version_gold,d.dimension,
        CASE WHEN d.dimension='demanda' THEN 'vehiculos_distintivo_0' ELSE 'distancia_minima_geometria_m' END nombre_variable,
        CASE WHEN d.dimension='demanda' THEN 'vehiculos' ELSE 'metros' END unidad,
        CASE WHEN d.dimension='demanda' THEN 'no_aplica' ELSE 'principal_admitida' END escenario_referencia,
        CASE WHEN d.dimension='demanda' THEN 'municipios_cv' ELSE 'candidatos_cv' END universo_referencia,
        count(v.dimension)::BIGINT n_total,
        count(*) FILTER(WHERE isfinite(valor) AND valor>=0)::BIGINT n_validos,
        count(*) FILTER(WHERE v.dimension IS NOT NULL AND valor IS NULL)::BIGINT n_nulos,
        count(*) FILTER(WHERE valor IS NOT NULL AND (NOT isfinite(valor) OR valor<0))::BIGINT n_invalidos,
        min(valor) FILTER(WHERE isfinite(valor) AND valor>=0) minimo,
        max(valor) FILTER(WHERE isfinite(valor) AND valor>=0) maximo,
        CASE WHEN maximo IS NULL THEN 'sin_valores_admisibles' WHEN maximo=0 THEN 'todos_cero'
            WHEN minimo=maximo THEN 'sin_variacion_positiva' ELSE 'con_variacion' END estado_referencia
        FROM (VALUES ('demanda'),('proximidad')) d(dimension) LEFT JOIN valores v USING(dimension)
        GROUP BY d.dimension''', [version])
    c.execute('''CREATE MACRO estado_norm(x,lo,hi) AS CASE WHEN x IS NULL THEN 'sin_valor'
        WHEN NOT isfinite(x) OR x<0 THEN 'valor_no_admisible'
        WHEN lo IS NULL OR hi IS NULL THEN 'sin_referencia'
        WHEN NOT isfinite(lo) OR NOT isfinite(hi) OR lo<0 OR hi<lo THEN 'referencia_no_admisible'
        WHEN x>hi THEN 'fuera_referencia' WHEN hi=0 THEN 'todos_cero'
        WHEN lo=hi THEN 'sin_variacion_positiva' ELSE 'calculable' END''')
    c.execute('''CREATE MACRO norma_log(x,lo,hi) AS CASE
        WHEN estado_norm(x,lo,hi)='todos_cero' THEN 0.0
        WHEN estado_norm(x,lo,hi) IN ('calculable','sin_variacion_positiva') THEN ln(1.0+x)/ln(1.0+hi) END''')
    c.execute('''CREATE MACRO norma_lineal(x,lo,hi) AS CASE
        WHEN estado_norm(x,lo,hi)='todos_cero' THEN 0.0
        WHEN estado_norm(x,lo,hi) IN ('calculable','sin_variacion_positiva') THEN x/hi END''')
    c.execute('''CREATE TABLE gold_variables_municipio AS SELECT ?::VARCHAR version_gold,m.codigo_municipio,
        m.vehiculos_distintivo_0,norma_log(m.vehiculos_distintivo_0,r.minimo,r.maximo) demanda_normalizada_log,
        norma_lineal(m.vehiculos_distintivo_0,r.minimo,r.maximo) demanda_normalizada_lineal,
        estado_norm(m.vehiculos_distintivo_0,r.minimo,r.maximo) estado_normalizacion
        FROM gold_municipios m CROSS JOIN gold_referencias_normalizacion r WHERE r.dimension='demanda' ''', [version])
    c.execute('''CREATE TABLE gold_variables_candidato AS SELECT ?::VARCHAR version_gold,x.candidato_id,e.escenario,
        x.codigo_municipio_contexto,x.estado_relacion_municipal estado_contexto_municipal,m.distancia_minima_geometria_m,
        CASE WHEN x.estado_geometria_variable<>'calculable' THEN x.estado_geometria_variable
            WHEN m.distancia_minima_geometria_m IS NULL THEN 'sin_recarga_admisible' ELSE 'calculable' END estado_distancia,
        b.distancia_geometria_borde_cv_m,
        CASE WHEN estado_distancia='calculable' THEN m.distancia_minima_geometria_m>b.distancia_geometria_borde_cv_m END
            posible_influencia_exterior_geometria,
        norma_log(m.distancia_minima_geometria_m,r.minimo,r.maximo) proximidad_normalizada_log,
        norma_lineal(m.distancia_minima_geometria_m,r.minimo,r.maximo) proximidad_normalizada_lineal,
        estado_norm(m.distancia_minima_geometria_m,r.minimo,r.maximo) estado_normalizacion
        FROM evaluables x CROSS JOIN escenarios e LEFT JOIN minimos m USING(candidato_id,escenario)
        LEFT JOIN distancias_borde_geometria b USING(candidato_id)
        CROSS JOIN gold_referencias_normalizacion r WHERE r.dimension='proximidad' ''', [version])


def gold_variables_validate(c, version):
    for name, keys in GOLD_VARIABLES_KEYS.items():
        comprobar_clave(c, name, keys)
        if c.execute(f'SELECT count(*) FROM {name} WHERE version_gold IS DISTINCT FROM ?', [version]).fetchone()[0]:
            raise ValueError('Versión de variables distinta')
    for source, target, fields in (('gold_municipios','gold_variables_municipio','codigo_municipio'),
            ('gold_proximidad_candidato_recarga','gold_variables_candidato','candidato_id,escenario')):
        exigir_cero(c, f'''SELECT count(*) FROM ((SELECT {fields} FROM {source} EXCEPT SELECT {fields} FROM {target})
            UNION ALL (SELECT {fields} FROM {target} EXCEPT SELECT {fields} FROM {source}))''', 'Cobertura de variables incompleta')
    exigir_cero(c, '''SELECT count(*) FROM gold_variables_municipio v JOIN gold_municipios m USING(codigo_municipio)
        WHERE v.vehiculos_distintivo_0 IS DISTINCT FROM m.vehiculos_distintivo_0''', 'Demanda original alterada')
    for dimension in ('demanda','proximidad'):
        observed = [r[0] for r in c.execute('SELECT valor FROM valores WHERE dimension=?',[dimension]).fetchall()]
        valid = [x for x in observed if x is not None and math.isfinite(x) and x>=0]
        lo,hi = (min(valid),max(valid)) if valid else (None,None)
        state = ('sin_valores_admisibles' if hi is None else 'todos_cero' if hi==0
                 else 'sin_variacion_positiva' if lo==hi else 'con_variacion')
        expected = (len(observed),len(valid),observed.count(None),len(observed)-len(valid)-observed.count(None),lo,hi,state)
        row = c.execute('''SELECT n_total,n_validos,n_nulos,n_invalidos,minimo,maximo,estado_referencia
            FROM gold_referencias_normalizacion WHERE dimension=?''',[dimension]).fetchone()
        if row != expected:
            raise ValueError('Referencia de normalización incoherente')
    if c.execute('SELECT count(*) FROM gold_referencias_normalizacion').fetchone()[0]!=2:
        raise ValueError('Dimensiones de referencia distintas')
    for table, raw, prefix, dimension in (('gold_variables_municipio','vehiculos_distintivo_0','demanda','demanda'),
            ('gold_variables_candidato','distancia_minima_geometria_m','proximidad','proximidad')):
        ref = c.execute('SELECT minimo,maximo FROM gold_referencias_normalizacion WHERE dimension=?', [dimension]).fetchone()
        for x, log, linear, state in c.execute(f'SELECT {raw},{prefix}_normalizada_log,{prefix}_normalizada_lineal,estado_normalizacion FROM {table}').fetchall():
            expected = ('sin_valor' if x is None else 'valor_no_admisible' if not math.isfinite(x) or x<0
                else 'sin_referencia' if ref[1] is None else 'fuera_referencia' if x>ref[1]
                else 'todos_cero' if ref[1]==0 else 'sin_variacion_positiva' if ref[0]==ref[1] else 'calculable')
            if state != expected:
                raise ValueError('Estado de normalización incoherente')
            if expected in ('calculable','todos_cero','sin_variacion_positiva'):
                a,b = (math.log1p(x)/math.log1p(ref[1]),x/ref[1]) if ref[1] else (0.0,0.0)
                if (log is None or linear is None or not math.isclose(log,a,rel_tol=1e-12,abs_tol=1e-12)
                        or not math.isclose(linear,b,rel_tol=1e-12,abs_tol=1e-12) or not 0<=log<=1 or not 0<=linear<=1):
                    raise ValueError('Normalización numérica incoherente')
            elif log is not None or linear is not None:
                raise ValueError('Normalización desconocida rellenada')
    exigir_cero(c, '''SELECT count(*) FROM gold_variables_candidato v JOIN gold_candidatos x USING(candidato_id)
        JOIN gold_proximidad_candidato_recarga p USING(candidato_id,escenario)
        LEFT JOIN gold_control_area_recarga a USING(candidato_id,escenario)
        JOIN distancias_borde_geometria b USING(candidato_id)
        JOIN evaluables e USING(candidato_id) LEFT JOIN minimos m USING(candidato_id,escenario) WHERE
        v.codigo_municipio_contexto IS DISTINCT FROM x.codigo_municipio_contexto
        OR v.estado_contexto_municipal IS DISTINCT FROM x.estado_relacion_municipal
        OR v.estado_distancia IS DISTINCT FROM CASE WHEN e.estado_geometria_variable<>'calculable' THEN e.estado_geometria_variable
            WHEN m.distancia_minima_geometria_m IS NULL THEN 'sin_recarga_admisible' ELSE 'calculable' END
        OR v.distancia_minima_geometria_m IS DISTINCT FROM m.distancia_minima_geometria_m
        OR (v.estado_distancia='calculable' AND (v.distancia_minima_geometria_m IS NULL
            OR NOT isfinite(v.distancia_minima_geometria_m) OR v.distancia_minima_geometria_m<0
            OR v.distancia_minima_geometria_m>p.distancia_plana_m+1e-8))
        OR (v.estado_distancia<>'calculable' AND v.distancia_minima_geometria_m IS NOT NULL)
        OR (x.representacion='punto' AND v.estado_distancia='calculable'
            AND v.distancia_minima_geometria_m IS DISTINCT FROM p.distancia_plana_m)
        OR (v.estado_distancia='calculable' AND x.representacion='area' AND
            (a.candidato_id IS NULL OR (v.distancia_minima_geometria_m=0)<>(a.instalaciones_intersectan>0)))
        OR v.posible_influencia_exterior_geometria IS DISTINCT FROM CASE WHEN v.estado_distancia='calculable'
            THEN v.distancia_minima_geometria_m>v.distancia_geometria_borde_cv_m END
        OR v.distancia_geometria_borde_cv_m IS DISTINCT FROM b.distancia_geometria_borde_cv_m''', 'Geometría o contexto de variables incoherente')
    exigir_cero(c, '''SELECT count(*) FROM gold_variables_candidato a JOIN gold_variables_candidato b USING(candidato_id)
        WHERE ((a.escenario='principal_admitida' AND b.escenario='incluye_pendientes') OR
            (a.escenario='incluye_pendientes' AND b.escenario='incluye_ausentes_historicas'))
        AND a.distancia_minima_geometria_m<b.distancia_minima_geometria_m''', 'Distancias no monótonas por escenario')
    return {'tablas': {n:c.execute('SELECT count(*) FROM '+n).fetchone()[0] for n in GOLD_VARIABLES_KEYS},
        'referencias': consultar(c,'SELECT * EXCLUDE(version_gold) FROM gold_referencias_normalizacion ORDER BY dimension'),
        'distancias': consultar(c,'''SELECT escenario,estado_distancia,count(*) elementos,
            count(*) FILTER(WHERE distancia_minima_geometria_m=0) ceros,median(distancia_minima_geometria_m) mediana_m,
            max(distancia_minima_geometria_m) maximo_m,count(*) FILTER(WHERE posible_influencia_exterior_geometria) avisos_borde
            FROM gold_variables_candidato GROUP BY 1,2 ORDER BY 1,2''')}


# %% Entradas, exportación y controles compartidos del recorrido principal

def guardar_json(path, contenido):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(contenido, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')


def conectar(temporal):
    c = duckdb.connect(str(temporal / 'trabajo.duckdb'), config={
        'threads': 2, 'memory_limit': '1GB', 'storage_compatibility_version': 'v1.5.0',
        'extension_directory': str(RAIZ / '.venv/duckdb_extensions'),
        'temp_directory': str(temporal / 'desborde')})
    try:
        c.execute('LOAD spatial')
        c.execute("SET TimeZone='UTC'")
    except Exception:
        c.close()
        raise
    return c


def comprobar_entradas():
    if duckdb.__version__ != '1.5.0':
        raise ValueError('Instalar las versiones de requirements.txt; se requiere DuckDB 1.5.0')
    with tempfile.TemporaryDirectory() as temporal:
        with conectar(Path(temporal)):
            pass
    archivos = [comprobar_archivo(RAIZ, Path(__file__))]
    for ruta in ENTRADAS.values():
        archivos.append(comprobar_archivo(RAIZ, RAIZ / ruta))
    catalogo = json.loads((RAIZ / ENTRADAS['catalogo']).read_text(encoding='utf-8'))
    bronze = bronze_build_manifest(RAIZ, catalogo)
    archivos.extend(bronze['archivos'])
    recarga_geo_evidence(RAIZ, RAIZ / ENTRADAS['metadatos_recarga'])
    archivos.append(comprobar_archivo(RAIZ, RAIZ / (ENTRADAS['metadatos_recarga'] + '.adquisicion.json')))
    evidencia = json.loads((RAIZ / ENTRADAS['metadatos_capacidad']).read_text(encoding='utf-8'))
    metadatos = metadatos_ide_verify(RAIZ, evidencia)
    archivos.extend(comprobar_archivo(RAIZ, RAIZ / e['ruta'], e) for e in evidencia['archivos'].values())
    recibo_path = RAIZ / (ENTRADAS['recarga_principal'] + '.adquisicion.json')
    archivos.append(comprobar_archivo(RAIZ, recibo_path))
    recibo = json.loads(recibo_path.read_text(encoding='utf-8'))
    principal = comprobar_archivo(RAIZ, RAIZ / ENTRADAS['recarga_principal'], recibo)
    archivos.append(principal)
    adquisicion = datetime.fromisoformat(recibo['fecha_utc'])
    if adquisicion.tzinfo is None:
        raise ValueError('Adquisición de recarga sin zona horaria')
    origenes = {e['ruta']: e['recurso_id'] + ':' + e['sha256'] for e in bronze['archivos']}
    origenes['recarga_principal'] = 'miteco_recarga:' + principal['sha256']
    return bronze, archivos, metadatos, adquisicion, origenes


def exportar_tablas(c, carpeta, claves):
    carpeta.mkdir(parents=True, exist_ok=False)
    salidas = {}
    for nombre, clave in claves.items():
        if clave:
            comprobar_clave(c, nombre, clave)
        path = carpeta / (nombre + '.parquet')
        c.execute(f'COPY {identificador_sql(nombre)} TO ? (FORMAT PARQUET, COMPRESSION ZSTD)', [str(path)])
        esquema = [(r[0], r[1]) for r in c.execute('DESCRIBE ' + identificador_sql(nombre)).fetchall()]
        restaurado = [(r[0], r[1]) for r in c.execute('DESCRIBE SELECT * FROM read_parquet(?)', [str(path)]).fetchall()]
        if esquema != restaurado:
            raise ValueError(f'Exportación alteró el esquema: {nombre}')
        geometrias = [campo for campo, tipo in esquema if tipo.startswith('GEOMETRY')]
        campos = ','.join('ST_AsWKB(' + identificador_sql(campo) + ') AS ' + identificador_sql(campo)
                          if campo in geometrias else identificador_sql(campo) for campo, _ in esquema)
        exigir_cero(c, f'''SELECT count(*) FROM (
            (SELECT {campos} FROM {identificador_sql(nombre)} EXCEPT ALL
             SELECT {campos} FROM read_parquet({texto_sql(str(path))})) UNION ALL
            (SELECT {campos} FROM read_parquet({texto_sql(str(path))}) EXCEPT ALL
             SELECT {campos} FROM {identificador_sql(nombre)}))''', f'Exportación alteró valores: {nombre}')
        if geometrias:
            metadatos = dict(c.execute('SELECT decode(key),decode(value) FROM parquet_kv_metadata(?)', [str(path)]).fetchall())
            geo = json.loads(metadatos['geo'])
            for campo in geometrias:
                if geo['columns'][campo]['crs']['id'] != {'authority': 'EPSG', 'code': 25830}:
                    raise ValueError(f'CRS GeoParquet incorrecto: {nombre}/{campo}')
        salidas[nombre] = {'ruta': path.relative_to(RAIZ).as_posix(), **huella(path),
            'clave': clave, 'esquema': [{'campo': a, 'tipo': b} for a, b in esquema],
            'filas': c.execute('SELECT count(*) FROM ' + identificador_sql(nombre)).fetchone()[0]}
    return salidas


def crear_vistas_revision(c):
    for nombre in RECARGA_KEYS:
        c.execute(f'CREATE TEMP VIEW {identificador_sql("recarga__" + nombre)} AS SELECT * FROM {identificador_sql(nombre)}')
    c.execute('CREATE TEMP VIEW recarga_espacial__recarga_instalaciones_geo AS SELECT * FROM recarga_instalaciones_geo')
    c.execute('CREATE TEMP VIEW base_territorial__recintos_municipales AS SELECT * FROM recintos')


def integrar_relaciones(c):
    c.execute('''CREATE TABLE relaciones_espaciales(tabla_entidad VARCHAR,clave_entidad VARCHAR,
        codigo_municipio VARCHAR,metodo VARCHAR,dimension_interseccion SMALLINT,
        longitud_interseccion_m DOUBLE,superficie_interseccion_m2 DOUBLE)''')
    c.execute('''CREATE TABLE estado_relaciones_espaciales(tabla_entidad VARCHAR,clave_entidad VARCHAR,
        estado_geometria VARCHAR,numero_municipios BIGINT,estado_relacion VARCHAR)''')
    calidad = {}
    for bloque, nombre, clave, tipo in INTEGRACION_SPATIAL:
        fuente = bloque + '__' + nombre
        print(f'  Relaciones municipales de {nombre}...', flush=True)
        if nombre == 'red_tramos':
            # Lotes provinciales: acotan temporales y evitan la caída nativa observada al cerrar DuckDB.
            for (provincia,) in c.execute(f'SELECT DISTINCT provincia_origen FROM {identificador_sql(fuente)} ORDER BY 1').fetchall():
                c.execute(f'CREATE OR REPLACE TEMP VIEW lote_red AS SELECT * FROM {identificador_sql(fuente)} '
                          f'WHERE provincia_origen={texto_sql(provincia)}')
                integracion_spatial_relations(c, 'lote_red', clave, nombre, tipo)
            c.execute('DROP VIEW lote_red')
        else:
            integracion_spatial_relations(c, fuente, clave, nombre, tipo)
        if c.execute('SELECT count(*) FROM ' + identificador_sql(fuente)).fetchone()[0] != c.execute(
                'SELECT count(*) FROM estado_relaciones_espaciales WHERE tabla_entidad=?', [nombre]).fetchone()[0]:
            raise ValueError('Se pierden entidades entre lotes espaciales')
        calidad[nombre] = consultar(c, 'SELECT estado_relacion,count(*) entidades,sum(numero_municipios) relaciones '
            f'FROM estado_relaciones_espaciales WHERE tabla_entidad={texto_sql(nombre)} GROUP BY 1 ORDER BY 1')
    for nombre in ('relaciones_espaciales', 'estado_relaciones_espaciales'):
        comprobar_clave(c, nombre, INTEGRACION_OUTPUT_KEYS[nombre])
    exigir_cero(c, '''SELECT count(*) FROM relaciones_espaciales ANTI JOIN
        base_territorial__municipios USING(codigo_municipio)''', 'Relación con municipio desconocido')
    return calidad


def crear_estado_fuentes(c):
    c.execute('''CREATE TABLE estado_fuentes(bloque VARCHAR,estado VARCHAR,
        integracion_municipal VARCHAR,limitacion VARCHAR)''')
    limitaciones = {
        'base_territorial': ('admitida', 'INE 2025 y recintos consultados 2026; sin acreditación de límites históricos.'),
        'demografia_parque': ('admitida_por_codigo', 'Periodos separados; 12066 y PP000 aparte; DGT sin municipio conservado; Distintivo 0 aproximado.'),
        'recarga': ('por_instalacion_validada', 'Sin suma de potencias como capacidad simultánea; fecha modificación no es adquisición.'),
        'recarga_espacial': ('parcial_con_incidencias', 'Conflictos documentados pendientes; admisión documental no certifica posición fina.'),
        'trafico_red': ('intersecciones_no_exclusivas', 'Cinco geometrías no admitidas; contactos/límites no reparten IMD/longitud; red no enrutable acreditada.'),
        'poi_osm': ('intersecciones_no_exclusivas', 'Sin fusión de punto/área; objetos exteriores conservados; POI no acredita disponibilidad.'),
        'capacidad_demanda': ('complementada_por_capa_espacial', 'Tabla histórica conservada; fecha/CRS posteriores en capacidad_demanda_espacial; capacidad no aditiva ni garantía de conexión.'),
        'capacidad_demanda_espacial': ('por_punto_validado', 'Publicación y periodo mensual separados; corte exacto desconocido; solo relaciones concordantes admitidas; capacidad orientativa y no aditiva.'),
    }
    c.executemany('INSERT INTO estado_fuentes VALUES (?,?,?,?)',
                  [(b, 'preparada_con_limitaciones', *valores) for b, valores in limitaciones.items()])


if __name__ == '__main__':
    main()
