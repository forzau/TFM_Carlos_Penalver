"""Fase 3 Silver: parque DGT, población y renta; originales intactos, sin red."""

import argparse
import csv
import json
import math
import re
import tempfile
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path

import duckdb
from openpyxl import load_workbook

from auditar_fuentes import ident, query
from catalogar_bronze import ROOT, fingerprint
from preparar_base_territorial import verify_inputs

TXT = 'Fuentes/parque_vehiculos_202512.txt'
XLSX = 'datos/bronze/dgt_municipal/2025/DatosMunicipalesGeneral_2025.xlsx'
PDF = 'datos/bronze/dgt_municipal/2025/ContenidoFichasGeneral2025.pdf'
POP = 'Fuentes/Poblacion.csv'
INCOME = 'Fuentes/Renta.csv'
INPUTS = {TXT: 'dgt_microdatos_202512', XLSX: 'dgt_municipal_2025',
          PDF: 'dgt_municipal_2025', POP: 'ine_poblacion', INCOME: 'ine_renta'}
PROVINCES = ('03', '12', '46')
SCHEMAS = {
    'poblacion_municipal': {'codigo_municipio': 'VARCHAR', 'sexo': 'VARCHAR', 'anio': 'SMALLINT',
        'poblacion': 'BIGINT', 'estado_valor': 'VARCHAR', 'valor_origen': 'VARCHAR', 'registro_origen': 'BIGINT'},
    'renta_territorial': {'nivel_territorial': 'VARCHAR', 'codigo_territorio': 'VARCHAR',
        'codigo_municipio': 'VARCHAR', 'codigo_distrito': 'VARCHAR', 'codigo_seccion': 'VARCHAR',
        'indicador': 'VARCHAR', 'anio': 'SMALLINT', 'valor_eur': 'DECIMAL(18,2)',
        'estado_valor': 'VARCHAR', 'valor_origen': 'VARCHAR', 'registro_origen': 'BIGINT'},
    'incidencias': {'tabla': 'VARCHAR', 'origen_id': 'VARCHAR', 'registro_origen': 'BIGINT',
        'codigo_municipio': 'VARCHAR', 'campo': 'VARCHAR', 'valor_origen': 'VARCHAR', 'motivo': 'VARCHAR'},
}
SCHEMAS['renta_codigos_historicos'] = {**SCHEMAS['renta_territorial'], 'motivo_separacion': 'VARCHAR'}
KEYS = {'parque_municipal': ['codigo_municipio', 'fecha_referencia'],
        'parque_sin_municipio': ['codigo_agrupacion_origen', 'fecha_referencia'],
        'parque_detalle_observado': ['origen_id', 'registro_origen'],
        'poblacion_municipal': ['codigo_municipio', 'sexo', 'anio'],
        'renta_territorial': ['nivel_territorial', 'codigo_territorio', 'indicador', 'anio'],
        'renta_codigos_historicos': ['nivel_territorial', 'codigo_territorio', 'indicador', 'anio']}


def ine_number(raw, integer=False):
    text = raw.strip()
    if text in ('', '.', '..', '...'):
        return None, 'no_informado'
    if not re.fullmatch(r'-?(?:[0-9]+|[0-9]{1,3}(?:\.[0-9]{3})+)(?:,[0-9]+)?', text):
        return None, 'error_conversion'
    value = Decimal(text.replace('.', '').replace(',', '.'))
    if integer:
        if value != value.to_integral_value() or not 0 <= value <= 2**63 - 1:
            return None, 'error_conversion'
        return int(value), 'informado'
    if (',' in text and len(text.split(',')[1]) > 2) or abs(value) >= Decimal('10000000000000000'):
        return None, 'error_conversion'
    return value, 'informado'


def territory_code(raw, width):
    match = re.match(r'^([0-9]{' + str(width) + r'})(?:\s|$)', raw)
    if not match:
        raise ValueError(f'Código territorial de longitud {width} no válido: {raw!r}')
    return match[1]


def csv_rows(path, headers):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.reader(stream, delimiter=';', strict=True)
        if next(reader) != headers:
            raise ValueError(f'Cabecera inesperada: {path.name}')
        for ordinal, values in enumerate(reader, 1):
            if len(values) != len(headers):
                raise ValueError(f'{path.name}: ancho inesperado en registro {ordinal}')
            yield ordinal, dict(zip(headers, values))


def prepare_ine(connection, population, income, codes, temporary, origins):
    counts = Counter()
    with (temporary / 'incidencias.csv').open('w', encoding='utf-8', newline='') as inc_stream:
        incidents = csv.writer(inc_stream)
        incidents.writerow(SCHEMAS['incidencias'])
        with (temporary / 'poblacion_municipal.csv').open('w', encoding='utf-8', newline='') as target:
            writer = csv.writer(target)
            writer.writerow(SCHEMAS['poblacion_municipal'])
            for ordinal, row in csv_rows(population, ['Total Nacional', 'Municipios', 'Sexo', 'Periodo', 'Total']):
                counts['poblacion_nacional'] += 1
                if row['Municipios'][:2] not in PROVINCES:
                    continue
                code = territory_code(row['Municipios'], 5)
                if code not in codes or row['Sexo'] not in ('Total', 'Hombres', 'Mujeres'):
                    raise ValueError(f'Población: territorio o sexo no admitido en registro {ordinal}')
                if not re.fullmatch(r'[0-9]{4}', row['Periodo']):
                    raise ValueError(f'Periodo de población no válido: registro {ordinal}')
                value, state = ine_number(row['Total'], integer=True)
                writer.writerow([code, row['Sexo'], int(row['Periodo']), value, state, row['Total'], ordinal])
                counts['poblacion_cv'] += 1
                if state == 'error_conversion':
                    incidents.writerow(['poblacion_municipal', origins[POP], ordinal, code, 'poblacion', row['Total'], state])
        with (temporary / 'renta_territorial.csv').open('w', encoding='utf-8', newline='') as target, (temporary / 'renta_codigos_historicos.csv').open('w', encoding='utf-8', newline='') as historic:
            writer, historical_writer = csv.writer(target), csv.writer(historic)
            writer.writerow(SCHEMAS['renta_territorial'])
            historical_writer.writerow(SCHEMAS['renta_codigos_historicos'])
            headers = ['Municipios', 'Distritos', 'Secciones', 'Indicadores de renta media', 'Periodo', 'Total']
            for ordinal, row in csv_rows(income, headers):
                counts['renta_nacional'] += 1
                if not any(row[field][:2] in PROVINCES for field in ('Municipios', 'Distritos', 'Secciones')):
                    continue
                code = territory_code(row['Municipios'], 5)
                if code not in codes and code != '12066':
                    raise ValueError(f'Renta: municipio ajeno a la referencia en registro {ordinal}: {code}')
                district = territory_code(row['Distritos'], 7) if row['Distritos'] else None
                section = territory_code(row['Secciones'], 10) if row['Secciones'] else None
                if (district and not district.startswith(code)) or (section and (not district or not section.startswith(district))):
                    raise ValueError(f'Jerarquía territorial de renta incoherente: registro {ordinal}')
                if not re.fullmatch(r'[0-9]{4}', row['Periodo']) or not row['Indicadores de renta media'].strip():
                    raise ValueError(f'Clave de renta incompleta: registro {ordinal}')
                level, territory = ('seccion', section) if section else ('distrito', district) if district else ('municipio', code)
                value, state = ine_number(row['Total'])
                values = [level, territory, code, district, section, row['Indicadores de renta media'],
                          int(row['Periodo']), value, state, row['Total'], ordinal]
                if code == '12066':
                    historical_writer.writerow(values + ['codigo_historico_fuera_de_referencia_2025'])
                    counts['renta_historica'] += 1
                else:
                    writer.writerow(values)
                    counts['renta_cv_vigente'] += 1
                if state == 'error_conversion':
                    incidents.writerow(['renta_codigos_historicos' if code == '12066' else 'renta_territorial',
                                        origins[INCOME], ordinal, code, 'valor_eur', row['Total'], state])
    for table, schema in SCHEMAS.items():
        # El CSV temporal usa vacío como NULL; el texto original vacío se restaura expresamente.
        connection.execute(f"CREATE TEMP TABLE {ident(table)} AS SELECT * REPLACE(coalesce(valor_origen,'') AS valor_origen) FROM read_csv(?,columns=?,header=true,delim=',',strict_mode=true,ignore_errors=false)", [str(temporary / f'{table}.csv'), schema])
    for table, counter in [('poblacion_municipal', 'poblacion_cv'),
                           ('renta_territorial', 'renta_cv_vigente'),
                           ('renta_codigos_historicos', 'renta_historica')]:
        if connection.execute(f'SELECT count(*) FROM {ident(table)}').fetchone()[0] != counts[counter]:
            raise ValueError(f'Pérdida de registros INE: {table}')
    return dict(counts)


def prepare_municipal_park(connection, path, codes):
    connection.execute('CREATE TEMP TABLE parque_municipal (codigo_municipio VARCHAR,fecha_referencia DATE,parque_total BIGINT,vehiculos_distintivo_0 BIGINT,electrificado_origen BIGINT)')
    connection.execute('CREATE TEMP TABLE parque_sin_municipio (codigo_agrupacion_origen VARCHAR,codigo_provincia VARCHAR,fecha_referencia DATE,parque_total BIGINT,vehiculos_distintivo_0 BIGINT,electrificado_origen BIGINT)')
    workbook = load_workbook(path, read_only=True, data_only=True)
    municipalities, unspecified, seen = [], [], set()
    try:
        rows = workbook['DatosMunicipalesGeneral_2025'].iter_rows(values_only=True)
        headers = next(rows)
        fields = ('Código INE', 'Parque Total', 'Distintivo 0', 'Electrificado')
        if len(set(headers)) != len(headers) or not set(fields).issubset(headers):
            raise ValueError('Cabecera municipal DGT inesperada')
        for ordinal, values in enumerate(rows, 1):
            if all(value is None for value in values):
                continue
            row = dict(zip(headers, values))
            code = row['Código INE']
            if not isinstance(code, str) or not re.fullmatch('[0-9]{5}', code) or code in seen:
                raise ValueError(f'Código municipal DGT no válido o repetido: registro {ordinal}')
            seen.add(code)
            if code[:2] not in PROVINCES:
                continue
            measures = [row[field] for field in fields[1:]]
            if any(type(value) not in (int, float) or not math.isfinite(value) or not 0 <= value <= 2**63 - 1 or value != int(value) for value in measures):
                raise ValueError(f'Recuentos DGT no admitidos: {code}')
            measures = [int(value) for value in measures]
            if code.endswith('000'):
                unspecified.append([code, code[:2], date(2025, 12, 31), *measures])
            elif code in codes:
                municipalities.append([code, date(2025, 12, 31), *measures])
            else:
                raise ValueError(f'Municipio DGT no vigente en la referencia: {code}')
    finally:
        workbook.close()
    if {row[0] for row in municipalities} != set(codes):
        raise ValueError('Cobertura municipal DGT incompleta')
    if {row[0] for row in unspecified} != {province + '000' for province in {code[:2] for code in codes}}:
        raise ValueError('Agrupaciones provinciales sin municipio DGT inesperadas')
    connection.executemany('INSERT INTO parque_municipal VALUES (?,?,?,?,?)', municipalities)
    connection.executemany('INSERT INTO parque_sin_municipio VALUES (?,?,?,?,?,?)', unspecified)
    return {'filas_nacionales': len(seen), 'municipios': len(municipalities), 'sin_municipio': len(unspecified)}


def prepare_detail(connection, path, origin):
    # Tabla de trabajo en disco: solo atributos necesarios, sin cargar el TXT en Python.
    connection.execute("""
        CREATE TABLE dgt_entrada AS SELECT PROVINCIA,MUNICIPIO,TIPO_DGT,PROPULSION,CATELECT,TIPO_DISTINTIVO,ordinality
        FROM read_csv(?,header=true,delim='|',encoding='latin-1',all_varchar=true,
            parallel=false,strict_mode=true,ignore_errors=false) WITH ORDINALITY
    """, [str(path)])
    national = connection.execute('SELECT count(*),min(ordinality),max(ordinality) FROM dgt_entrada').fetchone()
    if national[1:] != (1, national[0]):
        raise ValueError('Ordinales nacionales DGT incoherentes')
    connection.execute("""
        CREATE TABLE parque_detalle_observado AS SELECT ?::VARCHAR origen_id,
            ordinality::BIGINT registro_origen,PROVINCIA codigo_provincia,
            nullif(trim(MUNICIPIO),'') codigo_municipio,'2025-12'::VARCHAR periodo,
            TIPO_DGT tipo_dgt_codigo,PROPULSION propulsion_codigo,
            CATELECT categoria_electrica_codigo,TIPO_DISTINTIVO distintivo_codigo
        FROM dgt_entrada WHERE PROVINCIA IN ('03','12','46')
    """, [origin])
    bad = query(connection, """
        SELECT registro_origen,codigo_provincia,codigo_municipio FROM parque_detalle_observado
        WHERE codigo_municipio IS NOT NULL AND (substr(codigo_municipio,1,2)<>codigo_provincia
           OR codigo_municipio NOT IN (SELECT codigo_municipio FROM municipios)) LIMIT 10
    """)
    if bad:
        raise ValueError(f'Municipios publicados por DGT no admitidos: {bad}')
    return {'filas_nacionales': national[0], 'filas_fuera_ambito': connection.execute("SELECT count(*) FROM dgt_entrada WHERE PROVINCIA NOT IN ('03','12','46') OR PROVINCIA IS NULL").fetchone()[0]}


def check_key(connection, table, columns):
    key = ','.join(ident(column) for column in columns)
    nulls = ' OR '.join(f'{ident(column)} IS NULL' for column in columns)
    if connection.execute(f'SELECT count(*) FROM {ident(table)} WHERE {nulls}').fetchone()[0]:
        raise ValueError(f'Clave incompleta: {table}')
    if connection.execute(f'SELECT count(*) FROM (SELECT {key} FROM {ident(table)} GROUP BY {key} HAVING count(*)>1)').fetchone()[0]:
        raise ValueError(f'Clave repetida: {table}')


def validate(connection, codes):
    for table, key in KEYS.items():
        check_key(connection, table, key)
    park = query(connection, 'SELECT sum(parque_total) parque_total,sum(vehiculos_distintivo_0) distintivo_0,sum(electrificado_origen) electrificado FROM parque_municipal')[0]
    unspecified = query(connection, 'SELECT codigo_provincia,parque_total,vehiculos_distintivo_0,electrificado_origen FROM parque_sin_municipio ORDER BY 1')
    detail = query(connection, 'SELECT codigo_provincia,count(*) registros,count(*) FILTER(WHERE codigo_municipio IS NULL) sin_municipio,count(DISTINCT codigo_municipio) municipios_observados FROM parque_detalle_observado GROUP BY 1 ORDER BY 1')
    if sum(row['registros'] for row in detail) != park['parque_total'] + sum(row['parque_total'] for row in unspecified):
        raise ValueError('El detalle DGT no reconcilia con el total municipal y las agrupaciones PP000')
    disagreements = query(connection, """
        SELECT codigo_municipio,observados,parque_total FROM
        (SELECT codigo_municipio,count(*) observados FROM parque_detalle_observado
         WHERE codigo_municipio IS NOT NULL GROUP BY 1) d JOIN parque_municipal USING(codigo_municipio)
        WHERE observados<>parque_total
    """)
    if disagreements:
        raise ValueError(f'El detalle observado DGT no coincide con el parque municipal: {disagreements}')
    population = query(connection, """
        SELECT anio,sexo,count(*) filas,count(DISTINCT codigo_municipio) municipios,
            count(*) FILTER(WHERE estado_valor<>'informado') sin_valor
        FROM poblacion_municipal GROUP BY 1,2 ORDER BY 1,2
    """)
    if any(row['municipios'] != len(codes) for row in population):
        raise ValueError('Cobertura de población incompleta por año y sexo')
    sex_conflicts = query(connection, """
        SELECT codigo_municipio,anio FROM poblacion_municipal GROUP BY 1,2
        HAVING count(*)<>3 OR (count(poblacion)=3 AND
            max(poblacion) FILTER(WHERE sexo='Total')<>
            sum(poblacion) FILTER(WHERE sexo IN ('Hombres','Mujeres')))
    """)
    if sex_conflicts:
        raise ValueError(f'Población: Total no reconcilia con hombres y mujeres: {sex_conflicts[:10]}')
    income = query(connection, """
        SELECT nivel_territorial,anio,count(*) filas,count(DISTINCT codigo_municipio) municipios,
            count(*) FILTER(WHERE estado_valor='no_informado') sin_valor,
            count(*) FILTER(WHERE estado_valor='error_conversion') errores
        FROM renta_territorial GROUP BY 1,2 ORDER BY 1,2
    """)
    if any(row['municipios'] != len(codes) for row in income if row['nivel_territorial'] == 'municipio'):
        raise ValueError('Cobertura municipal de renta incompleta')
    return {'parque_municipal_totales': park, 'parque_sin_municipio': unspecified,
            'detalle_por_provincia': detail, 'municipios_detalle_reconciliados': True,
            'poblacion_por_anio_sexo': population, 'poblacion_total_reconciliada': True,
            'renta_por_nivel_anio': income,
            'categorias_electricas_detalle': query(connection, 'SELECT categoria_electrica_codigo,count(*) registros FROM parque_detalle_observado GROUP BY 1 ORDER BY 1'),
            'distintivos_detalle': query(connection, 'SELECT distintivo_codigo,count(*) registros FROM parque_detalle_observado GROUP BY 1 ORDER BY 1'),
            'incidencias_conversion': connection.execute('SELECT count(*) FROM incidencias').fetchone()[0]}


def municipal_base(connection, root):
    directory = root / 'resultados/silver/base_territorial'
    pointer = json.loads((directory / 'actual.json').read_text(encoding='utf-8'))
    report_path = (directory / pointer['informe']).resolve()
    if not report_path.is_relative_to(directory.resolve()):
        raise ValueError('Informe territorial fuera de su carpeta')
    report = json.loads(report_path.read_text(encoding='utf-8'))
    if report['estado'] != 'preparada':
        raise ValueError('La base territorial no consta como preparada')
    entry = next(item for item in report['salidas'] if item['archivo'] == 'municipios.parquet')
    path = report_path.parent / 'municipios.parquet'
    observed = fingerprint(path)
    if any(observed[key] != entry[key] for key in ('bytes', 'sha256')):
        raise ValueError('La tabla municipal no coincide con su informe Silver')
    connection.execute('CREATE TEMP TABLE municipios AS SELECT * FROM read_parquet(?)', [str(path)])
    codes = {row[0] for row in connection.execute('SELECT codigo_municipio FROM municipios').fetchall()}
    check_key(connection, 'municipios', ['codigo_municipio'])
    if len(codes) != 542 or any(not re.fullmatch('[0-9]{5}', code) for code in codes):
        raise ValueError('Referencia municipal Silver no válida')
    return codes, {'informe': report_path.relative_to(root.resolve()).as_posix(), 'informe_sha256': fingerprint(report_path)['sha256'],
                   'tabla': path.relative_to(root.resolve()).as_posix(), **observed}


def write_outputs(connection, target, origins):
    target.mkdir(exist_ok=False)
    outputs = []
    sources = {'parque_municipal': XLSX, 'parque_sin_municipio': XLSX,
               'parque_detalle_observado': TXT, 'poblacion_municipal': POP,
               'renta_territorial': INCOME, 'renta_codigos_historicos': INCOME}
    for table in [*KEYS, 'incidencias']:
        path = target / f'{table}.parquet'
        connection.execute(f'COPY {ident(table)} TO ? (FORMAT PARQUET,COMPRESSION ZSTD)', [str(path)])
        expected_schema = connection.execute(f'DESCRIBE {ident(table)}').fetchall()
        schema = connection.execute('DESCRIBE SELECT * FROM read_parquet(?)', [str(path)]).fetchall()
        if [(row[0], row[1]) for row in expected_schema] != [(row[0], row[1]) for row in schema]:
            raise ValueError(f'Tipos no conservados al exportar: {table}')
        changed = connection.execute(f"SELECT count(*) FROM ((SELECT * FROM {ident(table)} EXCEPT ALL SELECT * FROM read_parquet(?)) UNION ALL (SELECT * FROM read_parquet(?) EXCEPT ALL SELECT * FROM {ident(table)}))", [str(path), str(path)]).fetchone()[0]
        if changed:
            raise ValueError(f'Valores no conservados al exportar: {table}')
        outputs.append({'archivo': path.name, 'filas': connection.execute(f'SELECT count(*) FROM {ident(table)}').fetchone()[0],
                        'origen_id': origins.get(sources.get(table)), 'clave': KEYS.get(table),
                        'esquema': [{'campo': row[0], 'tipo': row[1]} for row in schema], **fingerprint(path)})
    return outputs


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--manifiesto-bronze', type=Path, required=True)
    args = parser.parse_args()
    manifest_path = args.manifiesto_bronze.resolve()
    manifest_stamp = fingerprint(manifest_path)
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    print('Verificando las cinco entradas Bronze de la fase 3...', flush=True)
    inputs = verify_inputs(ROOT, manifest, INPUTS)
    origins = {row['ruta']: row['recurso_id'] + ':' + row['sha256'] for row in inputs}
    directory = ROOT / 'resultados/silver/demografia_parque'
    directory.mkdir(parents=True, exist_ok=True)
    run = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%S%fZ')
    target = directory / run
    with tempfile.TemporaryDirectory(dir=directory) as temporary:
        temporary = Path(temporary)
        with duckdb.connect(str(temporary / 'trabajo.duckdb'), config={'threads': 2, 'memory_limit': '1GB', 'temp_directory': str(temporary / 'desborde')}) as connection:
            codes, base = municipal_base(connection, ROOT)
            park = prepare_municipal_park(connection, ROOT / XLSX, codes)
            print('Preparando población y renta mediante lectura secuencial...', flush=True)
            ine = prepare_ine(connection, ROOT / POP, ROOT / INCOME, codes, temporary, origins)
            print('Procesando el TXT DGT: atributos seleccionados en una base temporal en disco...', flush=True)
            detail = prepare_detail(connection, ROOT / TXT, origins[TXT])
            print('Validando claves, cobertura y cifras de control...', flush=True)
            profile = validate(connection, codes)
            print('Exportando y releyendo las tablas Parquet...', flush=True)
            outputs = write_outputs(connection, target, origins)
            runtime = connection.execute('SELECT version()').fetchone()[0]
    print('Comprobando que las entradas no han cambiado...', flush=True)
    if verify_inputs(ROOT, manifest, INPUTS) != inputs or fingerprint(manifest_path) != manifest_stamp:
        raise RuntimeError('Las entradas Bronze cambiaron durante el proceso')
    base_path = ROOT / base['tabla']
    if fingerprint(base_path)['sha256'] != base['sha256'] or fingerprint(ROOT / base['informe'])['sha256'] != base['informe_sha256']:
        raise RuntimeError('La referencia territorial cambió durante el proceso')
    report = {'version_reglas': 1, 'ejecucion': run,
              'estado': 'preparada_con_incidencias' if profile['incidencias_conversion'] else 'preparada',
              'fecha_proceso_utc': datetime.now(timezone.utc).isoformat(),
              'manifiesto_bronze': {'ruta': manifest_path.relative_to(ROOT.resolve()).as_posix(), **manifest_stamp},
              'base_territorial': base, 'entradas_verificadas': inputs,
              'origenes': origins, 'lectura': {'dgt_municipal': park, 'ine': ine, 'dgt_detalle': detail},
              'calidad': profile, 'salidas': outputs, 'duckdb': runtime,
              'codigo': [{'ruta': 'src/' + name, 'sha256': fingerprint(ROOT / 'src' / name)['sha256']}
                         for name in ('preparar_demografia_parque.py', 'preparar_base_territorial.py', 'catalogar_bronze.py', 'auditar_fuentes.py')],
              'limitaciones': ['Distintivo 0 es una aproximación, no el desglose BEV/PHEV.',
                               'El detalle DGT no publica municipio en todos los registros ni un identificador individual de vehículo.',
                               'Los símbolos de renta sin causa acreditada se conservan como no informados; no se imputan.',
                               'Se conservan los periodos por fuente sin fijar un año común para Gold.']}
    (target / 'informe.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    pointer = directory / 'actual.tmp'
    pointer.write_text(json.dumps({'ejecucion': run, 'informe': f'{run}/informe.json'}, indent=2) + '\n', encoding='utf-8')
    pointer.replace(directory / 'actual.json')
    print(f'Fase 3: {report["estado"]}. Informe: {target.relative_to(ROOT).as_posix()}/informe.json')


if __name__ == '__main__':
    main()
