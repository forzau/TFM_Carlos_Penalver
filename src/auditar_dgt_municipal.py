"""Comprueba la cobertura del Excel municipal DGT; no crea Silver ni modifica Fuentes/."""

import csv
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import urlopen

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
URL = 'https://www.dgt.es/export/sites/web-DGT/.galleries/downloads/dgt-en-cifras/informacion_municipal/Datos-municipales-general/DatosMunicipalesGeneral_2025.xlsx'
FIELDS = ('Parque Total', 'Electrificado', 'Distintivo 0')
PROVINCES = ('03', '12', '46')


def audit_workbook(path, population_path):
    with population_path.open(encoding='utf-8-sig', newline='') as stream:
        expected = {r['Municipios'][:5] for r in csv.DictReader(stream, delimiter=';')
                    if r['Municipios'][:2] in PROVINCES}
    if not expected or any(len(code) != 5 or not code.isdigit() for code in expected):
        raise ValueError('La referencia municipal de población está vacía o tiene códigos no válidos')
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook['DatosMunicipalesGeneral_2025']
        iterator = sheet.iter_rows(values_only=True)
        headers = list(next(iterator))
        if len(set(headers)) != len(headers):
            raise ValueError('Cabeceras repetidas en DGT')
        required = ('Código INE', 'Municipio') + FIELDS
        if not set(required).issubset(headers):
            raise ValueError('El esquema de DGT no contiene las columnas esperadas')
        counts = Counter()
        selected, unspecified = [], []
        national_rows = 0
        for values in iterator:
            if all(v is None for v in values):
                continue
            national_rows += 1
            row = dict(zip(headers, values))
            code = row['Código INE']
            if not isinstance(code, str) or len(code) != 5 or not code.isdigit():
                raise ValueError(f'Código INE no válido en fila {national_rows + 1}')
            counts[code] += 1
            if code[:2] not in PROVINCES:
                continue
            if code in {p + '000' for p in PROVINCES}:
                unspecified.append({k: row[k] for k in required})
            else:
                selected.append({k: row[k] for k in required})
        actual = {r['Código INE'] for r in selected}
        checks = {}
        for field in FIELDS:
            values = [r[field] for r in selected]
            invalid = [v for v in values if isinstance(v, bool) or not isinstance(v, (int, float)) or not (v >= 0 and v == int(v))]
            checks[field] = {'no_validos': len(invalid), 'total_municipios': sum(values) if not invalid else None}
        return {'filas_nacionales': national_rows, 'columnas': headers,
                'codigos_duplicados': {code: n for code, n in counts.items() if n > 1},
                'municipios_referencia': len(expected), 'municipios_cv': len(selected),
                'faltan': sorted(expected - actual), 'adicionales': sorted(actual - expected),
                'municipios_por_provincia': dict(Counter(r['Código INE'][:2] for r in selected)),
                'filas_sin_municipio': unspecified, 'indicadores': checks,
                'desglose_bev_phev_publicado': False,
                'advertencia': 'Electrificado y Distintivo 0 no acreditan un desglose municipal BEV/PHEV.'}
    finally:
        workbook.close()


def main():
    directory = ROOT / 'resultados/auditoria_dgt'
    path = ROOT / 'datos/bronze/dgt_municipal/2025/DatosMunicipalesGeneral_2025.xlsx'
    acquisition = None
    if not path.exists():
        path.parent.mkdir(parents=True, exist_ok=True)
        # Guardar solo tras completar la descarga; nunca sobrescribir una copia existente.
        with urlopen(URL, timeout=60) as response:
            content = response.read(20 * 1024 * 1024 + 1)
            if len(content) > 20 * 1024 * 1024 or not content.startswith(b'PK'):
                raise ValueError('La respuesta no es el XLSX esperado o supera 20 MB')
        with path.open('xb') as stream:
            stream.write(content)
        acquisition = datetime.now(timezone.utc).isoformat()
    before = hashlib.sha256(path.read_bytes()).hexdigest()
    result = audit_workbook(path, ROOT / 'Fuentes/Poblacion.csv')
    result.update({'url': URL, 'recurso': path.relative_to(ROOT).as_posix(),
                   'referencia_parque': '2025-12-31', 'fecha_inspeccion_utc': datetime.now(timezone.utc).isoformat(),
                   'fecha_descarga_en_esta_ejecucion': acquisition, 'sha256': before})
    if hashlib.sha256(path.read_bytes()).hexdigest() != before:
        raise RuntimeError('El original descargado cambió durante la inspección')
    directory.mkdir(parents=True, exist_ok=True)
    (directory / 'perfil_municipal_2025.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({k: result[k] for k in ['municipios_cv', 'faltan', 'adicionales', 'municipios_por_provincia', 'indicadores']}, ensure_ascii=False))
    if result['faltan'] or result['adicionales'] or result['codigos_duplicados'] or any(x['no_validos'] for x in result['indicadores'].values()):
        raise SystemExit(1)


if __name__ == '__main__':
    main()
